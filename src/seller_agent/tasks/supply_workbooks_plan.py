from __future__ import annotations

import csv
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from seller_agent.config import AppCredentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.ozon.adapter import OzonSellerAdapter
from seller_agent.marketplaces.wb.adapter import WbContentAdapter
from seller_agent.marketplaces.wb.fbw_supplies_adapter import WbFbwSuppliesAdapter
from seller_agent.marketplaces.wb.statistics_adapter import WbStatisticsAdapter
from seller_agent.reports.writer import ensure_dir, write_json


ACTIVE_OZON_SUPPLY_STATES = [
    "DATA_FILLING",
    "READY_TO_SUPPLY",
    "ACCEPTED_AT_SUPPLY_WAREHOUSE",
    "IN_TRANSIT",
    "ACCEPTANCE_AT_STORAGE_WAREHOUSE",
    "REPORTS_CONFIRMATION_AWAITING",
]
ACTIVE_WB_SUPPLY_STATUS_IDS = {2, 3, 4, 6}
WB_REGION_MAP = {
    "Центральный федеральный округ": "Центральный",
    "Приволжский федеральный округ": "Поволжье",
    "Северо-Западный федеральный округ": "Северо-Запад",
    "Сибирский федеральный округ": "Сибирь/ДВ",
    "Дальневосточный федеральный округ": "Сибирь/ДВ",
    "Уральский федеральный округ": "Урал",
    "Южный федеральный округ": "Юг/СК",
    "Северо-Кавказский федеральный округ": "Юг/СК",
}
MANUFACTURED_INCLUDE_RE = re.compile(
    r"(шеврон|нашив|петлиц|позывн|patch|патч|chev|nash|loop|pict|form|back|pz)",
    re.IGNORECASE,
)
MANUFACTURED_EXCLUDE_RE = re.compile(
    r"(головн|кепк[аи]|шапк|панам|подсум|фартук|фальшпогон|брелок)",
    re.IGNORECASE,
)


@dataclass
class PlanItem:
    marketplace: str
    cluster: str
    article: str
    sku: str
    barcode: str
    title: str
    pack_qty: int
    sales_90: int
    global_sales_90: int
    stock: int
    inbound: int
    score: float
    quantity: int = 0

    @property
    def physical(self) -> int:
        return self.quantity * self.pack_qty


def run_supply_workbooks_plan(
    *,
    data_dir: Path = Path("data"),
    run_id: str | None = None,
    target_days: int = 30,
    cycle_days: int = 5,
    daily_capacity: int = 200,
    ozon_weekly_physical: int = 1400,
    wb_target_physical: int = 1200,
    credentials: AppCredentials | None = None,
    today: date | None = None,
    ozon_adapter: OzonSellerAdapter | None = None,
    wb_stats_adapter: WbStatisticsAdapter | None = None,
    wb_content_adapter: WbContentAdapter | None = None,
    wb_supplies_adapter: WbFbwSuppliesAdapter | None = None,
) -> dict[str, Any]:
    started_at = datetime.now().astimezone()
    today = today or started_at.date()
    period_to = today - timedelta(days=1)
    period_from = period_to - timedelta(days=89)
    run_id = run_id or f"supply_workbooks_plan_{started_at.strftime('%Y%m%dT%H%M%S')}"
    run_dir = ensure_dir(data_dir / "runs" / today.isoformat() / run_id)
    raw_dir = ensure_dir(run_dir / "raw")
    processed_dir = ensure_dir(run_dir / "processed")
    report_path = run_dir / "supply_workbooks_report.md"
    summary_path = run_dir / "summary.json"

    inputs = {
        "target_days": target_days,
        "cycle_days": cycle_days,
        "daily_capacity": daily_capacity,
        "cycle_capacity_physical_pieces": cycle_days * daily_capacity,
        "ozon_weekly_physical": ozon_weekly_physical,
        "wb_target_physical": wb_target_physical,
        "period_from": period_from.isoformat(),
        "period_to": period_to.isoformat(),
        "period_days": 90,
    }
    artifacts = {
        "run_dir": str(run_dir),
        "report": str(report_path),
        "summary": str(summary_path),
        "run_manifest": str(run_dir / "manifest.json"),
    }
    errors: dict[str, str] = {}
    warnings: list[str] = []

    if credentials is None and not all((ozon_adapter, wb_stats_adapter, wb_content_adapter, wb_supplies_adapter)):
        result = _blocked_result(
            run_id=run_id,
            started_at=started_at,
            inputs=inputs,
            artifacts=artifacts,
            reason="missing_credentials_or_adapters",
        )
        _write_outputs(data_dir, run_dir, report_path, summary_path, result, processed_dir)
        return result

    unified_index = _load_unified_index(data_dir / "catalog/unified/products.csv")
    if not unified_index:
        errors["unified_catalog"] = "data/catalog/unified/products.csv is missing or empty"

    ozon: dict[str, Any] = {"status": "skipped"}
    wb: dict[str, Any] = {"status": "skipped"}
    ozon_plan_1: list[PlanItem] = []
    ozon_plan_2: list[PlanItem] = []
    wb_plan: list[PlanItem] = []

    if ozon_adapter or (credentials and credentials.ozon_seller):
        try:
            ozon_source = ozon_adapter or OzonSellerAdapter(credentials.ozon_seller)  # type: ignore[arg-type]
            ozon = _build_ozon_plan(
                ozon=ozon_source,
                unified_index=unified_index,
                raw_dir=raw_dir,
                period_from=period_from,
                period_to=period_to,
                target_physical=ozon_weekly_physical,
            )
            ozon_plan_1 = ozon["week1"]
            ozon_plan_2 = ozon["week2"]
        except Exception as exc:  # noqa: BLE001 - source failures must be reported, not hidden
            errors["ozon"] = _safe_error(exc)
    else:
        warnings.append("Ozon credentials/adapters are missing; Ozon supply workbook not generated.")

    if wb_stats_adapter or (credentials and credentials.wb):
        try:
            wb_stats = wb_stats_adapter or WbStatisticsAdapter(credentials.wb)  # type: ignore[arg-type]
            wb_content = wb_content_adapter or WbContentAdapter(credentials.wb)  # type: ignore[arg-type]
            wb_supplies = wb_supplies_adapter or WbFbwSuppliesAdapter(credentials.wb)  # type: ignore[arg-type]
            wb = _build_wb_plan(
                wb_stats=wb_stats,
                wb_content=wb_content,
                wb_supplies=wb_supplies,
                unified_index=unified_index,
                raw_dir=raw_dir,
                period_from=period_from,
                target_physical=wb_target_physical,
            )
            wb_plan = wb["plan"]
        except Exception as exc:  # noqa: BLE001 - source failures must be reported, not hidden
            errors["wb"] = _safe_error(exc)
    else:
        warnings.append("WB credentials/adapters are missing; WB supply workbook not generated.")

    if ozon_plan_1:
        path = run_dir / f"v_rabotu_ozon_week1_{sum(item.physical for item in ozon_plan_1)}_{today.isoformat()}.xlsx"
        _make_workbook(path, ozon_plan_1, "Ozon", "Неделя 1; read-only расчет, поставка в ЛК не создавалась.")
        artifacts["ozon_week1_xlsx"] = str(path)
        _write_csv_rows(processed_dir / "ozon_week1_cluster_rows.csv", _rows_from_plan(ozon_plan_1))
    if ozon_plan_2:
        path = run_dir / f"v_rabotu_ozon_week2_{sum(item.physical for item in ozon_plan_2)}_{today.isoformat()}.xlsx"
        _make_workbook(path, ozon_plan_2, "Ozon", "Неделя 2; week1 учтена как planned inbound.")
        artifacts["ozon_week2_xlsx"] = str(path)
        _write_csv_rows(processed_dir / "ozon_week2_cluster_rows.csv", _rows_from_plan(ozon_plan_2))
    if wb_plan:
        path = run_dir / f"v_rabotu_wb_{sum(item.physical for item in wb_plan)}_{today.isoformat()}.xlsx"
        _make_workbook(path, wb_plan, "WB", "Read-only расчет; поставка в ЛК не создавалась.")
        artifacts["wb_xlsx"] = str(path)
        _write_csv_rows(processed_dir / "wb_cluster_rows.csv", _rows_from_plan(wb_plan))

    source_status = {
        "current_stocks": _source_ok(ozon, wb, "stocks"),
        "sales_90_days": _source_ok(ozon, wb, "sales"),
        "sales_localization": _source_ok(ozon, wb, "localization"),
        "active_inbound_supplies": _source_ok(ozon, wb, "inbound"),
        "destination_cluster_mapping": _source_ok(ozon, wb, "destination_cluster"),
        "production_constraints": "ok",
    }
    generated_files = [value for key, value in artifacts.items() if key.endswith("_xlsx")]
    overall_status = (
        "error"
        if errors and not generated_files
        else "warning"
        if errors or warnings or any(value != "ok" for value in source_status.values())
        else "ok"
    )
    result = {
        "run_id": run_id,
        "started_at": started_at.isoformat(timespec="seconds"),
        "mode": "read_only",
        "overall_status": overall_status,
        "errors": errors,
        "warnings": warnings,
        "inputs": inputs,
        "source_status": source_status,
        "summary": {
            "period_from": period_from.isoformat(),
            "period_to": period_to.isoformat(),
            "ozon_week1_physical": sum(item.physical for item in ozon_plan_1),
            "ozon_week1_units": sum(item.quantity for item in ozon_plan_1),
            "ozon_week1_articles": len({item.article for item in ozon_plan_1}),
            "ozon_week2_physical": sum(item.physical for item in ozon_plan_2),
            "ozon_week2_units": sum(item.quantity for item in ozon_plan_2),
            "ozon_week2_articles": len({item.article for item in ozon_plan_2}),
            "wb_physical": sum(item.physical for item in wb_plan),
            "wb_units": sum(item.quantity for item in wb_plan),
            "wb_articles": len({item.article for item in wb_plan}),
            "generated_files": len(generated_files),
        },
        "ozon": _plan_summary(ozon, ozon_plan_1, ozon_plan_2),
        "wb": _plan_summary(wb, wb_plan),
        "artifacts": artifacts,
        "next_step": "Проверить Excel-файлы владельцем; создание поставок в ЛК не выполнялось.",
    }
    _write_outputs(data_dir, run_dir, report_path, summary_path, result, processed_dir)
    return result


def _build_ozon_plan(
    *,
    ozon: OzonSellerAdapter,
    unified_index: dict[str, dict[str, str]],
    raw_dir: Path,
    period_from: date,
    period_to: date,
    target_physical: int,
) -> dict[str, Any]:
    product_list = ozon.fetch_product_list(visibility="ALL")
    product_ids = [str(row.get("product_id")) for row in product_list if row.get("product_id")]
    info_rows = ozon.fetch_product_info(product_ids)
    stock_rows = ozon.fetch_product_stocks(product_ids)
    warehouse_rows = ozon.fetch_stock_on_warehouses()
    postings = ozon.fetch_fbo_postings(
        since=_period_start(period_from),
        to=_period_end(period_to),
        status="",
        limit=100,
    )
    inbound = _fetch_ozon_inbound(ozon)

    write_json(raw_dir / "ozon_product_list.json", product_list)
    write_json(raw_dir / "ozon_product_info.json", info_rows)
    write_json(raw_dir / "ozon_product_stocks.json", stock_rows)
    write_json(raw_dir / "ozon_stock_on_warehouses.json", warehouse_rows)
    write_json(raw_dir / "ozon_fbo_postings_90d.json", postings)
    write_json(raw_dir / "ozon_active_supply_orders.json", inbound["raw"])

    title_by_offer, sku_by_offer = _ozon_titles_and_skus(product_list, info_rows)
    stock_by_offer = _ozon_fbo_stock_by_offer(stock_rows)
    promised_by_offer = _ozon_promised_by_offer(warehouse_rows)
    inbound_by_offer = inbound["by_offer"]
    inbound_by_offer_cluster = inbound["by_offer_cluster"]

    cluster_sales: Counter[str] = Counter()
    product_cluster_sales: Counter[tuple[str, str]] = Counter()
    global_sales: Counter[str] = Counter()
    for posting in postings:
        if str(posting.get("status") or "").lower() == "cancelled":
            continue
        cluster = str((posting.get("financial_data") or {}).get("cluster_to") or "").strip()
        if not cluster:
            cluster = str((posting.get("analytics_data") or {}).get("warehouse_name") or "Не указан").strip()
        for product in posting.get("products") or []:
            offer = str(product.get("offer_id") or "").strip()
            if not offer:
                continue
            qty = _safe_int(product.get("quantity")) or 1
            title = title_by_offer.get(offer) or str(product.get("name") or "")
            row = unified_index.get(offer) or unified_index.get(str(product.get("sku") or ""))
            if not _is_manufactured(title, offer, product_group=(row or {}).get("product_group", "")):
                continue
            cluster_sales[cluster] += qty
            product_cluster_sales[(offer, cluster)] += qty
            global_sales[offer] += qty

    top_clusters = [cluster for cluster, _ in cluster_sales.most_common(5)]
    items: list[PlanItem] = []
    for (offer, cluster), sales_qty in product_cluster_sales.items():
        if cluster not in top_clusters:
            continue
        adjusted_sales = max(1, sales_qty - inbound_by_offer_cluster[(offer, cluster)])
        title = title_by_offer.get(offer, "")
        row = unified_index.get(offer) or unified_index.get(sku_by_offer.get(offer, ""))
        pack_qty = _infer_pack_qty(title, (row or {}).get("pack_qty", 1))
        stock = stock_by_offer[offer]
        inbound_qty = inbound_by_offer[offer] + promised_by_offer[offer]
        avg = global_sales[offer] / 90
        coverage_days = (stock + inbound_qty) / avg if avg else 999
        stock_pressure = max(0.25, min(3.0, (45 - min(45, coverage_days)) / 15 + 0.5))
        score = adjusted_sales * stock_pressure + max(0, 30 - coverage_days) * 0.1
        items.append(
            PlanItem(
                marketplace="Ozon",
                cluster=cluster,
                article=offer,
                sku=sku_by_offer.get(offer, ""),
                barcode="",
                title=title,
                pack_qty=pack_qty,
                sales_90=adjusted_sales,
                global_sales_90=global_sales[offer],
                stock=stock,
                inbound=inbound_qty,
                score=score,
            )
        )

    week1 = _allocate_physical([PlanItem(**asdict(item)) for item in items], target_physical)
    week1_by_offer = Counter({item.article: item.quantity for item in week1})
    week2_items: list[PlanItem] = []
    for item in items:
        adjusted = PlanItem(**asdict(item))
        adjusted.inbound += week1_by_offer[item.article]
        avg = adjusted.global_sales_90 / 90
        coverage_days = (adjusted.stock + adjusted.inbound) / avg if avg else 999
        stock_pressure = max(0.2, min(2.5, (60 - min(60, coverage_days)) / 20 + 0.35))
        adjusted.score = adjusted.sales_90 * stock_pressure + max(0, 45 - coverage_days) * 0.08
        week2_items.append(adjusted)
    week2 = _allocate_physical(week2_items, target_physical)
    return {
        "status": "ok",
        "source_status": {
            "stocks": "ok",
            "sales": "ok",
            "localization": "ok",
            "inbound": inbound["status"],
            "destination_cluster": "ok" if top_clusters else "warning",
        },
        "top_clusters": top_clusters,
        "cluster_sales": dict(cluster_sales),
        "week1": week1,
        "week2": week2,
        "inbound_summary": inbound["summary"],
        "sources": {
            "product_list": "/v3/product/list",
            "product_info": "/v3/product/info/list",
            "product_stocks": "/v4/product/info/stocks",
            "warehouse_stocks": "/v2/analytics/stock_on_warehouses",
            "sales": "/v2/posting/fbo/list",
            "active_supplies": "supply-order list/get/details/bundle",
        },
    }


def _build_wb_plan(
    *,
    wb_stats: WbStatisticsAdapter,
    wb_content: WbContentAdapter,
    wb_supplies: WbFbwSuppliesAdapter,
    unified_index: dict[str, dict[str, str]],
    raw_dir: Path,
    period_from: date,
    target_physical: int,
) -> dict[str, Any]:
    cards = wb_content.fetch_cards(limit=100)
    stocks = wb_stats.fetch_stocks_legacy(date_from="2019-01-01")
    sales = wb_stats.fetch_sales(date_from=period_from.isoformat(), flag=0)
    inbound = _fetch_wb_inbound(wb_supplies)

    write_json(raw_dir / "wb_cards.json", cards)
    write_json(raw_dir / "wb_stocks.json", stocks)
    write_json(raw_dir / "wb_sales_90d.json", sales)
    write_json(raw_dir / "wb_active_supplies.json", inbound["raw"])

    title_by_vendor: dict[str, str] = {}
    nm_by_vendor: dict[str, str] = {}
    barcode_by_vendor: dict[str, str] = {}
    subject_by_vendor: dict[str, str] = {}
    for card in cards:
        vendor = str(card.get("vendorCode") or "").strip()
        if not vendor:
            continue
        title_by_vendor[vendor] = str(card.get("title") or "").strip()
        nm_by_vendor[vendor] = str(card.get("nmID") or card.get("nmId") or "").strip()
        subject_by_vendor[vendor] = str(card.get("subjectName") or card.get("object") or "").strip()
        sizes = card.get("sizes") or []
        if sizes and isinstance(sizes, list):
            skus = sizes[0].get("skus") if isinstance(sizes[0], dict) else []
            if skus:
                barcode_by_vendor[vendor] = str(skus[0]).strip()

    stock_by_vendor: Counter[str] = Counter()
    for row in stocks:
        vendor = str(row.get("supplierArticle") or "").strip()
        if not vendor:
            continue
        stock_by_vendor[vendor] += _safe_int(row.get("quantity"))
        barcode = str(row.get("barcode") or "").strip()
        if barcode:
            barcode_by_vendor.setdefault(vendor, barcode)
        title_by_vendor.setdefault(vendor, str(row.get("subject") or ""))
        subject_by_vendor.setdefault(vendor, str(row.get("subject") or ""))

    cluster_sales: Counter[str] = Counter()
    product_cluster_sales: Counter[tuple[str, str]] = Counter()
    global_sales: Counter[str] = Counter()
    for row in sales:
        vendor = str(row.get("supplierArticle") or "").strip()
        if not vendor:
            continue
        title = title_by_vendor.get(vendor, "")
        unified = unified_index.get(vendor) or unified_index.get(str(row.get("nmId") or ""))
        subject = subject_by_vendor.get(vendor) or str(row.get("subject") or "")
        if not _is_manufactured(title, vendor, subject=subject, product_group=(unified or {}).get("product_group", "")):
            continue
        cluster = _wb_region_from_sale(row)
        cluster_sales[cluster] += 1
        product_cluster_sales[(vendor, cluster)] += 1
        global_sales[vendor] += 1

    top_clusters = [cluster for cluster, _ in cluster_sales.most_common(4)]
    items: list[PlanItem] = []
    for (vendor, cluster), sales_qty in product_cluster_sales.items():
        if cluster not in top_clusters:
            continue
        adjusted_sales = max(1, sales_qty - inbound["by_vendor_region"][(vendor, cluster)])
        title = title_by_vendor.get(vendor, "")
        unified = unified_index.get(vendor) or unified_index.get(nm_by_vendor.get(vendor, ""))
        pack_qty = _infer_pack_qty(title, (unified or {}).get("pack_qty", 1))
        stock = stock_by_vendor[vendor]
        inbound_qty = inbound["by_vendor"][vendor]
        avg = global_sales[vendor] / 90
        coverage_days = (stock + inbound_qty) / avg if avg else 999
        stock_pressure = max(0.25, min(3.0, (45 - min(45, coverage_days)) / 15 + 0.5))
        score = adjusted_sales * stock_pressure + max(0, 30 - coverage_days) * 0.1
        items.append(
            PlanItem(
                marketplace="WB",
                cluster=cluster,
                article=vendor,
                sku=nm_by_vendor.get(vendor, ""),
                barcode=barcode_by_vendor.get(vendor, ""),
                title=title,
                pack_qty=pack_qty,
                sales_90=adjusted_sales,
                global_sales_90=global_sales[vendor],
                stock=stock,
                inbound=inbound_qty,
                score=score,
            )
        )

    plan = _allocate_physical(items, target_physical, cluster_min_physical=200)
    return {
        "status": "ok",
        "source_status": {
            "stocks": "ok",
            "sales": "ok",
            "localization": "ok",
            "inbound": inbound["status"],
            "destination_cluster": "ok" if top_clusters else "warning",
        },
        "top_clusters": top_clusters,
        "cluster_sales": dict(cluster_sales),
        "plan": plan,
        "inbound_summary": inbound["summary"],
        "sources": {
            "cards": "/content/v2/get/cards/list",
            "stocks": "/api/v1/supplier/stocks",
            "sales": "/api/v1/supplier/sales?flag=0",
            "active_supplies": "supplies-api.wildberries.ru /api/v1/supplies + goods",
        },
    }


def _fetch_ozon_inbound(ozon: OzonSellerAdapter) -> dict[str, Any]:
    by_offer: Counter[str] = Counter()
    by_offer_cluster: Counter[tuple[str, str]] = Counter()
    raw: dict[str, Any] = {"orders": [], "details": [], "bundles": [], "errors": []}
    try:
        order_ids = ozon.fetch_supply_order_ids(states=ACTIVE_OZON_SUPPLY_STATES)
        orders = ozon.fetch_supply_orders(order_ids) if order_ids else []
        raw["orders"] = orders
        for order in orders:
            cluster = str(order.get("macrolocal_cluster_name") or order.get("macrolocal_cluster_id") or "").strip()
            for supply in order.get("supplies") or []:
                if not isinstance(supply, dict):
                    continue
                bundle_id = supply.get("bundle_id") or supply.get("bundleId")
                supply_id = supply.get("supply_id") or supply.get("supplyId")
                if not bundle_id and supply_id:
                    detail = ozon.fetch_supply_order_details(supply_id)
                    raw["details"].append(detail)
                    content = detail.get("content") if isinstance(detail, dict) else {}
                    bundle_id = (content or {}).get("bundle_id") or (content or {}).get("bundleId")
                    cluster = cluster or str(detail.get("storage_warehouse") or "").strip()
                if not bundle_id:
                    continue
                bundle_rows = ozon.fetch_supply_order_bundle(bundle_id)
                raw["bundles"].append({"bundle_id": bundle_id, "rows": bundle_rows})
                for row in bundle_rows:
                    offer = str(row.get("offer_id") or "").strip()
                    qty = _safe_int(row.get("quantity"))
                    if offer and qty > 0:
                        by_offer[offer] += qty
                        by_offer_cluster[(offer, cluster or "Не указан")] += qty
        status = "ok"
    except Exception as exc:  # noqa: BLE001
        raw["errors"].append(_safe_error(exc))
        status = "warning"
    return {
        "status": status,
        "by_offer": by_offer,
        "by_offer_cluster": by_offer_cluster,
        "raw": raw,
        "summary": {
            "orders": len(raw["orders"]),
            "bundle_groups": len(raw["bundles"]),
            "inbound_units": sum(by_offer.values()),
            "errors": raw["errors"],
        },
    }


def _fetch_wb_inbound(wb_supplies: WbFbwSuppliesAdapter) -> dict[str, Any]:
    by_vendor: Counter[str] = Counter()
    by_vendor_region: Counter[tuple[str, str]] = Counter()
    raw: dict[str, Any] = {"supplies": [], "details": [], "goods": [], "errors": []}
    try:
        supplies = wb_supplies.fetch_supplies()
        raw["supplies"] = supplies
        for supply in supplies:
            status_id = _safe_int(supply.get("statusID") or supply.get("statusId"))
            if status_id not in ACTIVE_WB_SUPPLY_STATUS_IDS:
                continue
            supply_id = supply.get("supplyID") or supply.get("supplyId")
            if not supply_id:
                continue
            detail = wb_supplies.fetch_supply(supply_id)
            goods = wb_supplies.fetch_supply_goods(supply_id)
            raw["details"].append(detail)
            raw["goods"].append({"supply_id": supply_id, "rows": goods})
            region = _normalize_wb_region(
                str(detail.get("warehouseName") or detail.get("actualWarehouseName") or supply.get("warehouseName") or "")
            )
            for row in goods:
                vendor = str(row.get("vendorCode") or "").strip()
                qty = _safe_int(row.get("quantity"))
                ready = _safe_int(row.get("readyForSaleQuantity"))
                inbound_qty = max(0, qty - ready)
                if vendor and inbound_qty > 0:
                    by_vendor[vendor] += inbound_qty
                    by_vendor_region[(vendor, region or "Не указан")] += inbound_qty
        status = "ok"
    except Exception as exc:  # noqa: BLE001
        raw["errors"].append(_safe_error(exc))
        status = "warning"
    return {
        "status": status,
        "by_vendor": by_vendor,
        "by_vendor_region": by_vendor_region,
        "raw": raw,
        "summary": {
            "supplies": len(raw["supplies"]),
            "active_goods_groups": len(raw["goods"]),
            "inbound_units": sum(by_vendor.values()),
            "errors": raw["errors"],
        },
    }


def _allocate_physical(items: list[PlanItem], target_physical: int, cluster_min_physical: int = 0) -> list[PlanItem]:
    if not items or target_physical <= 0:
        return []
    cluster_items: dict[str, list[PlanItem]] = defaultdict(list)
    for item in items:
        cluster_items[item.cluster].append(item)
    clusters = list(cluster_items)
    cluster_sales = {cluster: sum(max(1, item.sales_90) for item in rows) for cluster, rows in cluster_items.items()}
    total_sales = sum(cluster_sales.values()) or len(clusters)
    cluster_targets: dict[str, int] = {}
    remainder = target_physical
    for cluster in clusters:
        raw = target_physical * cluster_sales[cluster] / total_sales
        rounded = int(round(raw / 8.0) * 8)
        rounded = max(cluster_min_physical, rounded)
        cluster_targets[cluster] = rounded
        remainder -= rounded
    while remainder != 0:
        if remainder > 0:
            cluster = max(clusters, key=lambda item: cluster_sales[item])
            cluster_targets[cluster] += 8
            remainder -= 8
        else:
            candidates = [cluster for cluster in clusters if cluster_targets[cluster] - 8 >= cluster_min_physical]
            if not candidates:
                break
            cluster = max(candidates, key=lambda item: cluster_targets[item])
            cluster_targets[cluster] -= 8
            remainder += 8

    allocated: list[PlanItem] = []
    for cluster, rows in cluster_items.items():
        rows = sorted(rows, key=lambda item: (item.score, item.sales_90, item.global_sales_90), reverse=True)
        target = cluster_targets[cluster]
        total_score = sum(max(0.1, row.score) for row in rows)
        used = 0
        for row in rows:
            raw_physical = target * max(0.1, row.score) / total_score
            units = max(0, int(round(raw_physical / row.pack_qty / 8.0) * 8))
            row.quantity = units
            used += row.physical
        guard = 0
        while used < target and guard < 2000:
            guard += 1
            row = rows[(guard - 1) % len(rows)]
            row.quantity += 8
            used += 8 * row.pack_qty
        guard = 0
        while used > target and guard < 2000:
            guard += 1
            candidates = [row for row in rows if row.quantity >= 8 and used - 8 * row.pack_qty >= target]
            if not candidates:
                break
            row = min(candidates, key=lambda item: (item.score, item.sales_90))
            row.quantity -= 8
            used -= 8 * row.pack_qty
        allocated.extend(row for row in rows if row.quantity > 0)
    return [row for row in allocated if row.quantity > 0]


def _make_workbook(path: Path, items: list[PlanItem], marketplace: str, note: str) -> None:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Артикулы"
    totals: dict[str, dict[str, Any]] = {}
    for item in items:
        row = totals.setdefault(
            item.article,
            {
                "Маркетплейс": marketplace,
                "Артикул/SKU": item.article,
                "SKU/nmID": item.sku,
                "Баркод ВБ": item.barcode,
                "Название": item.title,
                "Общее количество": 0,
                "Шевронов в товаре": item.pack_qty,
                "Физических шевронов": 0,
                "Продажи 90д": item.global_sales_90,
                "Остаток": item.stock,
                "В пути/приемка/запланировано учтено": item.inbound,
                "Комментарий": note,
            },
        )
        row["Общее количество"] += item.quantity
        row["Физических шевронов"] += item.physical
    fields = list(next(iter(totals.values())).keys()) if totals else ["Маркетплейс", "Артикул/SKU", "Общее количество"]
    ws.append(fields)
    for row in sorted(totals.values(), key=lambda item: (-_safe_int(item["Физических шевронов"]), item["Артикул/SKU"])):
        ws.append([row.get(field, "") for field in fields])
    _sheet_setup(ws)

    ws2 = workbook.create_sheet("Кластеры")
    cluster_fields = [
        "Кластер/регион",
        "Артикул/SKU",
        "SKU/nmID",
        "Баркод ВБ",
        "Название",
        "Количество",
        "Шевронов в товаре",
        "Физических шевронов",
        "Продажи в кластере 90д",
        "Продажи товара 90д",
        "Остаток",
        "В пути/приемка/запланировано учтено",
    ]
    ws2.append(cluster_fields)
    for item in sorted(items, key=lambda row: (row.cluster, -row.physical, row.article)):
        ws2.append(
            [
                item.cluster,
                item.article,
                item.sku,
                item.barcode,
                item.title,
                item.quantity,
                item.pack_qty,
                item.physical,
                item.sales_90,
                item.global_sales_90,
                item.stock,
                item.inbound,
            ]
        )
    _sheet_setup(ws2)
    workbook.save(path)


def _sheet_setup(ws: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 10
        for cell in ws[letter]:
            max_len = max(max_len, min(60, len("" if cell.value is None else str(cell.value))))
        ws.column_dimensions[letter].width = max_len + 2


def _write_outputs(
    data_dir: Path,
    run_dir: Path,
    report_path: Path,
    summary_path: Path,
    result: dict[str, Any],
    processed_dir: Path,
) -> None:
    write_json(summary_path, result)
    _write_report(report_path, result)
    _write_csv_rows(processed_dir / "source_status.csv", [{"source": key, "status": value} for key, value in result.get("source_status", {}).items()])
    write_summary_run_manifest(
        data_dir=data_dir,
        run_dir=run_dir,
        summary=result,
        task="plan-supply-workbooks",
        mode="read_only",
        risk="low",
        marketplaces=["ozon", "wb"],
        inputs=result["inputs"],
    )


def _write_report(path: Path, result: dict[str, Any]) -> None:
    summary = result.get("summary", {})
    lines = [
        "# Файлы в работу: Ozon/WB",
        "",
        f"Run ID: `{result['run_id']}`",
        f"Статус: `{result['overall_status']}`",
        "Режим: read-only, поставки в ЛК не создавались.",
        f"Период продаж: `{summary.get('period_from', result['inputs'].get('period_from'))}` - `{summary.get('period_to', result['inputs'].get('period_to'))}`.",
        "",
        "## Итог",
        "",
        f"- Ozon неделя 1: `{summary.get('ozon_week1_physical', 0)}` физических шевронов, `{summary.get('ozon_week1_units', 0)}` товарных единиц, `{summary.get('ozon_week1_articles', 0)}` артикулов.",
        f"- Ozon неделя 2: `{summary.get('ozon_week2_physical', 0)}` физических шевронов, `{summary.get('ozon_week2_units', 0)}` товарных единиц, `{summary.get('ozon_week2_articles', 0)}` артикулов.",
        f"- WB: `{summary.get('wb_physical', 0)}` физических шевронов, `{summary.get('wb_units', 0)}` товарных единиц, `{summary.get('wb_articles', 0)}` артикулов.",
        "",
        "## Source Status",
        "",
    ]
    for key, value in sorted(result.get("source_status", {}).items()):
        lines.append(f"- `{key}`: `{value}`")
    for label in ("ozon", "wb"):
        data = result.get(label) or {}
        if not data:
            continue
        lines.extend(["", f"## {label.upper()} топ кластеров", ""])
        for cluster, qty in (data.get("cluster_sales_top") or {}).items():
            lines.append(f"- {cluster}: продажи 90д `{qty}`")
        inbound = data.get("inbound_summary") or {}
        if inbound:
            lines.extend(["", f"### {label.upper()} inbound", ""])
            for key, value in inbound.items():
                lines.append(f"- `{key}`: `{value}`")
    if result.get("errors"):
        lines.extend(["", "## Ошибки", ""])
        for key, value in sorted(result["errors"].items()):
            lines.append(f"- `{key}`: {value}")
    if result.get("warnings"):
        lines.extend(["", "## Ограничения", ""])
        for warning in result["warnings"]:
            lines.append(f"- {warning}")
    lines.extend(["", "## Файлы", ""])
    for key, value in sorted(result.get("artifacts", {}).items()):
        if key.endswith("_xlsx") or key in {"report", "summary"}:
            lines.append(f"- `{key}`: `{value}`")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _plan_summary(source: dict[str, Any], *plans: list[PlanItem]) -> dict[str, Any]:
    merged = [item for plan in plans for item in plan]
    return {
        "status": source.get("status", "skipped"),
        "top_clusters": source.get("top_clusters", []),
        "cluster_sales_top": {cluster: (source.get("cluster_sales") or {}).get(cluster, 0) for cluster in source.get("top_clusters", [])},
        "physical": sum(item.physical for item in merged),
        "units": sum(item.quantity for item in merged),
        "articles": len({item.article for item in merged}),
        "inbound_summary": source.get("inbound_summary", {}),
        "sources": source.get("sources", {}),
    }


def _source_ok(ozon: dict[str, Any], wb: dict[str, Any], key: str) -> str:
    statuses = []
    for data in (ozon, wb):
        source_status = data.get("source_status") if isinstance(data, dict) else {}
        if source_status and key in source_status:
            statuses.append(source_status[key])
    if not statuses:
        return "blocked"
    return "ok" if all(status == "ok" for status in statuses) else "warning"


def _blocked_result(
    *,
    run_id: str,
    started_at: datetime,
    inputs: dict[str, Any],
    artifacts: dict[str, str],
    reason: str,
) -> dict[str, Any]:
    return {
        "run_id": run_id,
        "started_at": started_at.isoformat(timespec="seconds"),
        "mode": "read_only",
        "overall_status": "blocked",
        "blocked_reason": reason,
        "inputs": inputs,
        "source_status": {
            "current_stocks": "blocked",
            "sales_90_days": "blocked",
            "sales_localization": "blocked",
            "active_inbound_supplies": "blocked",
            "destination_cluster_mapping": "blocked",
            "production_constraints": "ok",
        },
        "summary": {
            "period_from": inputs["period_from"],
            "period_to": inputs["period_to"],
            "generated_files": 0,
        },
        "errors": {},
        "warnings": [],
        "artifacts": artifacts,
        "next_step": "Передать credentials или тестовые adapter-и; write-операции не выполняются.",
    }


def _load_unified_index(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        return {}
    rows = _read_csv_rows(path)
    index: dict[str, dict[str, str]] = {}
    for row in rows:
        for key in ("ozon_offer_id", "ozon_sku", "wb_vendor_code", "wb_nm_id", "internal_sku"):
            value = str(row.get(key) or "").strip()
            if value:
                index[value] = row
    return index


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _write_csv_rows(path: Path, rows: list[dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    if not fieldnames:
        fieldnames = ["empty"]
        rows = []
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _rows_from_plan(items: list[PlanItem]) -> list[dict[str, Any]]:
    return [
        {
            "marketplace": item.marketplace,
            "cluster": item.cluster,
            "article": item.article,
            "sku_or_nmid": item.sku,
            "barcode": item.barcode,
            "title": item.title,
            "quantity": item.quantity,
            "pack_qty": item.pack_qty,
            "physical_chevrons": item.physical,
            "cluster_sales_90": item.sales_90,
            "global_sales_90": item.global_sales_90,
            "stock": item.stock,
            "inbound_accounted": item.inbound,
            "score": round(item.score, 4),
        }
        for item in items
    ]


def _ozon_titles_and_skus(product_list: list[dict[str, Any]], info_rows: list[dict[str, Any]]) -> tuple[dict[str, str], dict[str, str]]:
    title_by_offer: dict[str, str] = {}
    sku_by_offer: dict[str, str] = {}
    for row in info_rows:
        offer = str(row.get("offer_id") or "").strip()
        if offer:
            title_by_offer[offer] = str(row.get("name") or row.get("title") or "").strip()
            sku_by_offer[offer] = str(row.get("sku") or "").strip()
    for row in product_list:
        offer = str(row.get("offer_id") or "").strip()
        sku = str(row.get("sku") or "").strip()
        if offer and sku:
            sku_by_offer.setdefault(offer, sku)
    return title_by_offer, sku_by_offer


def _ozon_fbo_stock_by_offer(stock_rows: list[dict[str, Any]]) -> Counter[str]:
    result: Counter[str] = Counter()
    for row in stock_rows:
        offer = str(row.get("offer_id") or "").strip()
        for stock in row.get("stocks") or []:
            if str(stock.get("type") or stock.get("source") or "").lower() == "fbo":
                result[offer] += _safe_int(stock.get("present"))
    return result


def _ozon_promised_by_offer(rows: list[dict[str, Any]]) -> Counter[str]:
    result: Counter[str] = Counter()
    for row in rows:
        offer = str(row.get("item_code") or "").strip()
        if offer:
            result[offer] += _safe_int(row.get("promised_amount"))
    return result


def _infer_pack_qty(title: str, fallback: Any = 1) -> int:
    title_l = (title or "").lower()
    fallback_int = max(1, _safe_int(fallback) or 1)
    if "петлиц" in title_l:
        return 2
    if "комплект" in title_l:
        match = re.search(r"(\d+)\s*(?:шт|штук|предмет)", title_l)
        if match:
            return max(1, int(match.group(1)))
        return max(2, fallback_int)
    return fallback_int


def _is_manufactured(title: str, article: str, subject: str = "", product_group: str = "") -> bool:
    text = " ".join([title or "", article or "", subject or "", product_group or ""])
    if MANUFACTURED_EXCLUDE_RE.search(text):
        return False
    if product_group in {"chev", "nash", "loop"}:
        return True
    return bool(MANUFACTURED_INCLUDE_RE.search(text))


def _wb_region_from_sale(row: dict[str, Any]) -> str:
    oblast = str(row.get("oblastOkrugName") or "").strip()
    if oblast in WB_REGION_MAP:
        return WB_REGION_MAP[oblast]
    return _normalize_wb_region(str(row.get("oblastOkrugName") or row.get("regionName") or row.get("warehouseName") or "Не указан"))


def _normalize_wb_region(value: str) -> str:
    return (value or "").strip().replace("Юг-СК", "Юг/СК").replace("Сибирь - ДВ", "Сибирь/ДВ")


def _safe_int(value: Any) -> int:
    try:
        if value is None or value == "":
            return 0
        return int(float(str(value).replace(",", ".")))
    except Exception:
        return 0


def _period_start(day: date) -> str:
    return f"{day.isoformat()}T00:00:00.000Z"


def _period_end(day: date) -> str:
    return f"{day.isoformat()}T23:59:59.999Z"


def _safe_error(exc: BaseException) -> str:
    return str(exc).replace("\n", " ")[:800]
