from __future__ import annotations

from collections import Counter, defaultdict
import csv
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font

from seller_agent.config import AppCredentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.ozon.adapter import OzonSellerAdapter
from seller_agent.marketplaces.wb.finance_adapter import WbFinanceAdapter
from seller_agent.marketplaces.wb.promotion_adapter import WbPromotionAdapter
from seller_agent.marketplaces.wb.statistics_adapter import WbStatisticsAdapter
from seller_agent.reports.writer import ensure_dir, write_json


MOSCOW = ZoneInfo("Europe/Moscow")
REPORT_TYPES = {"short", "financial", "full"}
MARKETPLACES = {"ozon", "wb"}


def run_marketplace_period_report(
    *,
    credentials: AppCredentials,
    data_dir: Path,
    marketplace: str,
    report_type: str,
    date_from: str,
    date_to: str,
    run_id: str | None = None,
) -> dict[str, Any]:
    marketplace = str(marketplace).strip().lower()
    report_type = str(report_type).strip().lower()
    start = _parse_date(date_from)
    end = _parse_date(date_to)
    if marketplace not in MARKETPLACES:
        raise ValueError("marketplace must be ozon or wb")
    if report_type not in REPORT_TYPES:
        raise ValueError("report_type must be short, financial or full")
    if start > end:
        raise ValueError("date_from must be earlier than or equal to date_to")
    if end > datetime.now(MOSCOW).date():
        raise ValueError("date_to cannot be in the future")

    started_at = datetime.now(MOSCOW)
    run_id = run_id or f"marketplace_period_report_{marketplace}_{started_at.strftime('%Y%m%dT%H%M%S')}"
    run_dir = ensure_dir(data_dir / "runs" / started_at.date().isoformat() / run_id)
    catalog_rows = _load_catalog(data_dir / "catalog" / "unified" / "products.csv")
    catalog = _CatalogIndex(catalog_rows)

    if marketplace == "ozon":
        if credentials.ozon_seller is None:
            raise ValueError("Ozon Seller API credentials are not configured")
        metrics = _collect_ozon(
            adapter=OzonSellerAdapter(credentials.ozon_seller),
            catalog=catalog,
            start=start,
            end=end,
        )
    else:
        if credentials.wb is None:
            raise ValueError("WB API credentials are not configured")
        metrics = _collect_wb(
            statistics=WbStatisticsAdapter(credentials.wb),
            finance=WbFinanceAdapter(credentials.wb),
            promotion=WbPromotionAdapter(credentials.wb),
            catalog=catalog,
            start=start,
            end=end,
        )

    report_path = run_dir / "marketplace_period_report.md"
    xlsx_path = run_dir / "marketplace_period_report.xlsx"
    data_path = run_dir / "marketplace_period_report.json"
    _write_markdown(
        report_path,
        marketplace=marketplace,
        report_type=report_type,
        start=start,
        end=end,
        metrics=metrics,
    )
    _write_xlsx(
        xlsx_path,
        marketplace=marketplace,
        report_type=report_type,
        start=start,
        end=end,
        metrics=metrics,
    )
    write_json(data_path, metrics)

    warnings = list(metrics.get("warnings") or [])
    summary = {
        "run_id": run_id,
        "started_at": started_at.isoformat(timespec="seconds"),
        "overall_status": "warning" if warnings else "ok",
        "marketplace": marketplace,
        "report_type": report_type,
        "date_from": start.isoformat(),
        "date_to": end.isoformat(),
        "metrics": metrics.get("summary") or {},
        "warnings": warnings,
        "artifacts": {
            "report": str(report_path),
            "xlsx": str(xlsx_path),
            "json": str(data_path),
        },
    }
    write_json(run_dir / "summary.json", summary)
    manifest_artifacts = write_summary_run_manifest(
        data_dir=data_dir,
        run_dir=run_dir,
        summary=summary,
        task="marketplace-period-report",
        mode="read_only",
        risk="low",
        marketplaces=[marketplace],
        inputs={
            "marketplace": marketplace,
            "report_type": report_type,
            "date_from": start.isoformat(),
            "date_to": end.isoformat(),
        },
        lifecycle_status="closed",
        closed=True,
    )
    summary["artifacts"].update(manifest_artifacts)
    write_json(run_dir / "summary.json", summary)
    return summary


def _collect_ozon(
    *,
    adapter: OzonSellerAdapter,
    catalog: "_CatalogIndex",
    start: date,
    end: date,
) -> dict[str, Any]:
    date_from = f"{start.isoformat()}T00:00:00.000Z"
    date_to = f"{end.isoformat()}T23:59:59.999Z"
    operations = adapter.fetch_finance_transactions(date_from=date_from, date_to=date_to)
    analytics = adapter.fetch_analytics_data(
        date_from=start.isoformat(),
        date_to=end.isoformat(),
        metrics=["ordered_units", "revenue", "cancellations", "returns"],
        dimensions=["day"],
        limit=1000,
    )

    delivered = [row for row in operations if row.get("operation_type") == "OperationAgentDeliveredToCustomer"]
    returned_operations = [row for row in operations if row.get("operation_type") == "OperationItemReturn"]
    warnings = []
    return_source = "Ozon /v1/returns/list"
    try:
        return_rows = adapter.fetch_returns(
            logistic_return_date_from=date_from,
            logistic_return_date_to=date_to,
            return_schema="FBO",
        )
        return_metrics = _ozon_return_metrics(return_rows)
    except Exception as exc:  # noqa: BLE001 - finance fallback keeps the report usable.
        return_metrics = _ozon_finance_return_metrics(returned_operations)
        return_source = "Ozon /v3/finance/transaction/list (fallback по уникальным отправлениям)"
        warnings.append(
            "Реестр возвратов Ozon недоступен; количество рассчитано по уникальным "
            f"отправлениям финансовых операций: {_safe_error(exc)}"
        )
    product_totals: dict[str, dict[str, Any]] = defaultdict(lambda: {"units": 0, "physical": 0, "gross": 0.0})
    unmapped: Counter[str] = Counter()
    buyout_units = 0
    physical = 0
    buyout_gross = 0.0
    for row in delivered:
        items = row.get("items") if isinstance(row.get("items"), list) else []
        row_gross = _number(row.get("accruals_for_sale"))
        denominator = sum(max(1, _int(item.get("quantity"), 1)) for item in items if isinstance(item, dict)) or 1
        if not items:
            buyout_units += 1
            physical += 1
            unmapped["operation_without_items"] += 1
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            units = max(1, _int(item.get("quantity"), 1))
            sku = _text(item.get("sku"))
            pack_qty = catalog.ozon_pack_qty(sku)
            if sku and not catalog.has_ozon(sku):
                unmapped[sku] += units
            buyout_units += units
            physical += units * pack_qty
            buyout_gross += row_gross * units / denominator
            product = product_totals[sku or "unknown"]
            product["units"] += units
            product["physical"] += units * pack_qty
            product["gross"] += row_gross * units / denominator

    gross = sum(_number(row.get("accruals_for_sale")) for row in operations)
    net = sum(_number(row.get("amount")) for row in operations)
    expenses, operation_counts = _ozon_expenses(operations, gross=gross, net=net)
    daily = _ozon_analytics_daily(analytics)
    orders = sum(_int(row.get("ordered_units")) for row in daily)
    revenue = sum(_number(row.get("revenue")) for row in daily)
    cancellations = sum(_int(row.get("cancellations")) for row in daily)

    top_products = []
    for sku, values in sorted(product_totals.items(), key=lambda item: item[1]["gross"], reverse=True)[:20]:
        top_products.append(
            {
                "id": sku,
                "internal_sku": catalog.ozon_internal_sku(sku),
                "name": catalog.ozon_name(sku),
                **values,
            }
        )
    if end == datetime.now(MOSCOW).date():
        warnings.append("Текущий день может быть финансово неполным до закрытия операций Ozon.")
    if unmapped:
        warnings.append(f"Не найден pack_qty для {sum(unmapped.values())} выкупленных единиц Ozon; использовано значение 1.")
    return {
        "summary": {
            "orders": orders,
            "order_amount": _money(revenue),
            "buyout_units": buyout_units,
            "physical_pieces": physical,
            "buyout_gross": _money(buyout_gross),
            "returns": return_metrics["total"],
            "return_cancellations": return_metrics["cancellations"],
            "client_returns": return_metrics["client_returns"],
            "return_unknown": return_metrics["unknown"],
            "cancellations": cancellations,
            "gross": _money(gross),
            "expenses": _money(max(0.0, gross - net)),
            "net": _money(net),
            "net_per_piece": _money(net / physical if physical else 0),
        },
        "expenses": expenses,
        "operation_counts": dict(sorted(operation_counts.items())),
        "daily": daily,
        "top_products": top_products,
        "unmapped_ids": dict(unmapped.most_common(30)),
        "warnings": warnings,
        "sources": ["Ozon /v3/finance/transaction/list", return_source, "Ozon /v1/analytics/data"],
    }


def collect_ozon_period_metrics(
    *,
    adapter: OzonSellerAdapter,
    catalog_rows: list[dict[str, str]],
    start: date,
    end: date,
) -> dict[str, Any]:
    """Collect Ozon period metrics for other read-only financial workflows."""
    return _collect_ozon(
        adapter=adapter,
        catalog=_CatalogIndex(catalog_rows),
        start=start,
        end=end,
    )


def _collect_wb(
    *,
    statistics: WbStatisticsAdapter,
    finance: WbFinanceAdapter,
    promotion: WbPromotionAdapter,
    catalog: "_CatalogIndex",
    start: date,
    end: date,
) -> dict[str, Any]:
    warnings: list[str] = []
    if end == datetime.now(MOSCOW).date():
        warnings.append("Текущий день может быть финансово неполным до закрытия отчётов WB.")
    report_rows = finance.fetch_sales_reports(
        date_from=start.isoformat(), date_to=end.isoformat(), period="daily"
    )
    acquiring_reports: list[dict[str, Any]] = []
    try:
        acquiring_reports = finance.fetch_acquiring_reports(
            date_from=start.isoformat(),
            date_to=end.isoformat(),
        )
    except Exception as exc:  # noqa: BLE001 - sales report remains usable.
        warnings.append(f"Издержки WB на приём платежей не подтверждены: {_safe_error(exc)}")
    detail_rows = finance.fetch_sales_report_details(
        date_from=start.isoformat(), date_to=end.isoformat(), period="daily"
    )
    order_rows: list[dict[str, Any]] = []
    try:
        fetched_orders = statistics.fetch_orders(date_from=f"{start.isoformat()}T00:00:00", flag=0)
        order_rows = _filter_wb_rows_by_date(fetched_orders, start=start, end=end, keys=("date", "lastChangeDate"))
    except Exception as exc:  # noqa: BLE001 - finance report remains usable.
        warnings.append(f"Оперативные заказы WB недоступны: {_safe_error(exc)}")

    ad_spend = 0.0
    try:
        ad_spend = _wb_ad_spend(promotion, start=start, end=end)
    except Exception as exc:  # noqa: BLE001 - report must expose the missing expense.
        warnings.append(f"Расходы на рекламу WB не подтверждены: {_safe_error(exc)}")

    sale_rows = [row for row in detail_rows if _wb_doc_type(row) == "sale" and _int(row.get("quantity"), 1) > 0]
    return_rows = [row for row in detail_rows if _wb_doc_type(row) == "return"]
    physical = 0
    buyout_units = 0
    unmapped: Counter[str] = Counter()
    product_totals: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"units": 0, "physical": 0, "gross": 0.0, "for_pay": 0.0}
    )
    for row in sale_rows:
        units = max(1, _int(row.get("quantity"), 1))
        nm_id = _text(row.get("nmId"))
        vendor_code = _text(row.get("vendorCode"))
        pack_qty = catalog.wb_pack_qty(nm_id=nm_id, vendor_code=vendor_code)
        if not catalog.has_wb(nm_id=nm_id, vendor_code=vendor_code):
            unmapped[nm_id or vendor_code or "unknown"] += units
        buyout_units += units
        physical += units * pack_qty
        key = nm_id or vendor_code or "unknown"
        product = product_totals[key]
        product["units"] += units
        product["physical"] += units * pack_qty
        product["gross"] += _number(row.get("retailAmount"))
        product["for_pay"] += _number(row.get("forPay"))

    gross = sum(_number(row.get("retailAmountSum")) for row in report_rows)
    bank_payment = sum(_number(row.get("bankPaymentSum")) for row in report_rows)
    acquiring_expense = sum(
        _number(row.get("acquiringFeeSum")) + _number(row.get("acquiringFeeVatSum"))
        for row in acquiring_reports
    )
    expenses, adjustments = _wb_expenses(
        report_rows,
        gross=gross,
        bank_payment=bank_payment,
        ad_spend=ad_spend,
        acquiring_expense=acquiring_expense,
    )
    total_expenses = sum(expenses.values())
    net = gross - total_expenses
    cash_after_adjustments = bank_payment - acquiring_expense - ad_spend
    active_orders = [row for row in order_rows if not _truthy(row.get("isCancel"))]
    cancelled_orders = [row for row in order_rows if _truthy(row.get("isCancel"))]
    order_amount = sum(_number(row.get("priceWithDisc")) for row in active_orders)
    daily = _wb_daily(report_rows=report_rows, detail_rows=detail_rows)
    top_products = []
    for product_id, values in sorted(product_totals.items(), key=lambda item: item[1]["gross"], reverse=True)[:20]:
        top_products.append(
            {
                "id": product_id,
                "internal_sku": catalog.wb_internal_sku(product_id),
                "name": catalog.wb_name(product_id),
                **values,
            }
        )
    if unmapped:
        warnings.append(f"Не найден pack_qty для {sum(unmapped.values())} выкупленных единиц WB; использовано значение 1.")
    if not report_rows:
        warnings.append("WB не вернул сводные финансовые отчёты за выбранный период.")
    return {
        "summary": {
            "orders": len(active_orders),
            "order_amount": _money(order_amount),
            "buyout_units": buyout_units,
            "physical_pieces": physical,
            "buyout_gross": _money(sum(_number(row.get("retailAmount")) for row in sale_rows)),
            "returns": sum(max(1, abs(_int(row.get("quantity"), 1))) for row in return_rows),
            "cancellations": len(cancelled_orders),
            "gross": _money(gross),
            "expenses": _money(total_expenses),
            "net": _money(net),
            "net_per_piece": _money(net / physical if physical else 0),
            "advertising": _money(ad_spend),
            "bank_payment": _money(bank_payment),
            "adjustments": _money(sum(adjustments.values())),
            "cash_after_adjustments": _money(cash_after_adjustments),
        },
        "expenses": expenses,
        "adjustments": adjustments,
        "daily": daily,
        "top_products": top_products,
        "unmapped_ids": dict(unmapped.most_common(30)),
        "warnings": warnings,
        "sources": [
            "WB /api/finance/v1/sales-reports/list",
            "WB /api/finance/v1/acquiring/list",
            "WB /api/finance/v1/sales-reports/detailed",
            "WB /api/v1/supplier/orders",
            "WB /adv/v3/fullstats",
        ],
    }


def collect_wb_period_metrics(
    *,
    statistics: WbStatisticsAdapter,
    finance: WbFinanceAdapter,
    promotion: WbPromotionAdapter,
    catalog_rows: list[dict[str, str]],
    start: date,
    end: date,
) -> dict[str, Any]:
    """Collect WB period metrics for other read-only financial workflows."""
    return _collect_wb(
        statistics=statistics,
        finance=finance,
        promotion=promotion,
        catalog=_CatalogIndex(catalog_rows),
        start=start,
        end=end,
    )


def _ozon_expenses(
    operations: list[dict[str, Any]], *, gross: float, net: float
) -> tuple[dict[str, float], Counter[str]]:
    result = defaultdict(float)
    counts: Counter[str] = Counter()
    for row in operations:
        operation_type = _text(row.get("operation_type")) or "unknown"
        counts[operation_type] += 1
        amount = _number(row.get("amount"))
        services = row.get("services") if isinstance(row.get("services"), list) else []
        service_expense = max(0.0, -sum(_number(item.get("price")) for item in services if isinstance(item, dict)))
        if operation_type == "OperationAgentDeliveredToCustomer":
            result["Логистика"] += service_expense
            result["Комиссия"] += max(0.0, -_number(row.get("sale_commission")))
        elif operation_type == "OperationItemReturn":
            result["Возвраты"] += max(0.0, -amount)
        elif operation_type == "MarketplaceRedistributionOfAcquiringOperation":
            result["Эквайринг"] += max(0.0, -amount)
        elif operation_type == "OperationMarketplaceCostPerClick":
            result["Реклама"] += max(0.0, -amount)
        elif operation_type == "OperationMarketplaceServiceStorage":
            result["Хранение"] += max(0.0, -amount)
        elif operation_type == "MarketplaceServiceItemCrossdocking":
            result["Кросс-докинг"] += max(0.0, -amount)
        elif amount < 0:
            result["Прочее"] += -amount
    expected = max(0.0, gross - net)
    delta = expected - sum(result.values())
    if abs(delta) >= 0.01:
        result["Сверочная разница"] += delta
    return {key: _money(value) for key, value in sorted(result.items())}, counts


def _ozon_return_metrics(rows: list[dict[str, Any]]) -> dict[str, int]:
    result = Counter[str]()
    for row in rows:
        product = row.get("product") if isinstance(row.get("product"), dict) else {}
        quantity = max(1, _int(product.get("quantity"), 1))
        return_type = _text(row.get("type"))
        if return_type == "Cancellation":
            result["cancellations"] += quantity
        elif return_type == "ClientReturn":
            result["client_returns"] += quantity
        else:
            result["unknown"] += quantity
    result["total"] = result["cancellations"] + result["client_returns"] + result["unknown"]
    return {key: result[key] for key in ("total", "cancellations", "client_returns", "unknown")}


def _ozon_finance_return_metrics(rows: list[dict[str, Any]]) -> dict[str, int]:
    grouped: dict[tuple[str, str], int] = {}
    for row in rows:
        posting = row.get("posting") if isinstance(row.get("posting"), dict) else {}
        posting_number = _text(posting.get("posting_number")) or _text(row.get("posting_number"))
        items = row.get("items") if isinstance(row.get("items"), list) else []
        if not items:
            fallback_id = posting_number or _text(row.get("operation_id"))
            grouped[(fallback_id, "")] = max(grouped.get((fallback_id, ""), 0), 1)
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            sku = _text(item.get("sku"))
            key = (posting_number or _text(row.get("operation_id")), sku)
            grouped[key] = max(grouped.get(key, 0), max(1, _int(item.get("quantity"), 1)))
    total = sum(grouped.values())
    return {"total": total, "cancellations": 0, "client_returns": 0, "unknown": total}


def _wb_expenses(
    rows: list[dict[str, Any]],
    *,
    gross: float,
    bank_payment: float,
    ad_spend: float,
    acquiring_expense: float = 0.0,
) -> tuple[dict[str, float], dict[str, float]]:
    signed_values = {
        "Удержания площадки до логистики": max(0.0, gross - sum(_number(row.get("forPaySum")) for row in rows)),
        "Логистика": sum(_number(row.get("deliveryServiceSum")) for row in rows),
        "Хранение": sum(_number(row.get("paidStorageSum")) for row in rows),
        "Приёмка": sum(_number(row.get("paidAcceptanceSum")) for row in rows),
        "Удержания": sum(_number(row.get("deductionSum")) for row in rows),
        "Штрафы": sum(_number(row.get("penaltySum")) for row in rows),
        "Кешбэк и корректировки": sum(
            _number(row.get("cashbackAmountSum"))
            + _number(row.get("cashbackDiscountSum"))
            + _number(row.get("cashbackCommissionChangeSum"))
            for row in rows
        ),
        "Удержание по графику платежей": sum(_number(row.get("paymentSchedule")) for row in rows),
    }
    additional_payments = sum(_number(row.get("additionalPaymentSum")) for row in rows)
    expenses = {key: max(0.0, value) for key, value in signed_values.items()}
    if additional_payments < 0:
        expenses["Отрицательные дополнительные выплаты"] = abs(additional_payments)
    expenses["Издержки на приём платежей"] = max(0.0, acquiring_expense)
    expenses["Реклама"] = max(0.0, ad_spend)

    adjustments = {
        f"{key}: зачисление": abs(value)
        for key, value in signed_values.items()
        if value < 0
    }
    if additional_payments > 0:
        adjustments["Дополнительные выплаты"] = additional_payments

    sales_report_expenses = sum(max(0.0, value) for value in signed_values.values())
    if additional_payments < 0:
        sales_report_expenses += abs(additional_payments)
    expected_bank_payment = gross - sales_report_expenses + sum(adjustments.values())
    reconciliation_delta = bank_payment - expected_bank_payment
    if abs(reconciliation_delta) >= 0.01:
        adjustments["Неклассифицированная сверочная корректировка"] = reconciliation_delta
    return (
        {key: _money(value) for key, value in expenses.items()},
        {key: _money(value) for key, value in adjustments.items()},
    )


def _ozon_analytics_daily(payload: dict[str, Any]) -> list[dict[str, Any]]:
    result = payload.get("result") if isinstance(payload, dict) else {}
    rows = result.get("data") if isinstance(result, dict) else []
    parsed: list[dict[str, Any]] = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        dimensions = row.get("dimensions") if isinstance(row.get("dimensions"), list) else []
        day = _text(dimensions[0].get("id"))[:10] if dimensions and isinstance(dimensions[0], dict) else ""
        metrics = row.get("metrics") if isinstance(row.get("metrics"), list) else []
        parsed.append(
            {
                "date": day,
                "ordered_units": _int(metrics[0] if len(metrics) > 0 else 0),
                "revenue": _money(_number(metrics[1] if len(metrics) > 1 else 0)),
                "cancellations": _int(metrics[2] if len(metrics) > 2 else 0),
                "returns": _int(metrics[3] if len(metrics) > 3 else 0),
            }
        )
    return sorted(parsed, key=lambda row: row["date"])


def _wb_daily(*, report_rows: list[dict[str, Any]], detail_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    daily: dict[str, dict[str, Any]] = defaultdict(lambda: {"gross": 0.0, "bank_payment": 0.0, "sales": 0, "returns": 0})
    for row in report_rows:
        day = _text(row.get("dateFrom"))[:10]
        daily[day]["gross"] += _number(row.get("retailAmountSum"))
        daily[day]["bank_payment"] += _number(row.get("bankPaymentSum"))
    for row in detail_rows:
        day = _text(row.get("saleDt"))[:10] or _text(row.get("rrDate"))[:10]
        if _wb_doc_type(row) == "sale":
            daily[day]["sales"] += max(1, _int(row.get("quantity"), 1))
        elif _wb_doc_type(row) == "return":
            daily[day]["returns"] += max(1, abs(_int(row.get("quantity"), 1)))
    return [
        {"date": day, **{key: _money(value) if key in {"gross", "bank_payment"} else value for key, value in values.items()}}
        for day, values in sorted(daily.items())
        if day
    ]


def _wb_ad_spend(adapter: WbPromotionAdapter, *, start: date, end: date) -> float:
    count_data = adapter.fetch_campaign_count()
    ids: set[int] = set()
    for group in count_data.get("adverts") or []:
        if not isinstance(group, dict):
            continue
        for item in group.get("advert_list") or []:
            if isinstance(item, dict) and _int(item.get("advertId")):
                ids.add(_int(item.get("advertId")))
    if not ids:
        return 0.0
    spend = 0.0
    chunk_start = start
    while chunk_start <= end:
        chunk_end = min(end, chunk_start + timedelta(days=30))
        rows = adapter.fetch_fullstats(
            ids=sorted(ids), date_from=chunk_start.isoformat(), date_to=chunk_end.isoformat()
        )
        for campaign in rows:
            days = campaign.get("days") if isinstance(campaign.get("days"), list) else []
            spend += sum(_number(day.get("sum")) for day in days if isinstance(day, dict)) if days else _number(campaign.get("sum"))
        chunk_start = chunk_end + timedelta(days=1)
    return _money(spend)


class _CatalogIndex:
    def __init__(self, rows: list[dict[str, str]]) -> None:
        self.ozon: dict[str, dict[str, str]] = {}
        self.wb: dict[str, dict[str, str]] = {}
        for row in rows:
            for key in ("ozon_sku", "ozon_offer_id", "ozon_product_id"):
                value = _text(row.get(key))
                if value:
                    self.ozon[value] = row
            for key in ("wb_nm_id", "wb_vendor_code"):
                value = _text(row.get(key))
                if value:
                    self.wb[value] = row

    def has_ozon(self, value: str) -> bool:
        return value in self.ozon

    def ozon_pack_qty(self, value: str) -> int:
        return max(1, _int((self.ozon.get(value) or {}).get("pack_qty"), 1))

    def ozon_internal_sku(self, value: str) -> str:
        return _text((self.ozon.get(value) or {}).get("internal_sku"))

    def ozon_name(self, value: str) -> str:
        return _text((self.ozon.get(value) or {}).get("product_name"))

    def has_wb(self, *, nm_id: str, vendor_code: str) -> bool:
        return nm_id in self.wb or vendor_code in self.wb

    def _wb_row(self, value: str) -> dict[str, str]:
        return self.wb.get(value) or {}

    def wb_pack_qty(self, *, nm_id: str, vendor_code: str) -> int:
        row = self._wb_row(nm_id) or self._wb_row(vendor_code)
        return max(1, _int(row.get("pack_qty"), 1))

    def wb_internal_sku(self, value: str) -> str:
        return _text(self._wb_row(value).get("internal_sku"))

    def wb_name(self, value: str) -> str:
        return _text(self._wb_row(value).get("product_name"))


def _write_markdown(
    path: Path,
    *,
    marketplace: str,
    report_type: str,
    start: date,
    end: date,
    metrics: dict[str, Any],
) -> None:
    title = "Ozon" if marketplace == "ozon" else "Wildberries"
    report_labels = {"short": "Краткий", "financial": "Финансовый", "full": "Полный"}
    summary = metrics["summary"]
    lines = [
        f"# {title}: отчёт за период",
        "",
        f"- Период: `{start.strftime('%d.%m.%Y')} - {end.strftime('%d.%m.%Y')}`",
        f"- Вид отчёта: `{report_labels[report_type]}`",
        "- Режим: `read-only`",
        "",
        "## Основные показатели",
        "",
        f"- Заказы: **{summary['orders']} товаров** на **{_rub(summary['order_amount'])}**.",
        f"- Выкупы: **{summary['buyout_units']} товаров / {summary['physical_pieces']} физических изделий**.",
        (
            f"- Возвратные события: **{summary['returns']}** "
            f"(отмены/невыкупы: **{summary.get('return_cancellations', 0)}**; "
            f"возвраты после покупки: **{summary.get('client_returns', 0)}**; "
            f"другие: **{summary.get('return_unknown', 0)}**)."
            if marketplace == "ozon"
            else f"- Возвраты: **{summary['returns']}**; отмены: **{summary['cancellations']}**."
        ),
        f"- Продажи до расходов: **{_rub(summary['gross'])}**.",
        f"- Расходы: **{_rub(summary['expenses'])}**.",
        f"- После текущих расходов: **{_rub(summary['net'])}**.",
        f"- На одно физическое изделие: **{_rub(summary['net_per_piece'])}**.",
    ]
    if marketplace == "wb":
        lines.extend(
            [
                f"- Корректировки и зачисления WB: **{_rub(summary.get('adjustments', 0))}**.",
                f"- Начисление WB с корректировками после рекламы: **{_rub(summary.get('cash_after_adjustments', 0))}**.",
            ]
        )
    if report_type in {"financial", "full"}:
        lines.extend(["", "## Расходы", ""])
        for name, value in metrics.get("expenses", {}).items():
            lines.append(f"- {name}: **{_rub(value)}**.")
        if marketplace == "wb" and metrics.get("adjustments"):
            lines.extend(["", "## Корректировки И Зачисления", ""])
            for name, value in metrics["adjustments"].items():
                lines.append(f"- {name}: **{_rub(value)}**.")
    if report_type == "full":
        lines.extend(["", "## Товары", "", "| Товар | Единиц | Физических изделий | Продажи | К перечислению* |", "| --- | ---: | ---: | ---: | ---: |"])
        for row in metrics.get("top_products", []):
            label = row.get("internal_sku") or row.get("name") or row.get("id")
            lines.append(
                f"| {label} | {row.get('units', 0)} | {row.get('physical', 0)} | {_rub(row.get('gross', 0))} | {_rub(row.get('for_pay', 0)) if marketplace == 'wb' else 'н/д'} |"
            )
        lines.extend(["", "*Общие расходы площадки и реклама не распределяются по товарам.*"])
    if metrics.get("warnings"):
        lines.extend(["", "## Ограничения", ""])
        lines.extend(f"- {warning}" for warning in metrics["warnings"])
    lines.extend(["", "## Источники", ""])
    lines.extend(f"- `{source}`" for source in metrics.get("sources", []))
    lines.extend(["", "Изменений в кабинете маркетплейса не выполнялось.", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_xlsx(
    path: Path,
    *,
    marketplace: str,
    report_type: str,
    start: date,
    end: date,
    metrics: dict[str, Any],
) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Сводка"
    summary_sheet.append(["Маркетплейс", "Ozon" if marketplace == "ozon" else "Wildberries"])
    summary_sheet.append(["Период", f"{start.isoformat()} - {end.isoformat()}"])
    summary_sheet.append(["Тип", report_type])
    summary_sheet.append([])
    summary_rows = [
        ("orders", "Заказы, товаров"),
        ("order_amount", "Сумма заказов, руб."),
        ("buyout_units", "Выкупы, товаров"),
        ("physical_pieces", "Выкупы, физических изделий"),
        ("returns", "Возвратные события" if marketplace == "ozon" else "Возвраты"),
    ]
    if marketplace == "ozon":
        summary_rows.extend(
            [
                ("return_cancellations", "Отмены/невыкупы в возвратных событиях"),
                ("client_returns", "Возвраты после покупки"),
                ("return_unknown", "Другие возвратные события"),
            ]
        )
    summary_rows.extend(
        [
        ("cancellations", "Отмены"),
        ("gross", "Продажи до расходов, руб."),
        ("expenses", "Расходы, руб."),
        ("net", "После текущих расходов, руб."),
        ("net_per_piece", "На физическое изделие, руб."),
        ]
    )
    if marketplace == "wb":
        summary_rows.extend(
            [
                ("adjustments", "Корректировки и зачисления WB, руб."),
                ("cash_after_adjustments", "Начисление с корректировками после рекламы, руб."),
            ]
        )
    for key, label in summary_rows:
        summary_sheet.append([label, metrics["summary"].get(key, 0)])
    summary_sheet["A1"].font = Font(bold=True)
    summary_sheet.column_dimensions["A"].width = 38
    summary_sheet.column_dimensions["B"].width = 24

    if report_type in {"financial", "full"}:
        sheet = workbook.create_sheet("Расходы")
        sheet.append(["Статья", "Сумма, руб."])
        for key, value in metrics.get("expenses", {}).items():
            sheet.append([key, value])
        if marketplace == "wb" and metrics.get("adjustments"):
            sheet.append([])
            sheet.append(["Корректировки и зачисления", "Сумма, руб."])
            for key, value in metrics["adjustments"].items():
                sheet.append([key, value])
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 18

    if report_type == "full":
        products = workbook.create_sheet("Товары")
        headers = ["ID", "Внутренний артикул", "Название", "Единиц", "Физических изделий", "Продажи", "К перечислению"]
        products.append(headers)
        for row in metrics.get("top_products", []):
            products.append([row.get("id"), row.get("internal_sku"), row.get("name"), row.get("units"), row.get("physical"), row.get("gross"), row.get("for_pay")])
        products.freeze_panes = "A2"
        daily = workbook.create_sheet("По дням")
        daily_rows = metrics.get("daily", [])
        daily_headers = list(daily_rows[0]) if daily_rows else ["date"]
        daily.append(daily_headers)
        for row in daily_rows:
            daily.append([row.get(key) for key in daily_headers])
        daily.freeze_panes = "A2"
    ensure_dir(path.parent)
    workbook.save(path)


def _load_catalog(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _filter_wb_rows_by_date(
    rows: list[dict[str, Any]], *, start: date, end: date, keys: tuple[str, ...]
) -> list[dict[str, Any]]:
    filtered = []
    for row in rows:
        raw = next((_text(row.get(key)) for key in keys if _text(row.get(key))), "")[:10]
        try:
            value = date.fromisoformat(raw)
        except ValueError:
            continue
        if start <= value <= end:
            filtered.append(row)
    return filtered


def _wb_doc_type(row: dict[str, Any]) -> str:
    text = f"{_text(row.get('docTypeName'))} {_text(row.get('sellerOperName'))}".lower()
    if "возврат" in text:
        return "return"
    if "продаж" in text:
        return "sale"
    return "other"


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise ValueError(f"invalid ISO date: {value}") from exc


def _text(value: Any) -> str:
    return str(value or "").strip()


def _number(value: Any) -> float:
    try:
        return float(str(value or 0).replace(" ", "").replace(",", "."))
    except ValueError:
        return 0.0


def _int(value: Any, default: int = 0) -> int:
    try:
        return int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError):
        return default


def _money(value: Any) -> float:
    return round(_number(value), 2)


def _rub(value: Any) -> str:
    return f"{_number(value):,.2f} руб.".replace(",", " ")


def _truthy(value: Any) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "да"}


def _safe_error(exc: BaseException) -> str:
    return str(exc).replace("\n", " ").replace("\r", " ")[:300]
