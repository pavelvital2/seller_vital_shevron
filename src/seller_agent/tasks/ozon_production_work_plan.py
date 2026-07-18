from __future__ import annotations

from collections import Counter, defaultdict
import csv
from dataclasses import dataclass
from datetime import date, datetime, timedelta
import hashlib
import json
import math
from pathlib import Path
import re
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from seller_agent.config import AppCredentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.ozon.adapter import OzonSellerAdapter
from seller_agent.reports.writer import ensure_dir, write_json
from seller_agent.tasks.ozon_stock_supply_monitor import CONFIRMED_INBOUND_STATES, MONITORED_ORDER_STATES


MOSCOW = ZoneInfo("Europe/Moscow")
CAPACITY_TARGET_DAYS = 30
SALES_LOOKBACK_DAYS = 90
RECENT_SALES_DAYS = 30
MAX_CLUSTERS = 20
MANUFACTURED_GROUPS = {"chev", "nash", "loop"}
MANUFACTURED_RE = re.compile(r"(шеврон|нашив|петлиц|позывн|chev|nash|loop|patch|патч)", re.IGNORECASE)
EXCLUDED_RE = re.compile(r"(головн|кепк[аи]\s+как\s+товар|шапк|панам|подсум|фартук|фальшпогон|брелок)", re.IGNORECASE)
RUN_ID_RE = re.compile(r"^[a-zA-Z0-9_-]{8,180}$")


@dataclass
class OzonNeedRow:
    offer_id: str
    sku: str
    product_id: str
    title: str
    internal_sku: str
    cluster_id: str
    cluster: str
    cluster_rank: int
    product_type: str
    classification_confirmed: bool
    batch_multiple: int
    pack_qty: int
    sales_90: int
    sales_30: int
    avg_daily_90: float
    avg_daily_30: float
    forecast_daily: float
    free_stock: int
    reserved: int
    promised: int
    confirmed_inbound: int
    coverage_before: float
    target_days: int
    raw_need: float
    rounded_need: int
    quantity: int = 0

    @property
    def physical_quantity(self) -> int:
        return self.quantity * self.pack_qty

    @property
    def rounded_need_physical(self) -> int:
        return self.rounded_need * self.pack_qty

    @property
    def coverage_after(self) -> float:
        if self.forecast_daily <= 0:
            return 999.0
        return (self.free_stock + self.confirmed_inbound + self.quantity) / self.forecast_daily


def run_ozon_production_work_plan(
    *,
    data_dir: Path,
    mode: str,
    value: int,
    cluster_count: int,
    credentials: AppCredentials | None = None,
    run_id: str | None = None,
    today: date | None = None,
    adapter: OzonSellerAdapter | None = None,
) -> dict[str, Any]:
    normalized_mode = _text(mode).lower()
    if normalized_mode not in {"capacity", "coverage_days"}:
        raise ValueError("mode must be capacity or coverage_days")
    value = _int(value)
    cluster_count = _int(cluster_count)
    if value <= 0:
        raise ValueError("value must be a positive integer")
    if normalized_mode == "coverage_days" and value > 365:
        raise ValueError("coverage_days must not exceed 365")
    if normalized_mode == "capacity" and value > 100_000:
        raise ValueError("capacity must not exceed 100000 physical pieces")
    if cluster_count <= 0 or cluster_count > MAX_CLUSTERS:
        raise ValueError(f"cluster_count must be between 1 and {MAX_CLUSTERS}")

    started_at = datetime.now(MOSCOW)
    today = today or started_at.date()
    period_to = today - timedelta(days=1)
    period_from = period_to - timedelta(days=SALES_LOOKBACK_DAYS - 1)
    recent_from = period_to - timedelta(days=RECENT_SALES_DAYS - 1)
    target_days = value if normalized_mode == "coverage_days" else CAPACITY_TARGET_DAYS
    run_id = run_id or f"ozon_production_work_plan_{started_at.strftime('%Y%m%dT%H%M%S')}"
    run_dir = ensure_dir(data_dir / "runs" / today.isoformat() / run_id)
    raw_dir = ensure_dir(run_dir / "raw")
    processed_dir = ensure_dir(run_dir / "processed")

    if adapter is None:
        if credentials is None or credentials.ozon_seller is None:
            raise ValueError("Ozon Seller API credentials are not configured")
        adapter = OzonSellerAdapter(credentials.ozon_seller)

    catalog = _load_catalog(data_dir / "catalog" / "unified" / "products.csv")
    product_list = adapter.fetch_product_list(visibility="ALL")
    product_ids = [_text(row.get("product_id")) for row in product_list if _text(row.get("product_id"))]
    product_info = adapter.fetch_product_info(product_ids)
    clusters = adapter.fetch_fbo_clusters()
    warehouse_stocks = adapter.fetch_stock_on_warehouses()
    postings = adapter.fetch_fbo_postings(
        since=_period_start(period_from),
        to=_period_end(period_to),
        status="",
        limit=100,
    )
    write_json(raw_dir / "ozon_product_list.json", product_list)
    write_json(raw_dir / "ozon_product_info.json", product_info)
    write_json(raw_dir / "ozon_clusters.json", clusters)
    write_json(raw_dir / "ozon_stock_on_warehouses.json", warehouse_stocks)
    write_json(raw_dir / "ozon_fbo_postings_90d.json", postings)

    controls: list[dict[str, Any]] = []
    warnings: list[str] = []
    identity = _product_identity(product_list, product_info)
    cluster_maps = _cluster_maps(clusters)
    if not cluster_maps["by_id"] or not cluster_maps["warehouse_to_cluster"]:
        raise ValueError("Ozon cluster list did not return cluster and warehouse mapping")

    stock = _stock_indexes(
        warehouse_stocks,
        identity=identity,
        warehouse_to_cluster=cluster_maps["warehouse_to_cluster"],
    )
    controls.extend(stock["controls"])
    sales = _sales_indexes(
        postings,
        identity=identity,
        cluster_name_map=cluster_maps["by_name"],
        recent_from=recent_from,
        period_to=period_to,
    )
    controls.extend(sales["controls"])
    inbound = _fetch_confirmed_inbound(
        adapter,
        identity=identity,
        cluster_by_id=cluster_maps["by_id"],
    )
    controls.extend(inbound["controls"])
    warnings.extend(inbound["warnings"])
    write_json(raw_dir / "ozon_active_supply_orders.json", inbound["raw"])

    candidate_rows: list[OzonNeedRow] = []
    missing_catalog_clusters: dict[str, set[str]] = defaultdict(set)
    for offer_id, cluster in sorted(sales["sales_90"]):
        qty_90 = sales["sales_90"][(offer_id, cluster)]
        if qty_90 <= 0:
            continue
        item_identity = identity.get(f"offer:{offer_id}") or {}
        sku = _text(item_identity.get("sku"))
        product_id = _text(item_identity.get("product_id"))
        catalog_row = (
            catalog.get(f"offer:{offer_id}")
            or (catalog.get(f"sku:{sku}") if sku else None)
            or (catalog.get(f"product:{product_id}") if product_id else None)
        )
        if not catalog_row:
            missing_catalog_clusters[offer_id].add(cluster)
            continue
        title = _text(item_identity.get("title") or catalog_row.get("product_name"))
        if not _is_manufactured(title, offer_id, _text(catalog_row.get("product_group"))):
            continue
        pack_qty = max(1, _int(catalog_row.get("pack_qty"), 1))
        product_type, batch_multiple, confirmed_type = _classify_product(
            title=title,
            internal_sku=_text(catalog_row.get("internal_sku")),
            offer_id=offer_id,
        )
        qty_30 = sales["sales_30"][(offer_id, cluster)]
        avg_90 = qty_90 / SALES_LOOKBACK_DAYS
        avg_30 = qty_30 / RECENT_SALES_DAYS
        forecast_daily = avg_30 * 0.6 + avg_90 * 0.4
        free = stock["free"][(offer_id, cluster)]
        reserved = stock["reserved"][(offer_id, cluster)]
        promised = stock["promised"][(offer_id, cluster)]
        inbound_qty = inbound["by_offer_cluster"][(offer_id, cluster)]
        if forecast_daily <= 0:
            continue
        raw_need = max(0.0, forecast_daily * target_days - free - inbound_qty)
        if raw_need <= 0:
            continue
        cluster_id = cluster_maps["id_by_name"].get(_normalize(cluster), "")
        candidate_rows.append(
            OzonNeedRow(
                offer_id=offer_id,
                sku=sku,
                product_id=product_id,
                title=title,
                internal_sku=_text(catalog_row.get("internal_sku")) or offer_id,
                cluster_id=cluster_id,
                cluster=cluster,
                cluster_rank=0,
                product_type=product_type,
                classification_confirmed=confirmed_type,
                batch_multiple=batch_multiple,
                pack_qty=pack_qty,
                sales_90=qty_90,
                sales_30=qty_30,
                avg_daily_90=avg_90,
                avg_daily_30=avg_30,
                forecast_daily=forecast_daily,
                free_stock=free,
                reserved=reserved,
                promised=promised,
                confirmed_inbound=inbound_qty,
                coverage_before=(free + inbound_qty) / forecast_daily,
                target_days=target_days,
                raw_need=raw_need,
                rounded_need=0,
            )
        )

    for offer_id, affected_clusters in sorted(missing_catalog_clusters.items()):
        controls.append(
            _control(
                offer_id,
                ", ".join(sorted(affected_clusters)),
                "missing_catalog_mapping",
                "Нет подтвержденной строки unified catalog; товар исключен из автоплана.",
            )
        )

    priority_rows = _cluster_priorities(candidate_rows)
    selected_clusters = [row["cluster"] for row in priority_rows[:cluster_count]]
    selected_set = set(selected_clusters)
    rank_by_cluster = {cluster: rank for rank, cluster in enumerate(selected_clusters, start=1)}
    rows = [row for row in candidate_rows if row.cluster in selected_set]
    for row in rows:
        row.cluster_rank = rank_by_cluster[row.cluster]
    _assign_rounded_needs(rows)
    unknown_type_offers = {row.offer_id for row in rows if not row.classification_confirmed}
    for offer_id in sorted(unknown_type_offers):
        affected_clusters = sorted({row.cluster for row in rows if row.offer_id == offer_id})
        controls.append(
            _control(
                offer_id,
                ", ".join(affected_clusters),
                "unknown_product_type_default_8",
                "Тип не определен; по правилу владельца применена кратность 8.",
            )
        )

    if normalized_mode == "coverage_days":
        for row in rows:
            row.quantity = row.rounded_need
        unused_capacity = 0
    else:
        unused_capacity = _allocate_capacity(rows, value)
    plan_rows = [row for row in rows if row.quantity > 0]

    if len(selected_clusters) < cluster_count:
        warnings.append(
            f"Запрошено кластеров: {cluster_count}; с подтвержденной положительной потребностью найдено: "
            f"{len(selected_clusters)}."
        )
    if unknown_type_offers:
        warnings.append(f"Для {len(unknown_type_offers)} товаров тип изделия не определен; применена кратность 8.")
    if not plan_rows:
        warnings.append("После кластерного учета остатков, confirmed inbound и кратностей производственный план пуст.")
    if normalized_mode == "capacity" and unused_capacity:
        warnings.append(
            f"Не распределено {unused_capacity} физических изделий: следующий допустимый производственный блок "
            "не помещается либо потребность выбранных кластеров исчерпана."
        )

    final_cluster_totals = _cluster_totals(plan_rows, priority_rows)
    xlsx_path = run_dir / f"v_rabotu_ozon_{normalized_mode}_{value}_{len(selected_clusters)}cl_{today.isoformat()}.xlsx"
    markdown_path = run_dir / "ozon_production_work_plan.md"
    rows_csv = processed_dir / "ozon_production_rows.csv"
    controls_csv = processed_dir / "control_rows.csv"
    clusters_csv = processed_dir / "cluster_priorities.csv"
    _write_workbook(
        xlsx_path,
        plan_rows,
        controls,
        priority_rows,
        normalized_mode,
        value,
        cluster_count,
        period_from,
        period_to,
    )
    _write_csv(rows_csv, [_row_dict(row) for row in plan_rows])
    _write_csv(controls_csv, controls)
    _write_csv(clusters_csv, priority_rows)
    plan_checksum = _file_sha256(xlsx_path)
    physical_total = sum(row.physical_quantity for row in plan_rows)
    units_total = sum(row.quantity for row in plan_rows)
    summary: dict[str, Any] = {
        "task": "ozon-production-work-plan",
        "run_id": run_id,
        "started_at": started_at.isoformat(timespec="seconds"),
        "overall_status": "warning" if warnings or controls else "ok",
        "mode": "read_only",
        "calculation_mode": normalized_mode,
        "input_value": value,
        "requested_cluster_count": cluster_count,
        "target_days": target_days,
        "sales_period": {
            "from": period_from.isoformat(),
            "to": period_to.isoformat(),
            "days": SALES_LOOKBACK_DAYS,
        },
        "metrics": {
            "marketplace_units": units_total,
            "physical_pieces": physical_total,
            "articles": len({row.offer_id for row in plan_rows}),
            "clusters": len({row.cluster for row in plan_rows}),
            "selected_clusters": len(selected_clusters),
            "positive_need_clusters": len(priority_rows),
            "plan_rows": len(plan_rows),
            "control_rows": len(controls),
            "unknown_type_rows": len(unknown_type_offers),
            "requested_capacity_physical": value if normalized_mode == "capacity" else 0,
            "unused_capacity_physical": unused_capacity,
        },
        "cluster_totals": final_cluster_totals,
        "cluster_priorities": priority_rows,
        "warnings": warnings,
        "approval_status": "pending_owner_review",
        "plan_checksum": plan_checksum,
        "sources": [
            "Ozon POST /v3/product/list",
            "Ozon POST /v3/product/info/list",
            "Ozon POST /v2/cluster/list",
            "Ozon POST /v2/analytics/stock_on_warehouses",
            "Ozon POST /v2/posting/fbo/list",
            "Ozon supply-order list/get/bundle",
            "data/catalog/unified/products.csv",
        ],
        "artifacts": {
            "report": str(xlsx_path),
            "markdown": str(markdown_path),
            "rows_csv": str(rows_csv),
            "controls_csv": str(controls_csv),
            "clusters_csv": str(clusters_csv),
            "summary": str(run_dir / "summary.json"),
        },
        "next_step": "Проверить Excel и отдельно утвердить или отклонить план. Поставка в Ozon не создавалась.",
    }
    _write_report(markdown_path, summary)
    write_json(run_dir / "summary.json", summary)
    manifest = write_summary_run_manifest(
        data_dir=data_dir,
        run_dir=run_dir,
        summary=summary,
        task="ozon-production-work-plan",
        mode="read_only",
        risk="low",
        marketplaces=["ozon"],
        inputs={
            "mode": normalized_mode,
            "value": value,
            "cluster_count": cluster_count,
            "target_days": target_days,
        },
        pending_id=run_id,
        lifecycle_status="pending_review",
        closed=False,
    )
    summary["artifacts"].update(manifest)
    write_json(run_dir / "summary.json", summary)
    return summary


def record_ozon_work_plan_decision(*, data_dir: Path, plan_run_id: str, approved: bool) -> dict[str, Any]:
    plan_run_id = _text(plan_run_id)
    if not RUN_ID_RE.fullmatch(plan_run_id):
        raise ValueError("invalid plan_run_id")
    candidates = list((data_dir / "runs").glob(f"*/{plan_run_id}/summary.json"))
    if len(candidates) != 1:
        raise ValueError("Ozon work plan not found or ambiguous")
    summary_path = candidates[0]
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    if not isinstance(summary, dict) or summary.get("task") != "ozon-production-work-plan":
        raise ValueError("run is not an Ozon production work plan")
    report_path = Path(_text((summary.get("artifacts") or {}).get("report")))
    if not report_path.exists() or _file_sha256(report_path) != summary.get("plan_checksum"):
        raise ValueError("plan file checksum mismatch")
    status = "owner_approved" if approved else "owner_rejected"
    decision = {
        "plan_run_id": plan_run_id,
        "status": status,
        "decided_at": datetime.now(MOSCOW).isoformat(timespec="seconds"),
        "plan_checksum": summary.get("plan_checksum"),
        "marketplace_write_performed": False,
    }
    decision_path = summary_path.parent / "owner_decision.json"
    write_json(decision_path, decision)
    summary["approval_status"] = status
    summary["artifacts"]["owner_decision"] = str(decision_path)
    summary["next_step"] = (
        "План утвержден в работу; производство и поставка исполняются вне этого read-only контура."
        if approved
        else "План отклонен; производство и поставка по нему не запускаются."
    )
    manifest = write_summary_run_manifest(
        data_dir=data_dir,
        run_dir=summary_path.parent,
        summary=summary,
        task="ozon-production-work-plan",
        mode="read_only",
        risk="low",
        marketplaces=["ozon"],
        inputs={
            "mode": summary.get("calculation_mode"),
            "value": summary.get("input_value"),
            "cluster_count": summary.get("requested_cluster_count"),
        },
        pending_id=plan_run_id,
        approved_id=plan_run_id if approved else None,
        lifecycle_status="approved" if approved else "closed",
        closed=True,
    )
    summary["artifacts"].update(manifest)
    write_json(summary_path, summary)
    return decision


def _cluster_maps(rows: list[dict[str, Any]]) -> dict[str, dict[str, str]]:
    by_id: dict[str, str] = {}
    by_name: dict[str, str] = {}
    id_by_name: dict[str, str] = {}
    warehouse_to_cluster: dict[str, str] = {}
    for row in rows:
        cluster_id = _text(row.get("macrolocal_cluster_id"))
        data = row.get("data") if isinstance(row.get("data"), dict) else {}
        macrolocal = data.get("macrolocal_cluster") if isinstance(data.get("macrolocal_cluster"), dict) else {}
        name = _text(macrolocal.get("name"))
        if not cluster_id or not name:
            continue
        by_id[cluster_id] = name
        by_name[_normalize(name)] = name
        id_by_name[_normalize(name)] = cluster_id
        for warehouse in data.get("fulfillments") or []:
            if not isinstance(warehouse, dict):
                continue
            warehouse_name = _text(warehouse.get("name"))
            if warehouse_name:
                warehouse_to_cluster[_warehouse_key(warehouse_name)] = name
    return {
        "by_id": by_id,
        "by_name": by_name,
        "id_by_name": id_by_name,
        "warehouse_to_cluster": warehouse_to_cluster,
    }


def _stock_indexes(
    rows: list[dict[str, Any]],
    *,
    identity: dict[str, dict[str, str]],
    warehouse_to_cluster: dict[str, str],
) -> dict[str, Any]:
    free: Counter[tuple[str, str]] = Counter()
    reserved: Counter[tuple[str, str]] = Counter()
    promised: Counter[tuple[str, str]] = Counter()
    controls: list[dict[str, Any]] = []
    unknown_warehouses: Counter[str] = Counter()
    for row in rows:
        warehouse = _text(row.get("warehouse_name"))
        cluster = warehouse_to_cluster.get(_warehouse_key(warehouse), "")
        if not cluster:
            unknown_warehouses[warehouse or "Не указан"] += 1
            continue
        offer_id = _text(row.get("item_code") or row.get("offer_id"))
        sku = _text(row.get("sku"))
        if not offer_id:
            offer_id = _text((identity.get(f"sku:{sku}") or {}).get("offer_id"))
        if not offer_id:
            controls.append(_control("", cluster, "stock_offer_not_mapped", f"SKU {sku}: не найден offer_id."))
            continue
        key = (offer_id, cluster)
        free[key] += max(0, _int(row.get("free_to_sell_amount")))
        reserved[key] += max(0, _int(row.get("reserved_amount")))
        promised[key] += max(0, _int(row.get("promised_amount")))
    for warehouse, count in sorted(unknown_warehouses.items()):
        controls.append(
            _control("", "", "warehouse_cluster_not_mapped", f"Склад {warehouse}: {count} строк без mapping кластера.")
        )
    return {"free": free, "reserved": reserved, "promised": promised, "controls": controls}


def _sales_indexes(
    rows: list[dict[str, Any]],
    *,
    identity: dict[str, dict[str, str]],
    cluster_name_map: dict[str, str],
    recent_from: date,
    period_to: date,
) -> dict[str, Any]:
    sales_90: Counter[tuple[str, str]] = Counter()
    sales_30: Counter[tuple[str, str]] = Counter()
    controls: list[dict[str, Any]] = []
    unknown_clusters: Counter[str] = Counter()
    for posting in rows:
        if _text(posting.get("status")).lower() == "cancelled":
            continue
        financial = posting.get("financial_data") if isinstance(posting.get("financial_data"), dict) else {}
        source_cluster = _text(financial.get("cluster_to"))
        cluster = cluster_name_map.get(_normalize(source_cluster), "")
        if not cluster:
            unknown_clusters[source_cluster or "Не указан"] += 1
            continue
        posting_date = _parse_date(posting.get("created_at") or posting.get("in_process_at"))
        for product in posting.get("products") or []:
            if not isinstance(product, dict):
                continue
            offer_id = _text(product.get("offer_id"))
            sku = _text(product.get("sku"))
            if not offer_id:
                offer_id = _text((identity.get(f"sku:{sku}") or {}).get("offer_id"))
            if not offer_id:
                continue
            quantity = max(0, _int(product.get("quantity")))
            sales_90[(offer_id, cluster)] += quantity
            if posting_date and recent_from <= posting_date <= period_to:
                sales_30[(offer_id, cluster)] += quantity
    for cluster, count in sorted(unknown_clusters.items()):
        controls.append(
            _control("", cluster, "sales_cluster_not_mapped", f"Кластер продаж не найден в /v2/cluster/list: {count} posting-ов.")
        )
    return {"sales_90": sales_90, "sales_30": sales_30, "controls": controls}


def _fetch_confirmed_inbound(
    adapter: OzonSellerAdapter,
    *,
    identity: dict[str, dict[str, str]],
    cluster_by_id: dict[str, str],
) -> dict[str, Any]:
    by_offer_cluster: Counter[tuple[str, str]] = Counter()
    controls: list[dict[str, Any]] = []
    warnings: list[str] = []
    order_ids = adapter.fetch_supply_order_ids(states=MONITORED_ORDER_STATES)
    orders = adapter.fetch_supply_orders(order_ids) if order_ids else []
    raw: dict[str, Any] = {"order_ids": order_ids, "orders": orders, "bundles": []}
    virtual_orders = 0
    for order in orders:
        tags = order.get("order_tags") if isinstance(order.get("order_tags"), dict) else {}
        if bool(tags.get("is_virtual")):
            virtual_orders += 1
            continue
        order_id = _text(order.get("order_id") or order.get("order_number"))
        order_state = _text(order.get("state"))
        for supply in order.get("supplies") or []:
            if not isinstance(supply, dict):
                continue
            state = _text(supply.get("state")) or order_state
            if state not in CONFIRMED_INBOUND_STATES:
                continue
            cluster_id = _text(supply.get("macrolocal_cluster_id"))
            cluster = cluster_by_id.get(cluster_id, "")
            if not cluster:
                controls.append(
                    _control("", cluster_id, "inbound_cluster_not_mapped", f"Поставка {order_id}: неизвестный cluster_id.")
                )
                continue
            bundle_id = _text(supply.get("bundle_id"))
            if not bundle_id:
                controls.append(
                    _control("", cluster, "inbound_bundle_missing", f"Поставка {order_id}: состав не подтвержден.")
                )
                continue
            try:
                goods = adapter.fetch_supply_order_bundle(bundle_id)
            except Exception as exc:  # noqa: BLE001 - one failed supply must remain visible in control.
                controls.append(
                    _control("", cluster, "inbound_bundle_failed", f"Поставка {order_id}: {_safe_error(exc)}")
                )
                continue
            raw["bundles"].append({"order_id": order_id, "cluster_id": cluster_id, "bundle_id": bundle_id, "items": goods})
            for good in goods:
                offer_id = _text(good.get("offer_id"))
                sku = _text(good.get("sku"))
                if not offer_id:
                    offer_id = _text((identity.get(f"sku:{sku}") or {}).get("offer_id"))
                quantity = max(0, _int(good.get("quantity")))
                if offer_id and quantity:
                    by_offer_cluster[(offer_id, cluster)] += quantity
    if virtual_orders:
        warnings.append(f"Исключено виртуальных заявок-дублей: {virtual_orders}.")
    return {
        "by_offer_cluster": by_offer_cluster,
        "controls": controls,
        "warnings": warnings,
        "raw": raw,
    }


def _assign_rounded_needs(rows: list[OzonNeedRow]) -> None:
    by_offer: dict[str, list[OzonNeedRow]] = defaultdict(list)
    for row in rows:
        row.rounded_need = 0
        by_offer[row.offer_id].append(row)
    for offer_rows in by_offer.values():
        multiple = offer_rows[0].batch_multiple
        target = _round_up_units(sum(row.raw_need for row in offer_rows), multiple)
        allocated = 0
        while allocated < target:
            selected = max(
                offer_rows,
                key=lambda row: (
                    row.raw_need - row.rounded_need,
                    row.forecast_daily,
                    -row.cluster_rank,
                    row.cluster,
                ),
            )
            selected.rounded_need += multiple
            allocated += multiple


def _cluster_priorities(rows: list[OzonNeedRow]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "raw_need_marketplace_units": 0.0,
            "raw_need_physical_pieces": 0.0,
            "sales_90": 0,
            "sales_30": 0,
            "articles": set(),
            "cluster_id": "",
        }
    )
    for row in rows:
        target = grouped[row.cluster]
        target["cluster_id"] = row.cluster_id
        target["raw_need_marketplace_units"] += row.raw_need
        target["raw_need_physical_pieces"] += row.raw_need * row.pack_qty
        target["sales_90"] += row.sales_90
        target["sales_30"] += row.sales_30
        target["articles"].add(row.offer_id)
    ordered = sorted(
        grouped.items(),
        key=lambda item: (-item[1]["raw_need_physical_pieces"], -item[1]["sales_30"], item[0]),
    )
    return [
        {
            "priority": rank,
            "cluster": cluster,
            "cluster_id": values["cluster_id"],
            "raw_need_marketplace_units": round(values["raw_need_marketplace_units"], 2),
            "raw_need_physical_pieces": round(values["raw_need_physical_pieces"], 2),
            "sales_90": values["sales_90"],
            "sales_30": values["sales_30"],
            "articles": len(values["articles"]),
        }
        for rank, (cluster, values) in enumerate(ordered, start=1)
    ]


def _allocate_capacity(rows: list[OzonNeedRow], capacity_physical: int) -> int:
    remaining = capacity_physical
    cluster_need = Counter({row.cluster: 0 for row in rows})
    cluster_allocated: Counter[str] = Counter()
    for row in rows:
        cluster_need[row.cluster] += row.rounded_need_physical
    while True:
        fitting = [
            row
            for row in rows
            if row.quantity < row.rounded_need and row.batch_multiple * row.pack_qty <= remaining
        ]
        if not fitting:
            break
        selected = max(
            fitting,
            key=lambda row: (
                1 - cluster_allocated[row.cluster] / max(1, cluster_need[row.cluster]),
                cluster_need[row.cluster],
                1 - row.quantity / max(1, row.rounded_need),
                max(0.0, row.target_days - row.coverage_after),
                row.forecast_daily,
                -row.cluster_rank,
                row.offer_id,
            ),
        )
        physical_block = selected.batch_multiple * selected.pack_qty
        selected.quantity += selected.batch_multiple
        cluster_allocated[selected.cluster] += physical_block
        remaining -= physical_block
    return remaining


def _cluster_totals(rows: list[OzonNeedRow], priorities: list[dict[str, Any]]) -> list[dict[str, Any]]:
    priority_by_cluster = {row["cluster"]: row for row in priorities}
    grouped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"marketplace_units": 0, "physical_pieces": 0, "articles": set()}
    )
    for row in rows:
        grouped[row.cluster]["marketplace_units"] += row.quantity
        grouped[row.cluster]["physical_pieces"] += row.physical_quantity
        grouped[row.cluster]["articles"].add(row.offer_id)
    result: list[dict[str, Any]] = []
    for cluster, values in grouped.items():
        priority = priority_by_cluster.get(cluster, {})
        result.append(
            {
                "priority": priority.get("priority", 0),
                "cluster": cluster,
                "cluster_id": priority.get("cluster_id", ""),
                "marketplace_units": values["marketplace_units"],
                "physical_pieces": values["physical_pieces"],
                "articles": len(values["articles"]),
                "raw_need_physical_pieces": priority.get("raw_need_physical_pieces", 0),
            }
        )
    return sorted(result, key=lambda row: (row["priority"], row["cluster"]))


def _product_identity(
    product_list: list[dict[str, Any]], product_info: list[dict[str, Any]]
) -> dict[str, dict[str, str]]:
    index: dict[str, dict[str, str]] = {}
    for source in [*product_list, *product_info]:
        offer_id = _text(source.get("offer_id"))
        product_id = _text(source.get("product_id") or source.get("id"))
        sku = _text(source.get("sku") or source.get("fbo_sku") or source.get("fboSku"))
        value = {
            "offer_id": offer_id,
            "product_id": product_id,
            "sku": sku,
            "title": _text(source.get("name") or source.get("title")),
        }
        if offer_id:
            existing = index.setdefault(f"offer:{offer_id}", {})
            existing.update({key: item for key, item in value.items() if item})
            value = existing
        if product_id:
            index[f"product:{product_id}"] = value
        if sku:
            index[f"sku:{sku}"] = value
    return index


def _load_catalog(path: Path) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}
    if not path.exists():
        return index
    with path.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            for prefix, field in (
                ("offer", "ozon_offer_id"),
                ("sku", "ozon_sku"),
                ("product", "ozon_product_id"),
            ):
                identifier = _text(row.get(field))
                if identifier:
                    index[f"{prefix}:{identifier}"] = row
    return index


def _classify_product(*, title: str, internal_sku: str, offer_id: str) -> tuple[str, int, bool]:
    text = _normalize(" ".join((title, internal_sku, offer_id)))
    sku = _normalize(f"{internal_sku} {offer_id}")
    if ("комплект" in text and "позывн" in text) or ("kit" in sku and "pz" in sku):
        return "Комплект позывных", 6, True
    if "на спину" in text or "наспин" in text or "back" in sku:
        return "Наспинный", 2, True
    if "нагруд" in text or re.search(r"(^|_)ng(_|$)", internal_sku.lower()):
        return "Нагрудный", 12, True
    if "на кепк" in text or re.search(r"(^|_)kp(_|$)", internal_sku.lower()):
        return "На кепку", 9, True
    if "нарукав" in text or re.search(r"(^|_)nr(_|$)", internal_sku.lower()):
        return "Нарукавный", 8, True
    return "Тип не определен", 8, False


def _is_manufactured(title: str, offer_id: str, product_group: str) -> bool:
    text = " ".join((title, offer_id, product_group))
    if EXCLUDED_RE.search(text):
        return False
    return product_group.lower() in MANUFACTURED_GROUPS or bool(MANUFACTURED_RE.search(text))


def _row_dict(row: OzonNeedRow) -> dict[str, Any]:
    return {
        "cluster_priority": row.cluster_rank,
        "cluster": row.cluster,
        "cluster_id": row.cluster_id,
        "internal_sku": row.internal_sku,
        "offer_id": row.offer_id,
        "sku": row.sku,
        "product_id": row.product_id,
        "title": row.title,
        "product_type": row.product_type,
        "classification_confirmed": row.classification_confirmed,
        "batch_multiple": row.batch_multiple,
        "pack_qty": row.pack_qty,
        "quantity": row.quantity,
        "physical_quantity": row.physical_quantity,
        "sales_90": row.sales_90,
        "sales_30": row.sales_30,
        "avg_daily_90": round(row.avg_daily_90, 4),
        "avg_daily_30": round(row.avg_daily_30, 4),
        "forecast_daily": round(row.forecast_daily, 4),
        "free_stock_in_cluster": row.free_stock,
        "reserved_in_cluster": row.reserved,
        "promised_in_cluster_reference_only": row.promised,
        "confirmed_inbound_to_cluster": row.confirmed_inbound,
        "coverage_before": round(row.coverage_before, 1),
        "target_days": row.target_days,
        "raw_need": round(row.raw_need, 2),
        "rounded_need": row.rounded_need,
        "coverage_after": round(row.coverage_after, 1),
    }


def _write_workbook(
    path: Path,
    rows: list[OzonNeedRow],
    controls: list[dict[str, Any]],
    priorities: list[dict[str, Any]],
    mode: str,
    value: int,
    cluster_count: int,
    period_from: date,
    period_to: date,
) -> None:
    workbook = Workbook()
    articles = workbook.active
    articles.title = "Артикулы"
    grouped: dict[str, dict[str, Any]] = {}
    for row in rows:
        target = grouped.setdefault(
            row.offer_id,
            {
                "Внутренний артикул": row.internal_sku,
                "Артикул продавца Ozon": row.offer_id,
                "SKU Ozon": row.sku,
                "Название": row.title,
                "Тип изделия": row.product_type,
                "Кратность": row.batch_multiple,
                "Изделий в товаре": row.pack_qty,
                "Товарных единиц Ozon": 0,
                "Физических изделий": 0,
                "Кластеры назначения": set(),
                "Классификация подтверждена": "да" if row.classification_confirmed else "нет; кратность 8",
            },
        )
        target["Товарных единиц Ozon"] += row.quantity
        target["Физических изделий"] += row.physical_quantity
        target["Кластеры назначения"].add(row.cluster)
    article_rows: list[dict[str, Any]] = []
    for item in grouped.values():
        item["Кластеры назначения"] = ", ".join(sorted(item["Кластеры назначения"]))
        article_rows.append(item)
    _append_dict_rows(
        articles,
        sorted(article_rows, key=lambda row: (-_int(row["Физических изделий"]), str(row["Артикул продавца Ozon"]))),
    )

    cluster_sheet = workbook.create_sheet("Ozon кластеры")
    _append_dict_rows(
        cluster_sheet,
        [_row_dict(row) for row in sorted(rows, key=lambda item: (item.cluster_rank, -item.physical_quantity, item.offer_id))],
    )

    control = workbook.create_sheet("Контроль")
    parameters = [
        {"Параметр": "Режим", "Значение": "По производственной возможности" if mode == "capacity" else "По периоду покрытия"},
        {"Параметр": "Введенное значение", "Значение": value},
        {"Параметр": "Запрошено кластеров", "Значение": cluster_count},
        {"Параметр": "Период продаж", "Значение": f"{period_from.isoformat()} - {period_to.isoformat()}"},
        {"Параметр": "Формула", "Значение": "спрос товара в кластере - свободный остаток в кластере - confirmed inbound в кластер"},
        {"Параметр": "Статус", "Значение": "Ожидает проверки владельца; поставка Ozon не создавалась"},
    ]
    _append_dict_rows(control, parameters)
    control.append([])
    _append_dict_rows(control, priorities, start_row=control.max_row + 1)
    if controls:
        control.append([])
        _append_dict_rows(control, controls, start_row=control.max_row + 1)
    workbook.save(path)


def _write_report(path: Path, summary: dict[str, Any]) -> None:
    metrics = summary["metrics"]
    mode_label = "по производственной возможности" if summary["calculation_mode"] == "capacity" else "по периоду покрытия"
    lines = [
        "# Ozon: файл в работу",
        "",
        f"- Режим: `{mode_label}`.",
        f"- Введенное значение: `{summary['input_value']}`.",
        f"- Запрошено кластеров: `{summary['requested_cluster_count']}`; выбрано: `{metrics['selected_clusters']}`.",
        f"- Целевое покрытие: `{summary['target_days']}` дней.",
        f"- Период спроса: `{summary['sales_period']['from']} - {summary['sales_period']['to']}`.",
        f"- Товарных единиц Ozon: `{metrics['marketplace_units']}`.",
        f"- Физических изделий: `{metrics['physical_pieces']}`.",
        f"- Артикулов: `{metrics['articles']}`; кластеров назначения: `{metrics['clusters']}`.",
        f"- Контрольных строк: `{metrics['control_rows']}`.",
        "",
        "Потребность рассчитана отдельно для каждой пары товар × кластер. Общий остаток Ozon не вычитался.",
        "",
        "## Распределение",
        "",
    ]
    for row in summary["cluster_totals"]:
        lines.append(
            f"- #{row['priority']} {row['cluster']}: `{row['marketplace_units']}` ед. Ozon / "
            f"`{row['physical_pieces']}` физических изделий / `{row['articles']}` артикулов."
        )
    if summary["warnings"]:
        lines.extend(["", "## Ограничения", ""])
        lines.extend(f"- {warning}" for warning in summary["warnings"])
    lines.extend(
        [
            "",
            "## Следующий шаг",
            "",
            "Проверить Excel и отдельно утвердить или отклонить план. Производство и поставка Ozon автоматически не создавались.",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _append_dict_rows(sheet: Any, rows: list[dict[str, Any]], *, start_row: int = 1) -> None:
    fields: list[str] = []
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    if not fields:
        fields = ["Статус"]
        rows = [{"Статус": "Нет строк"}]
    for column, field in enumerate(fields, start=1):
        sheet.cell(start_row, column, field)
    for row_number, row in enumerate(rows, start=start_row + 1):
        for column, field in enumerate(fields, start=1):
            sheet.cell(row_number, column, row.get(field, ""))
    _sheet_setup(sheet, header_row=start_row)


def _sheet_setup(sheet: Any, *, header_row: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[header_row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        width = max(10, max((len(str(cell.value or "")) for cell in sheet[letter]), default=10))
        sheet.column_dimensions[letter].width = min(55, width + 2)


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields: list[str] = []
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    if not fields:
        fields = ["status"]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _control(offer_id: str, cluster: str, code: str, message: str) -> dict[str, Any]:
    return {"offer_id": offer_id, "cluster": cluster, "code": code, "message": message}


def _period_start(value: date) -> str:
    return f"{value.isoformat()}T00:00:00.000Z"


def _period_end(value: date) -> str:
    return f"{value.isoformat()}T23:59:59.999Z"


def _parse_date(value: Any) -> date | None:
    text = _text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def _round_up_units(value: float, multiple: int) -> int:
    if value <= 0:
        return 0
    return int(math.ceil(value / multiple) * multiple)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _warehouse_key(value: str) -> str:
    return re.sub(r"[^a-zа-я0-9]+", "_", _normalize(value)).strip("_")


def _normalize(value: Any) -> str:
    return " ".join(_text(value).replace("ё", "е").lower().split())


def _safe_error(exc: BaseException) -> str:
    return str(exc).replace("\n", " ").replace("\r", " ")[:500]


def _text(value: Any) -> str:
    return str(value or "").strip()


def _int(value: Any, default: int = 0) -> int:
    try:
        if value in (None, ""):
            return default
        return int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError):
        return default
