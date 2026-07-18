from __future__ import annotations

from collections import Counter, defaultdict
import csv
from dataclasses import dataclass
from datetime import date, datetime, timedelta
import hashlib
import json
import math
from pathlib import Path
import re
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from seller_agent.config import AppCredentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.wb.adapter import WbContentAdapter
from seller_agent.marketplaces.wb.analytics_adapter import WbAnalyticsAdapter
from seller_agent.marketplaces.wb.fbw_supplies_adapter import WbFbwSuppliesAdapter
from seller_agent.marketplaces.wb.statistics_adapter import WbStatisticsAdapter
from seller_agent.reports.writer import ensure_dir, write_json


MOSCOW = ZoneInfo("Europe/Moscow")
ACTIVE_INBOUND_STATUS_IDS = {2, 3, 4, 6}
CAPACITY_TARGET_DAYS = 30
SALES_LOOKBACK_DAYS = 90
RECENT_SALES_DAYS = 30

BUYER_REGION_MAP = {
    "центральный федеральный округ": "Центральный",
    "приволжский федеральный округ": "Приволжский",
    "северо-западный федеральный округ": "Северо-Западный",
    "сибирский федеральный округ": "Дальневосточный и Сибирский",
    "дальневосточный федеральный округ": "Дальневосточный и Сибирский",
    "уральский федеральный округ": "Уральский",
    "южный федеральный округ": "Южный и Северо-Кавказский",
    "северо-кавказский федеральный округ": "Южный и Северо-Кавказский",
}
MAX_DESTINATION_CLUSTERS = len(set(BUYER_REGION_MAP.values()))
WAREHOUSE_REGION_HINTS = {
    "шушар": "Северо-Западный",
    "санкт-петербург": "Северо-Западный",
    "электростал": "Центральный",
    "коледино": "Центральный",
    "владимир": "Центральный",
    "тула": "Центральный",
    "воронеж": "Центральный",
    "волгоград": "Южный и Северо-Кавказский",
    "краснодар": "Южный и Северо-Кавказский",
    "невинномысск": "Южный и Северо-Кавказский",
    "новосемейкино": "Приволжский",
    "казань": "Приволжский",
    "сарапул": "Приволжский",
    "екатеринбург": "Уральский",
    "омск": "Дальневосточный и Сибирский",
    "хабаровск": "Дальневосточный и Сибирский",
    "барнаул": "Дальневосточный и Сибирский",
    "чита": "Дальневосточный и Сибирский",
}
MANUFACTURED_GROUPS = {"chev", "nash", "loop"}
MANUFACTURED_RE = re.compile(r"(шеврон|нашив|петлиц|позывн|chev|nash|loop|patch|патч)", re.IGNORECASE)
EXCLUDED_RE = re.compile(r"(головн|кепк[аи]\s+как\s+товар|шапк|панам|подсум|фартук|фальшпогон|брелок)", re.IGNORECASE)
RUN_ID_RE = re.compile(r"^[a-zA-Z0-9_-]{8,180}$")


@dataclass
class NeedRow:
    vendor_code: str
    nm_id: str
    barcode: str
    title: str
    internal_sku: str
    region: str
    product_type: str
    classification_confirmed: bool
    batch_multiple: int
    pack_qty: int
    sales_90: int
    sales_30: int
    avg_daily_90: float
    avg_daily_30: float
    forecast_daily: float
    current_stock: int
    confirmed_inbound: int
    coverage_before: float
    target_days: int
    raw_need: float
    rounded_need: int
    cluster_rank: int = 0
    quantity: int = 0

    @property
    def physical_quantity(self) -> int:
        return self.quantity * self.pack_qty

    @property
    def coverage_after(self) -> float:
        if self.forecast_daily <= 0:
            return 999.0
        return (self.current_stock + self.confirmed_inbound + self.quantity) / self.forecast_daily


def run_wb_production_work_plan(
    *,
    data_dir: Path,
    mode: str,
    value: int,
    cluster_count: int,
    credentials: AppCredentials | None = None,
    run_id: str | None = None,
    today: date | None = None,
    analytics_adapter: WbAnalyticsAdapter | None = None,
    statistics_adapter: WbStatisticsAdapter | None = None,
    content_adapter: WbContentAdapter | None = None,
    supplies_adapter: WbFbwSuppliesAdapter | None = None,
) -> dict[str, Any]:
    normalized_mode = str(mode or "").strip().lower()
    if normalized_mode not in {"capacity", "coverage_days"}:
        raise ValueError("mode must be capacity or coverage_days")
    value = _int(value)
    cluster_count = _int(cluster_count)
    if value <= 0:
        raise ValueError("value must be a positive integer")
    if normalized_mode == "coverage_days" and value > 365:
        raise ValueError("coverage_days must not exceed 365")
    if normalized_mode == "capacity" and value > 100_000:
        raise ValueError("capacity must not exceed 100000 physical pieces")
    if cluster_count <= 0 or cluster_count > MAX_DESTINATION_CLUSTERS:
        raise ValueError(f"cluster_count must be between 1 and {MAX_DESTINATION_CLUSTERS}")

    started_at = datetime.now(MOSCOW)
    today = today or started_at.date()
    period_to = today - timedelta(days=1)
    period_from = period_to - timedelta(days=SALES_LOOKBACK_DAYS - 1)
    target_days = value if normalized_mode == "coverage_days" else CAPACITY_TARGET_DAYS
    run_id = run_id or f"wb_production_work_plan_{started_at.strftime('%Y%m%dT%H%M%S')}"
    run_dir = ensure_dir(data_dir / "runs" / today.isoformat() / run_id)
    raw_dir = ensure_dir(run_dir / "raw")
    processed_dir = ensure_dir(run_dir / "processed")

    if any(adapter is None for adapter in (analytics_adapter, statistics_adapter, content_adapter, supplies_adapter)):
        if credentials is None or credentials.wb is None:
            raise ValueError("WB API credentials are not configured")
        analytics_adapter = analytics_adapter or WbAnalyticsAdapter(credentials.wb)
        statistics_adapter = statistics_adapter or WbStatisticsAdapter(credentials.wb)
        content_adapter = content_adapter or WbContentAdapter(credentials.wb)
        supplies_adapter = supplies_adapter or WbFbwSuppliesAdapter(credentials.wb)

    catalog = _load_catalog(
        data_dir / "catalog" / "unified" / "products.csv",
        data_dir / "catalog" / "mapping" / "ozon_wb_internal_sku_confirmed.csv",
    )
    cards = content_adapter.fetch_cards(limit=100)
    stocks = analytics_adapter.fetch_wb_warehouse_stocks()
    sales = statistics_adapter.fetch_sales(date_from=period_from.isoformat(), flag=0)
    inbound = _fetch_inbound(supplies_adapter, stocks)
    write_json(raw_dir / "wb_cards.json", cards)
    write_json(raw_dir / "wb_warehouse_stocks.json", stocks)
    write_json(raw_dir / "wb_sales_90d.json", sales)
    write_json(raw_dir / "wb_active_inbound.json", inbound["raw"])

    warnings: list[str] = []
    controls: list[dict[str, Any]] = list(inbound["controls"])
    card_index = _card_index(cards)
    current_vendor_by_nm = {
        _text(card.get("nm_id")): vendor
        for vendor, card in card_index.items()
        if _text(card.get("nm_id"))
    }
    stock_index, anomalous_pairs, stock_warnings = _stock_index(stocks)
    warnings.extend(stock_warnings)
    sales_90, sales_30 = _sales_indexes(
        sales,
        period_from=period_from,
        period_to=period_to,
        current_vendor_by_nm=current_vendor_by_nm,
    )
    rows: list[NeedRow] = []
    possible_stockouts = 0
    unknown_types = 0

    vendors = sorted({vendor for vendor, _ in sales_90 if sales_90[(vendor, _)] > 0})
    for vendor in vendors:
        card = card_index.get(vendor) or {}
        nm_id = _text(card.get("nm_id"))
        catalog_row = catalog.get(f"vendor:{vendor}") or (catalog.get(f"nm:{nm_id}") if nm_id else None)
        if not catalog_row:
            controls.append(_control(vendor, nm_id, "missing_catalog_mapping", "Нет подтвержденной строки unified catalog."))
            continue
        title = _text(card.get("title") or catalog_row.get("product_name"))
        if not _is_manufactured(title, vendor, _text(catalog_row.get("product_group"))):
            continue
        barcode = _text(card.get("barcode") or catalog_row.get("wb_barcode"))
        if not nm_id:
            controls.append(_control(vendor, nm_id, "missing_nm_id", "Нет nmID WB; строка исключена из автоплана."))
            continue
        if not barcode:
            controls.append(_control(vendor, nm_id, "missing_barcode", "Нет штрихкода WB; строка исключена из автоплана."))
            continue
        pack_qty = max(1, _int(catalog_row.get("pack_qty"), 1))
        product_type, batch_multiple, confirmed_type = _classify_product(
            title=title,
            internal_sku=_text(catalog_row.get("internal_sku")),
            vendor_code=vendor,
        )
        if not confirmed_type:
            unknown_types += 1
            controls.append(
                _control(
                    vendor,
                    nm_id,
                    "unknown_product_type_default_8",
                    "Тип не определен; по правилу владельца применена кратность 8, строка оставлена в плане.",
                )
            )

        product_regions = sorted(region for item_vendor, region in sales_90 if item_vendor == vendor and sales_90[(vendor, region)] > 0)
        global_stock = sum(quantity for (item_nm, _), quantity in stock_index.items() if item_nm == nm_id)
        for region in product_regions:
            if (nm_id, region) in anomalous_pairs:
                controls.append(
                    _control(
                        vendor,
                        nm_id,
                        "stock_state_anomaly",
                        f"Регион {region}: WB переклассифицирует quantity/inWayFromClient; автопотребность временно не рассчитана.",
                    )
                )
                continue
            qty_90 = max(0, sales_90[(vendor, region)])
            qty_30 = max(0, sales_30[(vendor, region)])
            avg_90 = qty_90 / SALES_LOOKBACK_DAYS
            avg_30 = qty_30 / RECENT_SALES_DAYS
            forecast_daily = avg_30 * 0.6 + avg_90 * 0.4
            stock = stock_index[(nm_id, region)]
            inbound_qty = inbound["by_vendor_region"][(vendor, region)]
            if global_stock == 0 and qty_30 == 0 and qty_90 > 0:
                possible_stockouts += 1
                forecast_daily = max(forecast_daily, avg_90)
            if forecast_daily <= 0:
                continue
            raw_need = max(0.0, forecast_daily * target_days - stock - inbound_qty)
            rounded_need = _round_up_units(raw_need, batch_multiple)
            if rounded_need <= 0:
                continue
            coverage_before = (stock + inbound_qty) / forecast_daily
            rows.append(
                NeedRow(
                    vendor_code=vendor,
                    nm_id=nm_id,
                    barcode=barcode,
                    title=title,
                    internal_sku=_text(catalog_row.get("internal_sku")) or vendor,
                    region=region,
                    product_type=product_type,
                    classification_confirmed=confirmed_type,
                    batch_multiple=batch_multiple,
                    pack_qty=pack_qty,
                    sales_90=qty_90,
                    sales_30=qty_30,
                    avg_daily_90=avg_90,
                    avg_daily_30=avg_30,
                    forecast_daily=forecast_daily,
                    current_stock=stock,
                    confirmed_inbound=inbound_qty,
                    coverage_before=coverage_before,
                    target_days=target_days,
                    raw_need=raw_need,
                    rounded_need=rounded_need,
                )
            )

    cluster_priorities = _region_priorities(rows)
    selected_regions = [row["region"] for row in cluster_priorities[:cluster_count]]
    selected_set = set(selected_regions)
    rank_by_region = {region: rank for rank, region in enumerate(selected_regions, start=1)}
    rows = [row for row in rows if row.region in selected_set]
    for row in rows:
        row.cluster_rank = rank_by_region[row.region]

    if normalized_mode == "coverage_days":
        for row in rows:
            row.quantity = row.rounded_need
        unused_capacity = 0
    else:
        unused_capacity = _allocate_capacity(rows, value)

    plan_rows = [row for row in rows if row.quantity > 0]
    if len(selected_regions) < cluster_count:
        warnings.append(
            f"Запрошено кластеров: {cluster_count}; с подтвержденной положительной потребностью найдено: "
            f"{len(selected_regions)}."
        )
    if possible_stockouts:
        warnings.append(
            f"Для {possible_stockouts} строк возможен период отсутствия товара: история ежедневных остатков недоступна, "
            "поэтому нулевые недавние продажи не были приняты за отсутствие спроса."
        )
    if unknown_types:
        warnings.append(f"Для {unknown_types} товаров тип не определен; применена согласованная кратность 8.")
    if not plan_rows:
        warnings.append("После учета остатков, подтвержденных поставок и ограничений производственный план пуст.")
    if normalized_mode == "capacity" and unused_capacity:
        warnings.append(
            f"Не распределено {unused_capacity} физических изделий: следующий допустимый производственный блок "
            "не помещается либо подтвержденная потребность исчерпана."
        )

    xlsx_path = run_dir / f"v_rabotu_wb_{normalized_mode}_{value}_{len(selected_regions)}cl_{today.isoformat()}.xlsx"
    markdown_path = run_dir / "wb_production_work_plan.md"
    rows_csv = processed_dir / "wb_production_rows.csv"
    controls_csv = processed_dir / "control_rows.csv"
    _write_workbook(
        xlsx_path,
        plan_rows,
        controls,
        cluster_priorities,
        normalized_mode,
        value,
        cluster_count,
        period_from,
        period_to,
    )
    _write_csv(rows_csv, [_row_dict(row) for row in plan_rows])
    _write_csv(controls_csv, controls)
    plan_checksum = _file_sha256(xlsx_path)
    physical_total = sum(row.physical_quantity for row in plan_rows)
    units_total = sum(row.quantity for row in plan_rows)
    summary: dict[str, Any] = {
        "task": "wb-production-work-plan",
        "run_id": run_id,
        "started_at": started_at.isoformat(timespec="seconds"),
        "overall_status": "warning" if warnings or controls else "ok",
        "mode": "read_only",
        "calculation_mode": normalized_mode,
        "input_value": value,
        "requested_cluster_count": cluster_count,
        "target_days": target_days,
        "sales_period": {"from": period_from.isoformat(), "to": period_to.isoformat(), "days": SALES_LOOKBACK_DAYS},
        "metrics": {
            "marketplace_units": units_total,
            "physical_pieces": physical_total,
            "articles": len({row.vendor_code for row in plan_rows}),
            "regions": len({row.region for row in plan_rows}),
            "selected_clusters": len(selected_regions),
            "positive_need_clusters": len(cluster_priorities),
            "plan_rows": len(plan_rows),
            "control_rows": len(controls),
            "unknown_type_rows": unknown_types,
            "possible_stockout_rows": possible_stockouts,
            "requested_capacity_physical": value if normalized_mode == "capacity" else 0,
            "unused_capacity_physical": unused_capacity,
        },
        "region_totals": _region_totals(plan_rows, cluster_priorities),
        "cluster_priorities": cluster_priorities,
        "warnings": warnings,
        "approval_status": "pending_owner_review",
        "plan_checksum": plan_checksum,
        "sources": [
            "WB Content API /content/v2/get/cards/list",
            "WB Analytics API /api/analytics/v1/stocks-report/wb-warehouses",
            "WB Statistics API /api/v1/supplier/sales?flag=0",
            "WB Supplies API /api/v1/supplies and /goods",
            "data/catalog/unified/products.csv",
        ],
        "artifacts": {
            "report": str(xlsx_path),
            "markdown": str(markdown_path),
            "rows_csv": str(rows_csv),
            "controls_csv": str(controls_csv),
            "summary": str(run_dir / "summary.json"),
        },
        "next_step": "Проверить Excel и отдельно утвердить или отклонить план в боте. Поставка в WB не создавалась.",
    }
    _write_report(markdown_path, summary)
    write_json(run_dir / "summary.json", summary)
    manifest = write_summary_run_manifest(
        data_dir=data_dir,
        run_dir=run_dir,
        summary=summary,
        task="wb-production-work-plan",
        mode="read_only",
        risk="low",
        marketplaces=["wb"],
        inputs={
            "mode": normalized_mode,
            "value": value,
            "cluster_count": cluster_count,
            "target_days": target_days,
        },
        pending_id=run_id,
        lifecycle_status="pending_review",
        closed=False,
    )
    summary["artifacts"].update(manifest)
    write_json(run_dir / "summary.json", summary)
    return summary


def record_wb_work_plan_decision(*, data_dir: Path, plan_run_id: str, approved: bool) -> dict[str, Any]:
    plan_run_id = str(plan_run_id or "").strip()
    if not RUN_ID_RE.fullmatch(plan_run_id):
        raise ValueError("invalid plan_run_id")
    candidates = list((data_dir / "runs").glob(f"*/{plan_run_id}/summary.json"))
    if len(candidates) != 1:
        raise ValueError("WB work plan not found or ambiguous")
    summary_path = candidates[0]
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    if not isinstance(summary, dict) or summary.get("task") != "wb-production-work-plan":
        raise ValueError("run is not a WB production work plan")
    report_path = Path(str((summary.get("artifacts") or {}).get("report") or ""))
    if not report_path.exists() or _file_sha256(report_path) != summary.get("plan_checksum"):
        raise ValueError("plan file checksum mismatch")
    status = "owner_approved" if approved else "owner_rejected"
    decision = {
        "plan_run_id": plan_run_id,
        "status": status,
        "decided_at": datetime.now(MOSCOW).isoformat(timespec="seconds"),
        "plan_checksum": summary.get("plan_checksum"),
        "marketplace_write_performed": False,
    }
    decision_path = summary_path.parent / "owner_decision.json"
    write_json(decision_path, decision)
    summary["approval_status"] = status
    summary["artifacts"]["owner_decision"] = str(decision_path)
    summary["next_step"] = (
        "План утвержден в работу; производство и поставка исполняются вне этого read-only контура."
        if approved
        else "План отклонен; производство и поставка по нему не запускаются."
    )
    manifest = write_summary_run_manifest(
        data_dir=data_dir,
        run_dir=summary_path.parent,
        summary=summary,
        task="wb-production-work-plan",
        mode="read_only",
        risk="low",
        marketplaces=["wb"],
        inputs={
            "mode": summary.get("calculation_mode"),
            "value": summary.get("input_value"),
            "cluster_count": summary.get("requested_cluster_count"),
        },
        pending_id=plan_run_id,
        approved_id=plan_run_id if approved else None,
        lifecycle_status="approved" if approved else "closed",
        closed=True,
    )
    summary["artifacts"].update(manifest)
    write_json(summary_path, summary)
    return decision


def _card_index(cards: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for card in cards:
        vendor = _text(card.get("vendorCode"))
        if not vendor:
            continue
        barcode = ""
        for size in card.get("sizes") or []:
            if not isinstance(size, dict):
                continue
            skus = size.get("skus") or []
            if skus:
                barcode = _text(skus[0])
                break
        result[vendor] = {
            "nm_id": _text(card.get("nmID") or card.get("nmId")),
            "title": _text(card.get("title")),
            "barcode": barcode,
        }
    return result


def _stock_index(
    rows: list[dict[str, Any]],
) -> tuple[Counter[tuple[str, str]], set[tuple[str, str]], list[str]]:
    stocks: Counter[tuple[str, str]] = Counter()
    warehouse_totals: dict[str, Counter[str]] = defaultdict(Counter)
    row_keys: dict[str, set[tuple[str, str]]] = defaultdict(set)
    for row in rows:
        nm_id = _text(row.get("nmId") or row.get("nm_id"))
        region = _canonical_stock_region(_text(row.get("regionName") or row.get("region_name")))
        warehouse = _normalize(_text(row.get("warehouseName") or row.get("warehouse_name")))
        if not nm_id or not region:
            continue
        quantity = _int(row.get("quantity"))
        from_client = _int(row.get("inWayFromClient") or row.get("in_way_from_client"))
        stocks[(nm_id, region)] += quantity
        warehouse_totals[warehouse]["quantity"] += quantity
        warehouse_totals[warehouse]["from_client"] += from_client
        if from_client > 0:
            row_keys[warehouse].add((nm_id, region))
    anomalies: set[tuple[str, str]] = set()
    warnings: list[str] = []
    for warehouse, totals in warehouse_totals.items():
        quantity = totals["quantity"]
        from_client = totals["from_client"]
        if from_client >= 100 and from_client >= max(1, quantity * 2):
            anomalies.update(row_keys[warehouse])
            warnings.append(
                f"Склад {warehouse or 'не указан'}: inWayFromClient={from_client}, quantity={quantity}; "
                "затронутые товарные строки исключены из автоплана до стабилизации состояний WB."
            )
    return stocks, anomalies, warnings


def _sales_indexes(
    rows: list[dict[str, Any]],
    *,
    period_from: date,
    period_to: date,
    current_vendor_by_nm: dict[str, str],
) -> tuple[Counter[tuple[str, str]], Counter[tuple[str, str]]]:
    sales_90: Counter[tuple[str, str]] = Counter()
    sales_30: Counter[tuple[str, str]] = Counter()
    recent_from = period_to - timedelta(days=RECENT_SALES_DAYS - 1)
    for row in rows:
        day = _date(row.get("date"))
        if day is None or day < period_from or day > period_to:
            continue
        source_vendor = _text(row.get("supplierArticle"))
        nm_id = _text(row.get("nmId") or row.get("nmID"))
        vendor = current_vendor_by_nm.get(nm_id, source_vendor)
        region = _buyer_region(row)
        if not vendor or not region:
            continue
        direction = _sale_direction(row)
        if direction == 0:
            continue
        sales_90[(vendor, region)] += direction
        if day >= recent_from:
            sales_30[(vendor, region)] += direction
    return sales_90, sales_30


def _fetch_inbound(adapter: WbFbwSuppliesAdapter, stocks: list[dict[str, Any]]) -> dict[str, Any]:
    warehouse_regions: dict[str, str] = {}
    for row in stocks:
        warehouse = _normalize(_text(row.get("warehouseName")))
        region = _canonical_stock_region(_text(row.get("regionName")))
        if warehouse and region:
            warehouse_regions[warehouse] = region
    by_vendor_region: Counter[tuple[str, str]] = Counter()
    controls: list[dict[str, Any]] = []
    raw: dict[str, Any] = {"supplies": [], "details": [], "goods": []}
    supplies = adapter.fetch_supplies()
    raw["supplies"] = supplies
    for supply in supplies:
        status_id = _int(supply.get("statusID") or supply.get("statusId"))
        if status_id not in ACTIVE_INBOUND_STATUS_IDS:
            continue
        supply_id = supply.get("supplyID") or supply.get("supplyId")
        if not supply_id:
            continue
        detail = adapter.fetch_supply(supply_id)
        goods = adapter.fetch_supply_goods(supply_id)
        raw["details"].append({"supply_id": str(supply_id), "detail": detail})
        raw["goods"].append({"supply_id": str(supply_id), "rows": goods})
        warehouse = _text(
            detail.get("actualWarehouseName")
            or detail.get("warehouseName")
            or supply.get("actualWarehouseName")
            or supply.get("warehouseName")
        )
        region = _warehouse_region(warehouse, warehouse_regions)
        for good in goods:
            vendor = _text(good.get("vendorCode") or good.get("supplierArticle"))
            nm_id = _text(good.get("nmID") or good.get("nmId"))
            quantity = _int(good.get("quantity"))
            ready = _int(good.get("readyForSaleQuantity"))
            inbound_quantity = max(0, quantity - ready)
            if not vendor or inbound_quantity <= 0:
                continue
            if not region:
                controls.append(
                    _control(
                        vendor,
                        nm_id,
                        "unknown_inbound_destination_region",
                        f"Поставка {supply_id}, склад {warehouse or 'не указан'}: регион назначения не сопоставлен.",
                    )
                )
                continue
            by_vendor_region[(vendor, region)] += inbound_quantity
    return {"by_vendor_region": by_vendor_region, "controls": controls, "raw": raw}


def _allocate_capacity(rows: list[NeedRow], capacity_physical: int) -> int:
    remaining = capacity_physical
    eligible = [row for row in rows if row.rounded_need > 0]
    while eligible:
        fitting = [
            row
            for row in eligible
            if row.quantity < row.rounded_need and row.batch_multiple * row.pack_qty <= remaining
        ]
        if not fitting:
            break
        selected = max(
            fitting,
            key=lambda row: (
                (row.rounded_need - row.quantity) / max(1, row.rounded_need),
                max(0.0, row.target_days - row.coverage_after),
                row.forecast_daily,
                row.raw_need,
                row.vendor_code,
            ),
        )
        selected.quantity += selected.batch_multiple
        remaining -= selected.batch_multiple * selected.pack_qty
        eligible = [row for row in eligible if row.quantity < row.rounded_need]
    return remaining


def _region_priorities(rows: list[NeedRow]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "raw_need_marketplace_units": 0.0,
            "raw_need_physical_pieces": 0.0,
            "sales_90": 0,
            "sales_30": 0,
            "current_stock": 0,
            "confirmed_inbound": 0,
            "articles": set(),
        }
    )
    for row in rows:
        target = grouped[row.region]
        target["raw_need_marketplace_units"] += row.raw_need
        target["raw_need_physical_pieces"] += row.raw_need * row.pack_qty
        target["sales_90"] += row.sales_90
        target["sales_30"] += row.sales_30
        target["current_stock"] += row.current_stock
        target["confirmed_inbound"] += row.confirmed_inbound
        target["articles"].add(row.vendor_code)
    ordered = sorted(
        grouped.items(),
        key=lambda item: (-item[1]["raw_need_physical_pieces"], -item[1]["sales_30"], item[0]),
    )
    return [
        {
            "priority": rank,
            "region": region,
            "raw_need_marketplace_units": round(values["raw_need_marketplace_units"], 2),
            "raw_need_physical_pieces": round(values["raw_need_physical_pieces"], 2),
            "sales_90": values["sales_90"],
            "sales_30": values["sales_30"],
            "current_stock_in_cluster": values["current_stock"],
            "confirmed_inbound_to_cluster": values["confirmed_inbound"],
            "articles": len(values["articles"]),
        }
        for rank, (region, values) in enumerate(ordered, start=1)
    ]


def _classify_product(*, title: str, internal_sku: str, vendor_code: str) -> tuple[str, int, bool]:
    text = _normalize(" ".join((title, internal_sku, vendor_code)))
    sku = _normalize(f"{internal_sku} {vendor_code}")
    if "комплект" in text and "позывн" in text or ("kit" in sku and "pz" in sku):
        return "Комплект позывных", 6, True
    if "на спину" in text or "наспин" in text or "back" in sku:
        return "Наспинный", 2, True
    if "нагруд" in text or re.search(r"(^|_)ng(_|$)", internal_sku.lower()):
        return "Нагрудный", 12, True
    if "на кепк" in text or re.search(r"(^|_)kp(_|$)", internal_sku.lower()):
        return "На кепку", 9, True
    if "нарукав" in text or re.search(r"(^|_)nr(_|$)", internal_sku.lower()):
        return "Нарукавный", 8, True
    return "Тип не определен", 8, False


def _load_catalog(path: Path, confirmed_mapping_path: Path | None = None) -> dict[str, dict[str, str]]:
    index: dict[str, dict[str, str]] = {}
    if path.exists():
        with path.open(encoding="utf-8", newline="") as handle:
            for row in csv.DictReader(handle):
                vendor = _text(row.get("wb_vendor_code"))
                nm_id = _text(row.get("wb_nm_id"))
                if vendor:
                    index[f"vendor:{vendor}"] = row
                if nm_id:
                    index[f"nm:{nm_id}"] = row
    if confirmed_mapping_path and confirmed_mapping_path.exists():
        with confirmed_mapping_path.open(encoding="utf-8", newline="") as handle:
            for source in csv.DictReader(handle):
                vendor = _text(source.get("wb_vendor_code"))
                nm_id = _text(source.get("wb_nm_id"))
                if not vendor and not nm_id:
                    continue
                internal_sku = _text(source.get("internal_sku"))
                title = _text(source.get("wb_title") or source.get("ozon_title"))
                fallback = {
                    "internal_sku": internal_sku,
                    "product_name": title,
                    "product_group": _product_group_from_sku(internal_sku),
                    "pack_qty": str(_pack_qty_from_identity(internal_sku, title)),
                    "wb_vendor_code": vendor,
                    "wb_nm_id": nm_id,
                    "wb_barcode": _text(source.get("wb_barcode")),
                    "mapping_status": "owner_confirmed_fallback",
                }
                if vendor:
                    index.setdefault(f"vendor:{vendor}", fallback)
                if nm_id:
                    index.setdefault(f"nm:{nm_id}", fallback)
    return index


def _product_group_from_sku(internal_sku: str) -> str:
    sku = internal_sku.lower()
    if "loop" in sku:
        return "loop"
    if "nash" in sku:
        return "nash"
    return "chev"


def _pack_qty_from_identity(internal_sku: str, title: str) -> int:
    text = _normalize(f"{internal_sku} {title}")
    if "петлиц" in text or "loop" in internal_sku.lower():
        return 1
    match = re.search(r"(?:kit|комплект)[ _-]*([2-9])", text)
    if match:
        return int(match.group(1))
    match = re.search(r"([2-9])\s*(?:шт|штук)", text)
    return int(match.group(1)) if match else 1


def _is_manufactured(title: str, vendor_code: str, product_group: str) -> bool:
    text = " ".join((title, vendor_code, product_group))
    if EXCLUDED_RE.search(text):
        return False
    return product_group.lower() in MANUFACTURED_GROUPS or bool(MANUFACTURED_RE.search(text))


def _buyer_region(row: dict[str, Any]) -> str:
    federal = _normalize(_text(row.get("oblastOkrugName")))
    if federal in BUYER_REGION_MAP:
        return BUYER_REGION_MAP[federal]
    return _canonical_stock_region(_text(row.get("regionName")))


def _canonical_stock_region(value: str) -> str:
    text = _normalize(value)
    aliases = {
        "центральный": "Центральный",
        "приволжский": "Приволжский",
        "северо-западный": "Северо-Западный",
        "уральский": "Уральский",
        "южный и северо-кавказский": "Южный и Северо-Кавказский",
        "дальневосточный и сибирский": "Дальневосточный и Сибирский",
    }
    return aliases.get(text, "")


def _warehouse_region(warehouse: str, known: dict[str, str]) -> str:
    normalized = _normalize(warehouse)
    if normalized in known:
        return known[normalized]
    for needle, region in WAREHOUSE_REGION_HINTS.items():
        if needle in normalized:
            return region
    return ""


def _sale_direction(row: dict[str, Any]) -> int:
    sale_id = _text(row.get("saleID") or row.get("saleId")).upper()
    if sale_id.startswith("R"):
        return -1
    if sale_id.startswith("S"):
        return 1
    if row.get("isRealization") is True:
        return 1
    return 0


def _round_up_units(value: float, multiple: int) -> int:
    if value <= 0:
        return 0
    return int(math.ceil(value / multiple) * multiple)


def _row_dict(row: NeedRow) -> dict[str, Any]:
    return {
        "cluster_priority": row.cluster_rank,
        "region": row.region,
        "internal_sku": row.internal_sku,
        "vendor_code": row.vendor_code,
        "nm_id": row.nm_id,
        "barcode": row.barcode,
        "title": row.title,
        "product_type": row.product_type,
        "classification_confirmed": row.classification_confirmed,
        "batch_multiple": row.batch_multiple,
        "pack_qty": row.pack_qty,
        "quantity": row.quantity,
        "physical_quantity": row.physical_quantity,
        "sales_90": row.sales_90,
        "sales_30": row.sales_30,
        "avg_daily_90": round(row.avg_daily_90, 4),
        "avg_daily_30": round(row.avg_daily_30, 4),
        "forecast_daily": round(row.forecast_daily, 4),
        "current_stock": row.current_stock,
        "confirmed_inbound": row.confirmed_inbound,
        "coverage_before": round(row.coverage_before, 1),
        "target_days": row.target_days,
        "raw_need": round(row.raw_need, 2),
        "rounded_need": row.rounded_need,
        "coverage_after": round(row.coverage_after, 1),
    }


def _region_totals(rows: list[NeedRow], priorities: list[dict[str, Any]]) -> list[dict[str, Any]]:
    priority_by_region = {row["region"]: row for row in priorities}
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: {"marketplace_units": 0, "physical_pieces": 0, "articles": set()})
    for row in rows:
        grouped[row.region]["marketplace_units"] += row.quantity
        grouped[row.region]["physical_pieces"] += row.physical_quantity
        grouped[row.region]["articles"].add(row.vendor_code)
    return [
        {
            "priority": priority_by_region.get(region, {}).get("priority", 0),
            "region": region,
            "marketplace_units": values["marketplace_units"],
            "physical_pieces": values["physical_pieces"],
            "articles": len(values["articles"]),
        }
        for region, values in sorted(
            grouped.items(),
            key=lambda item: (priority_by_region.get(item[0], {}).get("priority", 999), item[0]),
        )
    ]


def _write_workbook(
    path: Path,
    rows: list[NeedRow],
    controls: list[dict[str, Any]],
    priorities: list[dict[str, Any]],
    mode: str,
    value: int,
    cluster_count: int,
    period_from: date,
    period_to: date,
) -> None:
    workbook = Workbook()
    articles = workbook.active
    articles.title = "Артикулы"
    grouped: dict[str, dict[str, Any]] = {}
    for row in rows:
        target = grouped.setdefault(
            row.vendor_code,
            {
                "Внутренний артикул": row.internal_sku,
                "Артикул продавца WB": row.vendor_code,
                "nmID": row.nm_id,
                "Баркод WB": row.barcode,
                "Название": row.title,
                "Тип изделия": row.product_type,
                "Кратность": row.batch_multiple,
                "Изделий в товаре": row.pack_qty,
                "Товарных единиц WB": 0,
                "Физических изделий": 0,
                "Кластеров назначения": set(),
                "Классификация подтверждена": "да" if row.classification_confirmed else "нет; кратность 8",
            },
        )
        target["Товарных единиц WB"] += row.quantity
        target["Физических изделий"] += row.physical_quantity
        target["Кластеров назначения"].add(row.region)
    article_rows: list[dict[str, Any]] = []
    for value_row in grouped.values():
        value_row["Кластеров назначения"] = ", ".join(sorted(value_row["Кластеров назначения"]))
        article_rows.append(value_row)
    _append_dict_rows(articles, sorted(article_rows, key=lambda row: (-_int(row["Физических изделий"]), str(row["Артикул продавца WB"]))))

    regions = workbook.create_sheet("ВБ регионы")
    _append_dict_rows(
        regions,
        [_row_dict(row) for row in sorted(rows, key=lambda item: (item.cluster_rank, -item.physical_quantity, item.vendor_code))],
    )

    control = workbook.create_sheet("Контроль")
    parameters = [
        {"Параметр": "Режим", "Значение": "По производственной возможности" if mode == "capacity" else "По периоду покрытия"},
        {"Параметр": "Введенное значение", "Значение": value},
        {"Параметр": "Запрошено кластеров", "Значение": cluster_count},
        {"Параметр": "Период продаж", "Значение": f"{period_from.isoformat()} - {period_to.isoformat()}"},
        {
            "Параметр": "Формула",
            "Значение": "спрос товара в кластере - остаток товара в этом кластере - confirmed inbound в этот кластер",
        },
        {"Параметр": "Статус", "Значение": "Ожидает проверки владельца; поставка WB не создавалась"},
    ]
    _append_dict_rows(control, parameters)
    control.append([])
    _append_dict_rows(control, priorities, start_row=control.max_row + 1)
    if controls:
        control.append([])
        _append_dict_rows(control, controls, start_row=control.max_row + 1)
    workbook.save(path)


def _append_dict_rows(sheet: Any, rows: list[dict[str, Any]], *, start_row: int = 1) -> None:
    fields: list[str] = []
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    if not fields:
        fields = ["Статус"]
        rows = [{"Статус": "Нет строк"}]
    for column, field in enumerate(fields, start=1):
        sheet.cell(start_row, column, field)
    for row_number, row in enumerate(rows, start=start_row + 1):
        for column, field in enumerate(fields, start=1):
            sheet.cell(row_number, column, row.get(field, ""))
    _sheet_setup(sheet, header_row=start_row)


def _sheet_setup(sheet: Any, *, header_row: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[header_row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        width = max(10, max((len(str(cell.value or "")) for cell in sheet[letter]), default=10))
        sheet.column_dimensions[letter].width = min(55, width + 2)


def _write_report(path: Path, summary: dict[str, Any]) -> None:
    metrics = summary["metrics"]
    mode_label = "по производственной возможности" if summary["calculation_mode"] == "capacity" else "по периоду покрытия"
    lines = [
        "# WB: файл в работу",
        "",
        f"- Режим: `{mode_label}`.",
        f"- Введенное значение: `{summary['input_value']}`.",
        f"- Запрошено кластеров: `{summary['requested_cluster_count']}`; выбрано: `{metrics['selected_clusters']}`.",
        f"- Целевое покрытие: `{summary['target_days']}` дней.",
        f"- Период спроса: `{summary['sales_period']['from']} - {summary['sales_period']['to']}`.",
        f"- Товарных единиц WB: `{metrics['marketplace_units']}`.",
        f"- Физических изделий: `{metrics['physical_pieces']}`.",
        f"- Артикулов: `{metrics['articles']}`; регионов назначения: `{metrics['regions']}`.",
        f"- Контрольных строк: `{metrics['control_rows']}`.",
        "",
        "Потребность рассчитана отдельно для каждой пары товар × WB-кластер. Общий остаток магазина не вычитался.",
        "",
        "## Распределение",
        "",
    ]
    for row in summary["region_totals"]:
        lines.append(
            f"- #{row['priority']} {row['region']}: `{row['marketplace_units']}` ед. WB / `{row['physical_pieces']}` физических изделий / `{row['articles']}` артикулов."
        )
    if summary["warnings"]:
        lines.extend(["", "## Ограничения", ""])
        lines.extend(f"- {warning}" for warning in summary["warnings"])
    lines.extend(
        [
            "",
            "## Следующий шаг",
            "",
            "Проверить Excel и отдельно утвердить или отклонить план. Производство и поставки WB автоматически не создавались.",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields: list[str] = []
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    if not fields:
        fields = ["status"]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _control(vendor_code: str, nm_id: str, code: str, message: str) -> dict[str, Any]:
    return {"vendor_code": vendor_code, "nm_id": nm_id, "code": code, "message": message}


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _date(value: Any) -> date | None:
    text = _text(value)[:10]
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _normalize(value: Any) -> str:
    return " ".join(_text(value).lower().replace("ё", "е").split())


def _text(value: Any) -> str:
    return str(value or "").strip()


def _int(value: Any, default: int = 0) -> int:
    try:
        if value in (None, ""):
            return default
        return int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError):
        return default
