from __future__ import annotations

import csv
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_CEILING
from pathlib import Path
from statistics import median
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font

from seller_agent.config import AppCredentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.ozon.adapter import OzonSellerAdapter
from seller_agent.reports.writer import ensure_dir, write_json
from seller_agent.tasks.marketplace_period_report import collect_ozon_period_metrics
from seller_agent.tasks.pricing_status import normalize_ozon_price_item


MOSCOW = ZoneInfo("Europe/Moscow")
PROMO_RESERVE = Decimal("0.20")
VISIBLE_DISCOUNT = Decimal("0.50")
PRICE_STEP = Decimal("10")


def run_ozon_pricing_margin(
    *,
    credentials: AppCredentials,
    data_dir: Path,
    unit_cost: Any,
    target_margin: Any,
    period_days: int,
    run_id: str | None = None,
    today: date | None = None,
    adapter: OzonSellerAdapter | None = None,
) -> dict[str, Any]:
    cost = _positive_decimal(unit_cost, "unit_cost")
    margin = _non_negative_decimal(target_margin, "target_margin")
    if period_days not in {15, 30}:
        raise ValueError("period_days must be 15 or 30")
    if credentials.ozon_seller is None and adapter is None:
        raise ValueError("Ozon Seller API credentials are not configured")

    current_day = today or datetime.now(MOSCOW).date()
    end = current_day - timedelta(days=1)
    start = end - timedelta(days=period_days - 1)
    started_at = datetime.now(MOSCOW)
    run_id = run_id or f"ozon_pricing_margin_{started_at.strftime('%Y%m%dT%H%M%S')}"
    run_dir = ensure_dir(data_dir / "runs" / started_at.date().isoformat() / run_id)

    catalog_rows = _load_catalog(data_dir / "catalog" / "unified" / "products.csv")
    if not catalog_rows:
        raise ValueError("Unified catalog is empty or missing")
    ozon = adapter or OzonSellerAdapter(credentials.ozon_seller)
    period_metrics = collect_ozon_period_metrics(
        adapter=ozon,
        catalog_rows=catalog_rows,
        start=start,
        end=end,
    )
    period_comparison: dict[str, Any] = {}
    for days in (15, 30):
        comparison_start = end - timedelta(days=days - 1)
        metrics_for_days = period_metrics if days == period_days else collect_ozon_period_metrics(
            adapter=ozon,
            catalog_rows=catalog_rows,
            start=comparison_start,
            end=end,
        )
        period_comparison[str(days)] = _period_snapshot(
            metrics_for_days,
            date_from=comparison_start,
            date_to=end,
        )
    price_items = [normalize_ozon_price_item(row) for row in ozon.fetch_product_info_prices()]
    stock_by_key: dict[str, int] = {}
    stock_confirmed = False
    try:
        stock_by_key = _ozon_stock_by_key(ozon.fetch_stock_on_warehouses())
        stock_confirmed = True
    except Exception as exc:  # noqa: BLE001 - stock gate is reported explicitly.
        warnings_from_stock = f"Ozon stock gate не подтвержден: {exc.__class__.__name__}"

    summary_metrics = period_metrics.get("summary") or {}
    buyout_units = int(summary_metrics.get("buyout_units") or 0)
    physical_pieces = int(summary_metrics.get("physical_pieces") or 0)
    gross = _decimal(summary_metrics.get("gross"))
    total_expenses = _decimal(summary_metrics.get("expenses"))
    expense_rows = _expense_rows(
        period_metrics.get("expenses") or {},
        buyout_units=buyout_units,
        physical_pieces=physical_pieces,
    )
    logistics = sum(
        _decimal(row["total"])
        for row in expense_rows
        if row["category"] == "Логистика"
    )

    warnings = list(period_metrics.get("warnings") or [])
    if not stock_confirmed:
        warnings.append(warnings_from_stock)
    if buyout_units <= 0 or physical_pieces <= 0 or gross <= 0:
        warnings.append("Недостаточно завершённых продаж Ozon для надёжного расчёта цен.")
    reconciliation = next(
        (_decimal(row["total"]) for row in expense_rows if row["category"] == "Сверочная разница"),
        Decimal("0"),
    )
    if abs(reconciliation) >= Decimal("0.01"):
        warnings.append(
            "В расходах есть сверочная разница; она включена в итог, но не распределена как отдельный тариф."
        )

    current_commissions = sorted(
        value
        for value in (_commission_fraction(row.get("sales_percent_fbo")) for row in price_items)
        if value is not None
    )
    current_commission_max = max(current_commissions, default=Decimal("0"))
    current_commission_median = (
        Decimal(str(median([float(value) for value in current_commissions])))
        if current_commissions
        else Decimal("0")
    )
    historical_variable_rate = (
        max(Decimal("0"), total_expenses - logistics) / gross if gross > 0 else Decimal("0")
    )
    model_variable_rate = max(historical_variable_rate, current_commission_max)
    retained_share = Decimal("1") - model_variable_rate
    fixed_logistics_per_product = logistics / buyout_units if buyout_units else Decimal("0")
    expense_per_product = total_expenses / buyout_units if buyout_units else Decimal("0")
    expense_per_physical = total_expenses / physical_pieces if physical_pieces else Decimal("0")

    calculation_ready = (
        buyout_units > 0
        and physical_pieces > 0
        and gross > 0
        and Decimal("0") < retained_share < Decimal("1")
        and bool(current_commissions)
    )
    if not current_commissions:
        warnings.append("Ozon не вернул текущую комиссию FBO; ценовая сетка заблокирована.")
    if retained_share <= 0:
        warnings.append("Расчётная доля удержаний достигла 100%; ценовая сетка заблокирована.")

    observed_pack_qty = sorted({_pack_qty(row) for row in catalog_rows if _is_active_ozon(row)})
    ladder_pack_qty = sorted(set(range(1, 6)).union(observed_pack_qty))
    price_ladder = (
        [
            _price_ladder_row(
                pack_qty=pack_qty,
                unit_cost=cost,
                target_margin=margin,
                fixed_logistics=fixed_logistics_per_product,
                variable_rate=model_variable_rate,
            )
            for pack_qty in ladder_pack_qty
        ]
        if calculation_ready
        else []
    )
    ladder_by_pack = {int(row["pack_qty"]): row for row in price_ladder}
    products = _product_rows(
        catalog_rows,
        price_items,
        ladder_by_pack,
        stock_by_key=stock_by_key,
        stock_confirmed=stock_confirmed,
    )
    products_without_prices = sum(1 for row in products if not row["price_data_found"])
    if products_without_prices:
        warnings.append(
            f"Для {products_without_prices} активных товаров unified catalog не найдены текущие цены Ozon."
        )

    metrics = {
        "period_days": period_days,
        "date_from": start.isoformat(),
        "date_to": end.isoformat(),
        "catalog_products": len([row for row in catalog_rows if _is_active_ozon(row)]),
        "price_products": len(price_items),
        "buyout_units": buyout_units,
        "physical_pieces": physical_pieces,
        "gross": _money(gross),
        "total_expenses": _money(total_expenses),
        "net": _money(summary_metrics.get("net")),
        "expense_per_sold_product": _money(expense_per_product),
        "expense_per_physical_item": _money(expense_per_physical),
        "fixed_logistics_per_sold_product": _money(fixed_logistics_per_product),
        "historical_variable_rate_pct": _percent(historical_variable_rate),
        "current_commission_median_pct": _percent(current_commission_median),
        "current_commission_max_pct": _percent(current_commission_max),
        "model_variable_rate_pct": _percent(model_variable_rate),
        "retained_share_pct": _percent(retained_share),
        "calculation_ready": calculation_ready,
        "stock_gate_confirmed": stock_confirmed,
        "products_with_stock": sum(1 for row in products if row["stock_gate"] == "pass"),
        "products_without_stock": sum(1 for row in products if row["stock_gate"] == "blocked_no_stock"),
        "products_in_actions": sum(1 for row in products if row["in_action"] is True),
        "action_membership_unconfirmed": sum(1 for row in products if row["in_action"] is None),
    }
    overall_status = "warning" if warnings else "ok"

    report_path = run_dir / "ozon_pricing_margin.xlsx"
    markdown_path = run_dir / "ozon_pricing_margin.md"
    json_path = run_dir / "ozon_pricing_margin.json"
    _write_markdown(
        markdown_path,
        metrics=metrics,
        unit_cost=cost,
        target_margin=margin,
        expense_rows=expense_rows,
        price_ladder=price_ladder,
        warnings=warnings,
    )
    _write_xlsx(
        report_path,
        metrics=metrics,
        unit_cost=cost,
        target_margin=margin,
        expense_rows=expense_rows,
        price_ladder=price_ladder,
        products=products,
        warnings=warnings,
    )
    write_json(
        json_path,
        {
            "metrics": metrics,
            "inputs": {"unit_cost": _money(cost), "target_margin": _money(margin)},
            "expenses": expense_rows,
            "price_ladder": price_ladder,
            "products": products,
            "period_comparison": period_comparison,
            "warnings": warnings,
            "sources": period_metrics.get("sources") or [],
        },
    )

    summary = {
        "run_id": run_id,
        "started_at": started_at.isoformat(timespec="seconds"),
        "overall_status": overall_status,
        "marketplace": "ozon",
        "mode": "dry_run",
        "inputs": {
            "unit_cost": _money(cost),
            "target_margin": _money(margin),
            "period_days": period_days,
        },
        "metrics": metrics,
        "expenses": expense_rows,
        "price_ladder": price_ladder,
        "period_comparison": period_comparison,
        "products": products,
        "warnings": warnings,
        "artifacts": {
            "report": str(report_path),
            "markdown": str(markdown_path),
            "json": str(json_path),
        },
    }
    write_json(run_dir / "summary.json", summary)
    manifest_artifacts = write_summary_run_manifest(
        data_dir=data_dir,
        run_dir=run_dir,
        summary=summary,
        task="ozon-pricing-margin",
        mode="dry_run",
        risk="low",
        marketplaces=["ozon"],
        inputs={
            "unit_cost": _money(cost),
            "target_margin": _money(margin),
            "period_days": period_days,
            "date_from": start.isoformat(),
            "date_to": end.isoformat(),
        },
        lifecycle_status="pending_review",
        closed=False,
    )
    summary["artifacts"].update(manifest_artifacts)
    write_json(run_dir / "summary.json", summary)
    return summary


def _price_ladder_row(
    *,
    pack_qty: int,
    unit_cost: Decimal,
    target_margin: Decimal,
    fixed_logistics: Decimal,
    variable_rate: Decimal,
) -> dict[str, Any]:
    retained_share = Decimal("1") - variable_rate
    required_payout = (unit_cost + target_margin) * pack_qty
    minimum = _ceil_price((required_payout + fixed_logistics) / retained_share)
    discounted = _ceil_price(minimum / (Decimal("1") - PROMO_RESERVE))
    base = _ceil_price(discounted / (Decimal("1") - VISIBLE_DISCOUNT))
    expected_expenses = minimum * variable_rate + fixed_logistics
    expected_margin_per_item = (minimum - expected_expenses - unit_cost * pack_qty) / pack_qty
    return {
        "pack_qty": pack_qty,
        "cost_total": _money(unit_cost * pack_qty),
        "target_margin_total": _money(target_margin * pack_qty),
        "minimum_price": _money(minimum),
        "discounted_price": _money(discounted),
        "base_price": _money(base),
        "promo_floor_at_20_pct": _money(discounted * (Decimal("1") - PROMO_RESERVE)),
        "expected_expenses_at_minimum": _money(expected_expenses),
        "expected_margin_per_item": _money(expected_margin_per_item),
    }


def _expense_rows(values: dict[str, Any], *, buyout_units: int, physical_pieces: int) -> list[dict[str, Any]]:
    rows = []
    for category, raw_value in values.items():
        total = _decimal(raw_value)
        rows.append(
            {
                "category": str(category),
                "total": _money(total),
                "per_sold_product": _money(total / buyout_units) if buyout_units else None,
                "per_physical_item": _money(total / physical_pieces) if physical_pieces else None,
            }
        )
    return rows


def _product_rows(
    catalog_rows: list[dict[str, str]],
    price_items: list[dict[str, Any]],
    ladder_by_pack: dict[int, dict[str, Any]],
    *,
    stock_by_key: dict[str, int] | None = None,
    stock_confirmed: bool = False,
) -> list[dict[str, Any]]:
    stock_by_key = stock_by_key or {}
    price_by_key: dict[str, dict[str, Any]] = {}
    for item in price_items:
        for field in ("offer_id", "product_id", "sku"):
            key = str(item.get(field) or "").strip()
            if key:
                price_by_key[key] = item
    rows = []
    for row in catalog_rows:
        if not _is_active_ozon(row):
            continue
        item = next(
            (
                price_by_key[value]
                for value in (
                    str(row.get("ozon_offer_id") or "").strip(),
                    str(row.get("ozon_product_id") or "").strip(),
                    str(row.get("ozon_sku") or "").strip(),
                )
                if value in price_by_key
            ),
            {},
        )
        pack_qty = _pack_qty(row)
        target = ladder_by_pack.get(pack_qty) or {}
        stock_qty = max(
            [
                stock_by_key.get(str(value or "").strip(), 0)
                for value in (row.get("ozon_offer_id"), row.get("ozon_product_id"), row.get("ozon_sku"))
            ]
            or [0]
        )
        marketing_price = _number_or_none(item.get("marketing_seller_price"))
        current_price = _number_or_none(item.get("price"))
        action_confirmed = item.get("marketing_seller_price") not in (None, "")
        in_action = bool(marketing_price is not None and current_price is not None and marketing_price < current_price) if action_confirmed else None
        rows.append(
            {
                "internal_sku": row.get("internal_sku") or "",
                "product_name": row.get("product_name") or "",
                "pack_qty": pack_qty,
                "ozon_offer_id": row.get("ozon_offer_id") or "",
                "current_price": current_price,
                "current_old_price": _number_or_none(item.get("old_price")),
                "current_min_price": _number_or_none(item.get("min_price")),
                "current_marketing_seller_price": _number_or_none(item.get("marketing_seller_price")),
                "target_minimum_price": target.get("minimum_price"),
                "target_discounted_price": target.get("discounted_price"),
                "target_base_price": target.get("base_price"),
                "price_data_found": bool(item),
                "stock_qty": stock_qty if stock_confirmed else None,
                "stock_gate": "pass" if stock_confirmed and stock_qty > 0 else "blocked_no_stock" if stock_confirmed else "not_confirmed",
                "in_action": in_action,
                "action_check_status": "confirmed_from_marketing_seller_price" if action_confirmed else "not_confirmed",
            }
        )
    return rows


def _period_snapshot(metrics: dict[str, Any], *, date_from: date, date_to: date) -> dict[str, Any]:
    summary = metrics.get("summary") if isinstance(metrics.get("summary"), dict) else {}
    buyout_units = int(summary.get("buyout_units") or 0)
    physical = int(summary.get("physical_pieces") or 0)
    expenses = _decimal(summary.get("expenses"))
    return {
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "buyout_units": buyout_units,
        "physical_pieces": physical,
        "total_expenses": _money(expenses),
        "expense_per_sold_product": _money(expenses / buyout_units) if buyout_units else None,
        "expense_per_physical_item": _money(expenses / physical) if physical else None,
        "warnings": list(metrics.get("warnings") or []),
    }


def _ozon_stock_by_key(rows: list[dict[str, Any]]) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in rows:
        quantity = int(
            _decimal(
                row.get("free_to_sell_amount")
                or row.get("available_stock_count")
                or row.get("present")
                or row.get("quantity")
            )
        )
        for key in ("sku", "offer_id", "product_id", "item_code"):
            value = str(row.get(key) or "").strip()
            if value:
                result[value] = result.get(value, 0) + quantity
    return result


def _write_markdown(
    path: Path,
    *,
    metrics: dict[str, Any],
    unit_cost: Decimal,
    target_margin: Decimal,
    expense_rows: list[dict[str, Any]],
    price_ladder: list[dict[str, Any]],
    warnings: list[str],
) -> None:
    lines = [
        "# Ozon: цены и маржа",
        "",
        f"- Период расходов: `{metrics['date_from']} - {metrics['date_to']}` ({metrics['period_days']} дней)",
        f"- Себестоимость одного физического изделия: **{_rub(unit_cost)}**",
        f"- Целевая маржа одного физического изделия: **{_rub(target_margin)}**",
        "- Режим: `dry-run`; цены Ozon не изменялись.",
        "",
        "## Расходы Ozon",
        "",
        f"- Выкупы: **{metrics['buyout_units']} товаров / {metrics['physical_pieces']} физических изделий**.",
        f"- Всего расходов: **{_rub(metrics['total_expenses'])}**.",
        f"- На проданный товар/комплект: **{_rub(metrics['expense_per_sold_product'])}**.",
        f"- На одно физическое изделие: **{_rub(metrics['expense_per_physical_item'])}**.",
        f"- Средняя логистика на товар/комплект: **{_rub(metrics['fixed_logistics_per_sold_product'])}**.",
        f"- Историческая переменная доля расходов: **{metrics['historical_variable_rate_pct']}%**.",
        f"- Текущая комиссия FBO: медиана **{metrics['current_commission_median_pct']}%**, максимум **{metrics['current_commission_max_pct']}%**.",
        f"- Для расчёта взята консервативная переменная доля: **{metrics['model_variable_rate_pct']}%**.",
        "",
        "| Статья | Всего | На товар/комплект | На физическое изделие |",
        "| --- | ---: | ---: | ---: |",
    ]
    for row in expense_rows:
        lines.append(
            f"| {row['category']} | {_rub(row['total'])} | {_rub_or_na(row['per_sold_product'])} | {_rub_or_na(row['per_physical_item'])} |"
        )
    lines.extend(["", "## Расчётная ценовая сетка", ""])
    if price_ladder:
        lines.extend(
            [
                "| Изделий | Минимальная | Со скидкой | Базовая | Маржа/изделие |",
                "| ---: | ---: | ---: | ---: | ---: |",
            ]
        )
        for row in price_ladder:
            lines.append(
                f"| {row['pack_qty']} | {_rub(row['minimum_price'])} | {_rub(row['discounted_price'])} | {_rub(row['base_price'])} | {_rub(row['expected_margin_per_item'])} |"
            )
    else:
        lines.append("Расчёт заблокирован из-за недостаточных или неполных данных.")
    if warnings:
        lines.extend(["", "## Ограничения", ""])
        lines.extend(f"- {warning}" for warning in warnings)
    lines.extend(
        [
            "",
            "## Методика",
            "",
            "- Налог не включён.",
            "- Логистика учитывается как фиксированный средний расход на проданный товар/комплект.",
            "- Остальные подтверждённые расходы входят в переменную долю и повторно не добавляются.",
            "- Цена со скидкой оставляет 20% резерв до минимальной цены; базовая цена вдвое выше цены со скидкой.",
            "",
            "Изменений в Ozon не выполнялось.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_xlsx(
    path: Path,
    *,
    metrics: dict[str, Any],
    unit_cost: Decimal,
    target_margin: Decimal,
    expense_rows: list[dict[str, Any]],
    price_ladder: list[dict[str, Any]],
    products: list[dict[str, Any]],
    warnings: list[str],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    summary.append(["Ozon: цены и маржа", "read-only"])
    rows = [
        ("Период", f"{metrics['date_from']} - {metrics['date_to']}"),
        ("Дней", metrics["period_days"]),
        ("Себестоимость изделия, руб.", float(unit_cost)),
        ("Маржа изделия, руб.", float(target_margin)),
        ("Выкупы, товаров", metrics["buyout_units"]),
        ("Выкупы, физических изделий", metrics["physical_pieces"]),
        ("Расходы всего, руб.", metrics["total_expenses"]),
        ("Расходы на товар/комплект, руб.", metrics["expense_per_sold_product"]),
        ("Расходы на физическое изделие, руб.", metrics["expense_per_physical_item"]),
        ("Логистика на товар/комплект, руб.", metrics["fixed_logistics_per_sold_product"]),
        ("Историческая переменная доля, %", metrics["historical_variable_rate_pct"]),
        ("Текущая комиссия FBO, медиана, %", metrics["current_commission_median_pct"]),
        ("Текущая комиссия FBO, максимум, %", metrics["current_commission_max_pct"]),
        ("Расчётная переменная доля, %", metrics["model_variable_rate_pct"]),
    ]
    for row in rows:
        summary.append(list(row))
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 28
    summary["A1"].font = Font(bold=True)

    expenses = workbook.create_sheet("Расходы")
    expenses.append(["Статья", "Всего, руб.", "На товар/комплект, руб.", "На изделие, руб."])
    for row in expense_rows:
        expenses.append([row["category"], row["total"], row["per_sold_product"], row["per_physical_item"]])
    expenses.freeze_panes = "A2"

    ladder = workbook.create_sheet("Ценовая сетка")
    ladder_headers = [
        "Изделий в товаре", "Себестоимость", "Целевая маржа", "Минимальная цена",
        "Цена со скидкой", "Базовая цена", "Цена после резерва 20%",
        "Расходы при минимальной цене", "Маржа на изделие",
    ]
    ladder.append(ladder_headers)
    for row in price_ladder:
        ladder.append(list(row.values()))
    ladder.freeze_panes = "A2"

    product_sheet = workbook.create_sheet("Товары")
    product_headers = list(products[0]) if products else ["internal_sku"]
    product_sheet.append(product_headers)
    for row in products:
        product_sheet.append([row.get(key) for key in product_headers])
    product_sheet.freeze_panes = "A2"

    if warnings:
        warning_sheet = workbook.create_sheet("Ограничения")
        warning_sheet.append(["Предупреждение"])
        for warning in warnings:
            warning_sheet.append([warning])
        warning_sheet.column_dimensions["A"].width = 120
    ensure_dir(path.parent)
    workbook.save(path)


def _load_catalog(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def _is_active_ozon(row: dict[str, str]) -> bool:
    value = str(row.get("active_ozon") or "").strip().lower()
    if value:
        return value in {"1", "true", "yes", "да"}
    return bool(str(row.get("ozon_offer_id") or "").strip())


def _pack_qty(row: dict[str, Any]) -> int:
    try:
        return max(1, int(float(str(row.get("pack_qty") or 1).replace(",", "."))))
    except ValueError:
        return 1


def _positive_decimal(value: Any, field: str) -> Decimal:
    result = _decimal_checked(value, field)
    if result <= 0:
        raise ValueError(f"{field} must be greater than zero")
    return result


def _non_negative_decimal(value: Any, field: str) -> Decimal:
    result = _decimal_checked(value, field)
    if result < 0:
        raise ValueError(f"{field} must be zero or greater")
    return result


def _decimal_checked(value: Any, field: str) -> Decimal:
    try:
        return Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field} must be a number") from exc


def _decimal(value: Any) -> Decimal:
    try:
        return Decimal(str(value or 0).replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _commission_fraction(value: Any) -> Decimal | None:
    if value in {None, ""}:
        return None
    commission = _decimal(value)
    if commission < 0:
        return None
    return commission / Decimal("100") if commission > Decimal("1") else commission


def _ceil_price(value: Decimal) -> Decimal:
    return (value / PRICE_STEP).to_integral_value(rounding=ROUND_CEILING) * PRICE_STEP


def _money(value: Any) -> float:
    return float(_decimal(value).quantize(Decimal("0.01")))


def _percent(value: Decimal) -> float:
    return float((value * Decimal("100")).quantize(Decimal("0.01")))


def _number_or_none(value: Any) -> float | None:
    if value in {None, ""}:
        return None
    return _money(value)


def _rub(value: Any) -> str:
    return f"{_money(value):,.2f} руб.".replace(",", " ")


def _rub_or_na(value: Any) -> str:
    return "не распределено" if value is None else _rub(value)
