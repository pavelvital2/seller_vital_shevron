from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, ROUND_CEILING
import json
from pathlib import Path
import re
import subprocess
from typing import Any

from openpyxl import Workbook, load_workbook

from takterra_agent.config import AppCredentials
from takterra_agent.reports.writer import ensure_dir, write_json


PROJECT_ROOT = Path(__file__).resolve().parents[3]


@dataclass
class Scheme:
    threshold: int
    fallback_no_promo: int
    fallback_over_threshold: int

    @property
    def label(self) -> str:
        return f"{self.threshold}-{self.fallback_no_promo}-{self.fallback_over_threshold}"


@dataclass
class PromoAggregate:
    nm_id: int
    vendor_code: str = ""
    title: str = ""
    subject: str = ""
    promo_ids: set[int] = field(default_factory=set)
    promo_names: set[str] = field(default_factory=set)
    statuses: set[str] = field(default_factory=set)
    min_upload_discount: int | None = None
    max_plan_price: Decimal | None = None
    rows_count: int = 0


def parse_scheme(value: str) -> Scheme:
    match = re.fullmatch(r"(\d+)-(\d+)-(\d+)", value.strip())
    if not match:
        raise ValueError("Scheme must look like 65-50-50")
    threshold, fallback_no_promo, fallback_over = [int(part) for part in match.groups()]
    return Scheme(
        threshold=threshold,
        fallback_no_promo=fallback_no_promo,
        fallback_over_threshold=fallback_over,
    )


def decimal_value(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    return Decimal(text)


def int_value(value: Any) -> int | None:
    dec = decimal_value(value)
    if dec is None:
        return None
    return int(dec)


def ceil_percent_from_price(base_price: Decimal, plan_price: Decimal) -> int:
    if base_price <= 0:
        return 0
    raw = (Decimal("1") - (plan_price / base_price)) * Decimal("100")
    return int(raw.to_integral_value(rounding=ROUND_CEILING))


def _run_snapshot(*, day: str, raw_dir: Path, prices_dir: Path) -> dict[str, Any]:
    completed = subprocess.run(
        [
            "node",
            "scripts/actions/wb_download_active_actions.js",
            "--date",
            day,
            "--out-dir",
            str(raw_dir),
            "--prices-dir",
            str(prices_dir),
        ],
        cwd=PROJECT_ROOT,
        check=False,
        capture_output=True,
        text=True,
        timeout=900,
    )
    if completed.returncode != 0:
        raise RuntimeError((completed.stderr or completed.stdout or "WB snapshot failed")[-2000:])
    return json.loads(completed.stdout)


def promo_id_from_filename(path: Path) -> int:
    match = re.search(r"promo-(\d+)-", path.name)
    if not match:
        raise ValueError(f"Cannot parse promo id from {path.name}")
    return int(match.group(1))


def promo_meta_by_id(snapshot_path: Path) -> dict[int, str]:
    snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
    return {int(item["actionID"]): item["name"] for item in snapshot.get("promos", [])}


def _clean_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ")).strip()


def _header_index(header: tuple[Any, ...]) -> dict[str, int]:
    return {_clean_header(name): pos for pos, name in enumerate(header)}


def _find_column(idx: dict[str, int], aliases: tuple[str, ...]) -> int | None:
    for alias in aliases:
        key = _clean_header(alias)
        if key in idx:
            return idx[key]
    return None


def read_promo_files(actions_dir: Path, excel_dir: Path) -> dict[int, PromoAggregate]:
    promo_names = promo_meta_by_id(actions_dir / "cabinet-actions-snapshot.json")
    aggregates: dict[int, PromoAggregate] = {}

    for path in sorted(excel_dir.glob("*.xlsx")):
        promo_id = promo_id_from_filename(path)
        promo_name = promo_names.get(promo_id, str(promo_id))
        # WB promo exports can contain incorrect worksheet dimensions.
        # Normal mode reads all columns; read_only mode may expose only A1.
        workbook = load_workbook(path, read_only=False, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        idx = _header_index(header)
        columns = {
            "status": _find_column(idx, ("Товар уже участвует в акции", "Item in promo (Yes / No)")),
            "subject": _find_column(idx, ("Предмет", "Subcategory")),
            "title": _find_column(idx, ("Наименование", "Item name")),
            "vendor_code": _find_column(idx, ("Артикул поставщика", "Seller item No.")),
            "nm_id": _find_column(idx, ("Артикул WB", "WB item No.")),
            "plan_price": _find_column(idx, ("Плановая цена для акции", "Target promo price")),
            "upload_discount": _find_column(idx, ("Загружаемая скидка для участия в акции", "Recommended promo discount")),
        }
        missing = [name for name, pos in columns.items() if pos is None]
        if missing:
            raise ValueError(f"{path.name}: missing columns {missing}")

        for row in rows:
            nm_id = int_value(row[columns["nm_id"]])
            if nm_id is None:
                continue
            plan_price = decimal_value(row[columns["plan_price"]])
            upload_discount = int_value(row[columns["upload_discount"]])
            if plan_price is None or upload_discount is None:
                continue

            item = aggregates.setdefault(nm_id, PromoAggregate(nm_id=nm_id))
            item.vendor_code = item.vendor_code or str(row[columns["vendor_code"]] or "")
            item.title = item.title or str(row[columns["title"]] or "")
            item.subject = item.subject or str(row[columns["subject"]] or "")
            item.promo_ids.add(promo_id)
            item.promo_names.add(promo_name)
            item.statuses.add(str(row[columns["status"]] or ""))
            item.rows_count += 1
            if item.min_upload_discount is None or upload_discount < item.min_upload_discount:
                item.min_upload_discount = upload_discount
            if item.max_plan_price is None or plan_price > item.max_plan_price:
                item.max_plan_price = plan_price
        workbook.close()

    return aggregates


def read_prices(path: Path) -> list[dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    goods = data.get("goods", [])
    if not isinstance(goods, list) or not goods:
        raise ValueError(f"No goods in {path}")
    return [row for row in goods if isinstance(row, dict)]


def build_rows(*, actions_dir: Path, prices_path: Path, scheme: Scheme) -> tuple[list[dict[str, Any]], dict[str, int]]:
    promo_by_nm = read_promo_files(actions_dir, actions_dir / "excel")
    goods = read_prices(prices_path)
    rows: list[dict[str, Any]] = []
    summary: dict[str, int] = defaultdict(int)

    for good in goods:
        nm_id = int(good["nmID"])
        base_price = decimal_value((good.get("prices") or [None])[0])
        current_discount = int_value(good.get("discount")) or 0
        current_discounted_price = decimal_value((good.get("discountedPrices") or [None])[0])
        agg = promo_by_nm.get(nm_id)

        if agg and base_price:
            min_discount = agg.min_upload_discount
            max_plan_price = agg.max_plan_price
            calc_discount = ceil_percent_from_price(base_price, max_plan_price)
            pre_threshold = min(min_discount, calc_discount)
            if pre_threshold > scheme.threshold:
                final_discount = scheme.fallback_over_threshold
                reason = f"скидка до порога > {scheme.threshold}% -> {scheme.fallback_over_threshold}%"
                summary["threshold_triggered"] += 1
            else:
                final_discount = pre_threshold
                reason = f"скидка до порога <= {scheme.threshold}%"
        else:
            min_discount = None
            max_plan_price = None
            calc_discount = None
            pre_threshold = None
            final_discount = scheme.fallback_no_promo
            reason = f"товара нет в активных акциях -> {scheme.fallback_no_promo}%"

        final_price = ""
        if base_price is not None:
            final_price = str((base_price * (Decimal("100") - Decimal(final_discount)) / Decimal("100")).quantize(Decimal("0.01")))
        delta = final_discount - current_discount
        if delta > 0:
            action = "увеличить скидку"
            summary["increase"] += 1
        elif delta < 0:
            action = "снизить скидку"
            summary["decrease"] += 1
        else:
            action = "не менять"
            summary["no_change"] += 1

        promo_count = len(agg.promo_ids) if agg else 0
        summary["in_promos" if promo_count else "outside_promos"] += 1
        if promo_count > 1:
            summary["multiple_promos"] += 1
        if delta != 0:
            summary["to_change"] += 1

        rows.append(
            {
                "Артикул WB": nm_id,
                "Артикул поставщика": good.get("vendorCode") or (agg.vendor_code if agg else ""),
                "Предмет": good.get("subject") or (agg.subject if agg else ""),
                "Наименование": good.get("title") or (agg.title if agg else ""),
                "Базовая цена": str(base_price or ""),
                "Текущая скидка": current_discount,
                "Текущая цена со скидкой": str(current_discounted_price or ""),
                "Акций": promo_count,
                "Акции": "; ".join(sorted(agg.promo_names)) if agg else "",
                "Статусы в файлах акций": ", ".join(sorted(agg.statuses)) if agg else "",
                "MIN загружаемая скидка": min_discount if min_discount is not None else "",
                "MAX плановая цена": str(max_plan_price or ""),
                "Расчетная скидка от плановой цены": calc_discount if calc_discount is not None else "",
                "Скидка до порога": pre_threshold if pre_threshold is not None else "",
                "Финальная скидка": final_discount,
                "Финальная цена": final_price,
                "Дельта, п.п.": delta,
                "Действие": action,
                "Причина": reason,
            }
        )

    summary["total_goods"] = len(goods)
    summary["promo_unique_total"] = len(promo_by_nm)
    return rows, dict(summary)


def write_csv_rows(rows: list[dict[str, Any]], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()), delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx_rows(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "WB actions discounts"
    headers = list(rows[0].keys())
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 12), 70)
    workbook.save(path)


def _price_value(value: str) -> int | float:
    decimal = Decimal(str(value))
    if decimal == decimal.to_integral_value():
        return int(decimal)
    return float(decimal)


def build_payload(rows: list[dict[str, Any]]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    changed_rows = [row for row in rows if int(row["Дельта, п.п."]) != 0]
    payload = {
        "dry_run": True,
        "upload_endpoint": "https://discounts-prices-api.wildberries.ru/api/v2/upload/task",
        "data": [
            {
                "nmID": int(row["Артикул WB"]),
                "price": _price_value(str(row["Базовая цена"])),
                "discount": int(row["Финальная скидка"]),
            }
            for row in changed_rows
        ],
        "note": "Preview only. Do not upload without explicit owner confirmation.",
    }
    return payload, changed_rows


def write_report(path: Path, *, run_id: str, scheme: Scheme, summary: dict[str, Any], artifacts: dict[str, str]) -> None:
    lines = [
        "# WB Actions Discount Dry Run",
        "",
        "Mode: read-only/dry-run. No discount upload was sent to WB.",
        f"Run ID: `{run_id}`",
        f"Scheme: `{scheme.label}`",
        "",
        "## Summary",
        "",
    ]
    for key in (
        "total_goods",
        "in_promos",
        "outside_promos",
        "multiple_promos",
        "threshold_triggered",
        "to_change",
        "increase",
        "decrease",
        "no_change",
        "promo_unique_total",
    ):
        lines.append(f"- `{key}`: {summary.get(key, 0)}")

    lines.extend(
        [
            "",
            "## Rule",
            "",
            f"- threshold: `{scheme.threshold}`",
            f"- fallback for goods outside active promos: `{scheme.fallback_no_promo}`",
            f"- fallback for over-threshold goods: `{scheme.fallback_over_threshold}`",
            "",
            "## Artifacts",
            "",
        ]
    )
    for key, value in sorted(artifacts.items()):
        lines.append(f"- `{key}`: `{value}`")
    lines.extend(["", "Apply is forbidden without explicit owner approval.", ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def run_wb_actions_discount_plan(
    *,
    credentials: AppCredentials,
    data_dir: Path = Path("data"),
    run_id: str | None = None,
    scheme_text: str = "65-50-50",
    actions_dir: Path | None = None,
    prices_json: Path | None = None,
) -> dict[str, Any]:
    _ = credentials
    scheme = parse_scheme(scheme_text)
    started_at = datetime.now()
    day = started_at.strftime("%Y-%m-%d")
    run_id = run_id or f"wb_actions_discount_plan_{scheme.label}_{started_at.strftime('%Y%m%dT%H%M%S')}"
    run_dir = ensure_dir(data_dir / "runs" / day / run_id)
    raw_actions_dir = ensure_dir(run_dir / "raw" / "actions")
    raw_prices_dir = ensure_dir(run_dir / "raw" / "prices")

    if actions_dir and prices_json:
        raw_actions_dir = actions_dir.resolve()
        prices_path = prices_json.resolve()
        snapshot = {
            "activePromos": len(json.loads((raw_actions_dir / "cabinet-actions-snapshot.json").read_text(encoding="utf-8")).get("promos", [])),
            "futurePromos": len(json.loads((raw_actions_dir / "cabinet-actions-snapshot.json").read_text(encoding="utf-8")).get("futurePromos", [])),
            "excelFiles": len(list((raw_actions_dir / "excel").glob("*.xlsx"))),
            "pricesCount": len(read_prices(prices_path)),
            "pricesPath": str(prices_path),
        }
    else:
        snapshot = _run_snapshot(day=day, raw_dir=raw_actions_dir, prices_dir=raw_prices_dir)
        prices_path = Path(snapshot["pricesPath"])
    rows, summary = build_rows(actions_dir=raw_actions_dir, prices_path=prices_path, scheme=scheme)
    rows.sort(key=lambda row: (row["Действие"] == "не менять", abs(int(row["Дельта, п.п."])) * -1, int(row["Артикул WB"])))
    payload, changed_rows = build_payload(rows)

    csv_path = run_dir / f"wb-discount-calculation-active-actions-{scheme.label}.csv"
    xlsx_path = run_dir / f"wb-discount-calculation-active-actions-{scheme.label}.xlsx"
    report_path = run_dir / f"wb-discount-calculation-active-actions-{scheme.label}.md"
    payload_path = run_dir / f"wb-upload-payload-preview-{scheme.label}.json"
    changed_csv_path = run_dir / f"wb-upload-changed-rows-{scheme.label}.csv"

    write_csv_rows(rows, csv_path)
    write_xlsx_rows(rows, xlsx_path)
    write_json(payload_path, payload)
    if changed_rows:
        write_csv_rows(changed_rows, changed_csv_path)
    else:
        changed_csv_path.write_text("", encoding="utf-8")

    summary = {
        **summary,
        "scheme": scheme.label,
        "changed_rows": len(changed_rows),
        "active_promos": int(snapshot.get("activePromos") or 0),
        "future_promos": int(snapshot.get("futurePromos") or 0),
        "excel_files": int(snapshot.get("excelFiles") or 0),
        "prices_count": int(snapshot.get("pricesCount") or 0),
    }
    artifacts = {
        "run_dir": str(run_dir),
        "report": str(report_path),
        "csv": str(csv_path),
        "xlsx": str(xlsx_path),
        "payload_preview": str(payload_path),
        "changed_rows_csv": str(changed_csv_path),
        "actions_snapshot": str(raw_actions_dir / "cabinet-actions-snapshot.json"),
        "prices_snapshot": str(prices_path),
        "summary": str(run_dir / "summary.json"),
    }
    result = {
        "run_id": run_id,
        "started_at": started_at.isoformat(timespec="seconds"),
        "mode": "dry-run",
        "summary": summary,
        "artifacts": artifacts,
        "apply_performed": False,
    }
    write_report(report_path, run_id=run_id, scheme=scheme, summary=summary, artifacts=artifacts)
    write_json(run_dir / "summary.json", result)
    return result
