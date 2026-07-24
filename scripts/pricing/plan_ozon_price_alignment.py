#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
from collections import Counter
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from html import escape
import json
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from seller_agent.config import load_credentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.ozon.adapter import OzonSellerAdapter
from seller_agent.reports.writer import ensure_dir, write_json
from seller_agent.safety.approvals import canonical_checksum
from seller_agent.tasks.pricing_status import normalize_ozon_price_item


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data"
CATALOG_PATH = DATA_DIR / "catalog/unified/products.csv"
PRIORITY_PATH = (
    DATA_DIR
    / "runs/2026-07-19/ozon_non_elastic_parser_review_20260719T081214"
    / "ozon_non_elastic_products.csv"
)
MOSCOW = ZoneInfo("Europe/Moscow")
PRICE_GRID = {
    1: (Decimal("530"), Decimal("650"), Decimal("1300")),
    2: (Decimal("860"), Decimal("1100"), Decimal("2200")),
    3: (Decimal("1180"), Decimal("1450"), Decimal("2900")),
    4: (Decimal("1500"), Decimal("1800"), Decimal("3600")),
    5: (Decimal("1820"), Decimal("2200"), Decimal("4400")),
}
REPORT_FIELDS = [
    "apply_stage",
    "selection_group",
    "selection_reason",
    "internal_sku",
    "offer_id",
    "product_id",
    "ozon_sku",
    "product_name",
    "pack_qty",
    "stock",
    "orders_30d",
    "revenue_30d",
    "elastic_active",
    "elastic_price_api_confirmed",
    "current_action_price",
    "current_min_price_control_only",
    "agreed_min_price_control_only",
    "min_price_matches_grid",
    "current_price",
    "target_price",
    "current_old_price",
    "target_old_price",
    "price_change",
    "old_price_change",
    "risk_code",
]


def _args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build an approval-only Ozon discount/base price alignment plan."
    )
    parser.add_argument("--elastic-run-id", required=True)
    parser.add_argument("--target-size", type=int, default=120)
    parser.add_argument("--run-id", default="")
    return parser.parse_args()


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return [
            {key: (value or "").strip() for key, value in row.items()}
            for row in csv.DictReader(handle)
        ]


def _write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _dec(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).replace(" ", "").replace("\u00a0", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _int(value: Any) -> int:
    return int(_dec(value))


def _money(value: Any) -> str:
    amount = _dec(value)
    if amount == amount.to_integral_value():
        return str(int(amount))
    return str(amount.quantize(Decimal("0.01")))


def _truthy(value: Any) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "y"}


def _is_callsign(row: dict[str, Any]) -> bool:
    text = " ".join(
        str(row.get(key) or "").lower()
        for key in ("internal_sku", "offer_id", "product_name")
    )
    internal_sku = str(row.get("internal_sku") or "").lower()
    offer_id = str(row.get("offer_id") or "").lower()
    return (
        "позывн" in text
        or internal_sku.startswith(("chev_pz_", "chev_kit2_pz_"))
        or offer_id.startswith(("pzol", "pzmh", "pzkit"))
        or "_pz_" in internal_sku
    )


def _elastic_in_price(row: dict[str, Any]) -> bool:
    actions = ((row.get("marketing_actions") or {}).get("actions") or [])
    return any("эластич" in str(action.get("title") or "").lower() for action in actions)


def _stock_from_product_info(row: dict[str, Any]) -> int:
    stocks = ((row.get("stocks") or {}).get("stocks") or [])
    return sum(
        _int(item.get("present"))
        for item in stocks
        if str(item.get("source") or "").lower() == "fbo"
    )


def _sales_by_sku(
    adapter: OzonSellerAdapter,
    *,
    date_from: str,
    date_to: str,
) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    result: dict[str, dict[str, Any]] = {}
    raw_rows: list[dict[str, Any]] = []
    for offset in range(0, 5000, 1000):
        payload = adapter.fetch_analytics_data(
            date_from=date_from,
            date_to=date_to,
            metrics=["ordered_units", "revenue"],
            dimensions=["sku"],
            limit=1000,
            offset=offset,
        )
        rows = ((payload.get("result") or {}).get("data") or []) if isinstance(payload, dict) else []
        raw_rows.extend(row for row in rows if isinstance(row, dict))
        for row in rows:
            dimensions = row.get("dimensions") or []
            metrics = row.get("metrics") or []
            if not dimensions:
                continue
            sku = str(dimensions[0].get("id") or "")
            result[sku] = {
                "orders": _int(metrics[0] if metrics else 0),
                "revenue": _dec(metrics[1] if len(metrics) > 1 else 0),
            }
        if len(rows) < 1000:
            break
    return result, raw_rows


def _latest_priority() -> dict[str, dict[str, str]]:
    if not PRIORITY_PATH.exists():
        return {}
    return {
        row.get("offer_id", ""): row
        for row in _read_csv(PRIORITY_PATH)
        if row.get("offer_id")
    }


def _row_sort_key(row: dict[str, Any]) -> tuple[Any, ...]:
    priority_rank = {"recovery_a": 0, "test_b": 1}.get(str(row.get("priority_code") or ""), 2)
    return (
        priority_rank,
        -_int(row.get("orders_30d")),
        -_dec(row.get("revenue_30d")),
        -_int(row.get("stock")),
        str(row.get("offer_id") or ""),
    )


def _xlsx(path: Path, summary: dict[str, Any], rows: list[dict[str, Any]], excluded: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Сводка"
    summary_values = [
        ("Run ID", summary["run_id"]),
        ("Статус", "На согласовании, изменений в Ozon нет"),
        ("Товаров в пакете", summary["selected_rows"]),
        ("Всего со старыми базовыми/скидочными ценами", summary["eligible_stale_rows"]),
        ("Останется после применения пакета", summary["remaining_stale_after_plan"]),
        ("Активный Elastic со старыми ценами", summary["elastic_stale_selected"]),
        ("Позывные вне Elastic", summary["callsign_stale_selected"]),
        ("Следующая приоритетная пачка", summary["next_batch_selected"]),
        ("Период продаж", summary["sales_period"]),
        ("Контроль min_price", "Только чтение; в payload отсутствует"),
        ("Checksum", summary["actions_checksum"]),
    ]
    for index, (label, value) in enumerate(summary_values, 1):
        summary_sheet.cell(index, 1, label).font = Font(bold=True)
        summary_sheet.cell(index, 2, value)
    summary_sheet.column_dimensions["A"].width = 38
    summary_sheet.column_dimensions["B"].width = 70

    plan_sheet = workbook.create_sheet("План на согласование")
    plan_sheet.append(REPORT_FIELDS)
    for cell in plan_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="185A4A")
    for row in rows:
        plan_sheet.append([row.get(field, "") for field in REPORT_FIELDS])
    plan_sheet.freeze_panes = "A2"
    plan_sheet.auto_filter.ref = plan_sheet.dimensions

    excluded_fields = [
        "offer_id", "internal_sku", "product_name", "pack_qty", "stock",
        "orders_30d", "current_price", "current_old_price", "reason",
    ]
    excluded_sheet = workbook.create_sheet("Не включены")
    excluded_sheet.append(excluded_fields)
    for cell in excluded_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6B7280")
    for row in excluded:
        excluded_sheet.append([row.get(field, "") for field in excluded_fields])
    excluded_sheet.freeze_panes = "A2"
    excluded_sheet.auto_filter.ref = excluded_sheet.dimensions

    grid_sheet = workbook.create_sheet("Ценовая сетка")
    grid_sheet.append(["Количество изделий", "min_price (контроль)", "price", "old_price"])
    for cell in grid_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="185A4A")
    for pack_qty, values in PRICE_GRID.items():
        grid_sheet.append([pack_qty, *[_money(value) for value in values]])

    for sheet in (plan_sheet, excluded_sheet, grid_sheet):
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            width = min(60, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[letter].width = width
    workbook.save(path)


def _markdown(summary: dict[str, Any]) -> str:
    pack_lines = "\n".join(
        f"- {pack_qty} шт.: `{count}` товаров"
        for pack_qty, count in sorted(summary["pack_distribution"].items(), key=lambda item: int(item[0]))
    )
    group_lines = "\n".join(
        f"- {group}: `{count}`"
        for group, count in summary["group_distribution"].items()
    )
    return f"""# Ozon: согласование базовых цен и цен со скидкой

## Итог

Подготовлен точный dry-run на `{summary['selected_rows']}` товаров. Изменения в Ozon не выполнялись.

- всего товаров со старой базовой ценой или ценой со скидкой: `{summary['eligible_stale_rows']}`;
- после применения этого пакета останется: `{summary['remaining_stale_after_plan']}`;
- активный Elastic со старыми обычными ценами: `{summary['elastic_stale_selected']}`;
- позывные вне Elastic со старыми ценами: `{summary['callsign_stale_selected']}`;
- следующая приоритетная пачка: `{summary['next_batch_selected']}`;
- текущих участников Elastic: `{summary['elastic_active_rows']}`;
- расхождений active-list и `marketing_actions` в price API: `{summary['elastic_confirmation_mismatches']}`.

## Группы

{group_lines}

## По количеству изделий

{pack_lines}

## Что будет меняться после отдельного согласования

Только цена со скидкой (`price`) и базовая зачеркнутая цена (`old_price`).
`min_price` показан в таблице исключительно для контроля, отсутствует в
payload и не должен изменяться. Управление участием в Elastic в payload также
отсутствует.

Порядок будущего apply:

1. fresh price/Elastic drift-check;
2. этап A: Elastic и позывные;
3. verify `price`, `old_price`, неизменность `min_price`, Elastic и цены покупателя;
4. этап B: следующая приоритетная пачка;
5. повторный verify.

## Источники

- Ozon Elastic run: `{summary['elastic_source_run_id']}`;
- Ozon Seller API `/v5/product/info/prices`;
- Ozon Seller API `/v1/analytics/data`, период `{summary['sales_period']}`;
- unified catalog `{CATALOG_PATH.relative_to(PROJECT_ROOT)}`.

## Риски и ограничения

- повышение обычной цены участника Elastic не должно менять цену покупателя в акции, но это обязательно проверяется после этапа A;
- строки с отсутствующим остатком не использовались для добора следующей пачки;
- `min_price` не является полем изменения и любое его расхождение при fresh-check блокирует конкретную строку;
- пакет действует только для checksum `{summary['actions_checksum']}`.
"""


def _html(summary: dict[str, Any], rows: list[dict[str, Any]]) -> str:
    table_rows = []
    for row in rows:
        search = escape(" ".join(str(value) for value in row.values()).lower())
        table_rows.append(
            "<tr data-search=\"" + search + "\">"
            + "".join(
                f"<td>{escape(str(row.get(field, '')))}</td>"
                for field in (
                    "apply_stage", "selection_group", "offer_id", "product_name", "pack_qty",
                    "stock", "orders_30d", "elastic_active", "current_action_price",
                    "current_min_price_control_only", "current_price", "target_price",
                    "current_old_price", "target_old_price", "risk_code",
                )
            )
            + "</tr>"
        )
    return f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Ozon: базовые цены и цены со скидкой</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#f3f5f7;color:#182026;font:14px/1.45 Arial,sans-serif;letter-spacing:0}}main{{max-width:1500px;margin:auto;padding:20px}}h1{{font-size:26px;margin:0 0 6px}}h2{{font-size:18px;margin:0 0 12px}}code{{overflow-wrap:anywhere;word-break:break-all}}.meta{{color:#667085}}.band{{background:#fff;border:1px solid #d6dce2;border-radius:6px;padding:16px;margin:14px 0;min-width:0}}.metrics{{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:10px}}.metric{{border-left:4px solid #176b57;background:#f7faf9;padding:10px;min-width:0;overflow-wrap:anywhere}}.metric b{{display:block;font-size:23px}}.warning{{border-left:4px solid #b54708;padding-left:10px}}input{{width:100%;max-width:520px;padding:10px;border:1px solid #aeb7c2;border-radius:4px}}.table{{overflow:auto;max-height:680px;border:1px solid #d6dce2;margin-top:10px;max-width:100%}}table{{border-collapse:collapse;width:100%;min-width:1500px}}th,td{{padding:7px;border-bottom:1px solid #e2e6ea;text-align:left;white-space:nowrap}}th{{position:sticky;top:0;background:#e9eef1;z-index:1}}.tag{{display:inline-block;padding:2px 7px;border-radius:4px;background:#e4f2ed}}@media(max-width:800px){{main{{padding:10px}}h1{{font-size:21px}}.metrics{{grid-template-columns:repeat(2,minmax(0,1fr))}}}}
</style></head><body><main><h1>Ozon: базовые цены и цены со скидкой</h1><div class="meta">{escape(summary['run_id'])} · dry-run · изменений в Ozon нет</div>
<section class="band"><div class="metrics"><div class="metric">Всего со старыми ценами<b>{summary['eligible_stale_rows']}</b></div><div class="metric">В текущем пакете<b>{summary['selected_rows']}</b></div><div class="metric">Из них в Elastic<b>{summary['elastic_stale_selected']}</b></div><div class="metric">Останется после пакета<b>{summary['remaining_stale_after_plan']}</b></div></div><p class="warning"><b>Минимальная цена не меняется.</b> План обновляет только цену со скидкой (<code>price</code>) и базовую зачеркнутую цену (<code>old_price</code>). Участие в Elastic также не меняется.</p></section>
<section class="band"><h2>Состав пакета</h2><p>Период продаж: {escape(summary['sales_period'])}. Текущих участников Elastic: {summary['elastic_active_rows']}. Checksum: <code>{escape(summary['actions_checksum'])}</code>.</p><input id="search" type="search" placeholder="Поиск по артикулу, названию или группе"><div class="table"><table><thead><tr>{''.join(f'<th>{escape(label)}</th>' for label in ['Этап','Группа','offer_id','Товар','Штук','Остаток','Заказы 30д','Elastic','Цена акции','Минимальная (контроль)','Со скидкой сейчас','Со скидкой цель','Базовая сейчас','Базовая цель','Риск'])}</tr></thead><tbody>{''.join(table_rows)}</tbody></table></div></section>
<section class="band"><h2>Будущий порядок применения</h2><ol><li>Fresh drift-check цен и Elastic.</li><li>Этап A: Elastic и позывные, затем verify.</li><li>Этап B: следующая пачка, затем verify.</li><li>Любое изменение min_price, состава Elastic или цены покупателя блокирует строку и выносится на новое согласование.</li></ol></section>
</main><script>const q=document.getElementById('search');q.addEventListener('input',()=>{{const v=q.value.toLowerCase();document.querySelectorAll('tbody tr').forEach(r=>r.hidden=!r.dataset.search.includes(v));}});</script></body></html>"""


def main() -> int:
    args = _args()
    started = datetime.now(MOSCOW)
    run_id = args.run_id or f"ozon_price_alignment_plan_{started:%Y%m%dT%H%M%S}"
    run_dir = ensure_dir(DATA_DIR / "runs" / started.date().isoformat() / run_id)
    raw_dir = ensure_dir(run_dir / "raw")
    processed_dir = ensure_dir(run_dir / "processed")

    elastic_dir = DATA_DIR / "runs" / started.date().isoformat() / args.elastic_run_id
    elastic_csv = elastic_dir / "ozon_elastic_dry_run.csv"
    product_info_path = elastic_dir / "processed/product_info_safe_snapshot.json"
    if not elastic_csv.exists() or not product_info_path.exists():
        raise FileNotFoundError(f"Elastic artifacts are incomplete: {elastic_dir}")
    elastic_rows = _read_csv(elastic_csv)
    product_info = json.loads(product_info_path.read_text(encoding="utf-8"))
    active_by_offer = {
        row["offer_id"]: row for row in elastic_rows if row.get("source_group") == "active"
    }

    catalog_rows = _read_csv(CATALOG_PATH)
    catalog_by_offer = {
        row["ozon_offer_id"]: row
        for row in catalog_rows
        if row.get("ozon_offer_id") and _truthy(row.get("active_ozon"))
    }
    scope_offers = sorted(set(catalog_by_offer) | set(active_by_offer))

    credentials = load_credentials()
    if credentials.ozon_seller is None:
        raise RuntimeError("Ozon Seller credentials are missing")
    adapter = OzonSellerAdapter(credentials.ozon_seller)
    fresh_price_rows = adapter.fetch_product_info_prices_by_offer_ids(scope_offers)
    write_json(raw_dir / "fresh_prices.json", fresh_price_rows)
    fresh_by_offer = {
        str(item.get("offer_id") or ""): normalize_ozon_price_item(item)
        for item in fresh_price_rows
        if item.get("offer_id")
    }

    sales_to = started.date() - timedelta(days=1)
    sales_from = sales_to - timedelta(days=29)
    sales, raw_sales = _sales_by_sku(
        adapter,
        date_from=sales_from.isoformat(),
        date_to=sales_to.isoformat(),
    )
    write_json(raw_dir / "sales_30d.json", raw_sales)
    priority = _latest_priority()

    eligible: list[dict[str, Any]] = []
    excluded: list[dict[str, Any]] = []
    confirmation_mismatches: list[dict[str, Any]] = []
    for offer_id, catalog in catalog_by_offer.items():
        current = fresh_by_offer.get(offer_id)
        if current is None:
            excluded.append({
                "offer_id": offer_id,
                "internal_sku": catalog.get("internal_sku", ""),
                "product_name": catalog.get("product_name", ""),
                "reason": "missing_in_fresh_price_snapshot",
            })
            continue
        pack_qty = _int(catalog.get("pack_qty"))
        if pack_qty not in PRICE_GRID:
            excluded.append({
                "offer_id": offer_id,
                "internal_sku": catalog.get("internal_sku", ""),
                "product_name": catalog.get("product_name", ""),
                "pack_qty": pack_qty,
                "reason": "unknown_pack_qty",
            })
            continue
        agreed_min, grid_price, grid_old = PRICE_GRID[pack_qty]
        current_price = _dec(current.get("price"))
        current_old = _dec(current.get("old_price"))
        current_min = _dec(current.get("min_price"))
        if current_price <= 0 or current_old <= 0:
            excluded.append({
                "offer_id": offer_id,
                "internal_sku": catalog.get("internal_sku", ""),
                "product_name": catalog.get("product_name", ""),
                "pack_qty": pack_qty,
                "current_price": _money(current_price),
                "current_old_price": _money(current_old),
                "reason": "invalid_current_price",
            })
            continue
        if current_price >= grid_price and current_old >= grid_old:
            continue
        target_price = max(current_price, grid_price)
        target_old = max(current_old, grid_old)
        if target_price < current_min:
            excluded.append({
                "offer_id": offer_id,
                "internal_sku": catalog.get("internal_sku", ""),
                "product_name": catalog.get("product_name", ""),
                "pack_qty": pack_qty,
                "current_price": _money(current_price),
                "current_old_price": _money(current_old),
                "reason": "target_price_below_current_min_price",
            })
            continue

        product_id = str(current.get("product_id") or catalog.get("ozon_product_id") or "")
        info = product_info.get(product_id) if isinstance(product_info, dict) else {}
        info = info if isinstance(info, dict) else {}
        stock = _stock_from_product_info(info)
        sku = str(current.get("sku") or catalog.get("ozon_sku") or "")
        sku_sales = sales.get(sku, {})
        direct_active = offer_id in active_by_offer
        price_api_active = _elastic_in_price(fresh_price_rows[next(
            index for index, item in enumerate(fresh_price_rows) if str(item.get("offer_id") or "") == offer_id
        )]) if offer_id in fresh_by_offer else False
        if direct_active != price_api_active:
            confirmation_mismatches.append({
                "offer_id": offer_id,
                "active_action_api": direct_active,
                "active_price_api": price_api_active,
            })
        prior = priority.get(offer_id, {})
        base = {
            "internal_sku": catalog.get("internal_sku", ""),
            "offer_id": offer_id,
            "product_id": product_id,
            "ozon_sku": sku,
            "product_name": catalog.get("product_name") or info.get("name") or "",
            "pack_qty": pack_qty,
            "stock": stock,
            "orders_30d": _int(sku_sales.get("orders")),
            "revenue_30d": _money(sku_sales.get("revenue")),
            "elastic_active": direct_active,
            "elastic_price_api_confirmed": price_api_active,
            "current_action_price": _money(current.get("marketing_seller_price")),
            "current_min_price_control_only": _money(current_min),
            "agreed_min_price_control_only": _money(agreed_min),
            "min_price_matches_grid": current_min == agreed_min,
            "current_price": _money(current_price),
            "target_price": _money(target_price),
            "current_old_price": _money(current_old),
            "target_old_price": _money(target_old),
            "price_change": f"{_money(current_price)} -> {_money(target_price)}",
            "old_price_change": f"{_money(current_old)} -> {_money(target_old)}",
            "currency_code": str(current.get("currency_code") or "RUB"),
            "priority_code": prior.get("recommendation_code", ""),
            "callsign": _is_callsign({
                "internal_sku": catalog.get("internal_sku"),
                "offer_id": offer_id,
                "product_name": catalog.get("product_name"),
            }),
            "risk_code": "" if current_min == agreed_min else "min_price_differs_from_grid_control_only",
        }
        eligible.append(base)

    elastic_stale = sorted([row for row in eligible if row["elastic_active"]], key=_row_sort_key)
    callsign_stale = sorted(
        [row for row in eligible if not row["elastic_active"] and row["callsign"]],
        key=_row_sort_key,
    )
    selected: list[dict[str, Any]] = []
    selected_offers: set[str] = set()

    def select(rows: list[dict[str, Any]], stage: str, group: str, reason: str) -> None:
        for row in rows:
            if row["offer_id"] in selected_offers:
                continue
            selected.append({**row, "apply_stage": stage, "selection_group": group, "selection_reason": reason})
            selected_offers.add(row["offer_id"])

    select(
        elastic_stale,
        "A",
        "elastic_active_stale",
        "Фактически участвует в Elastic, но обычные price/old_price ниже согласованной сетки.",
    )
    select(
        callsign_stale,
        "A",
        "callsign_stale",
        "Позывной: согласованная сетка применяется без увеличения CPC.",
    )

    remaining = [
        row for row in eligible
        if row["offer_id"] not in selected_offers and not row["callsign"] and row["stock"] > 0
    ]
    remaining.sort(key=_row_sort_key)
    needed = max(0, args.target_size - len(selected))
    for row in remaining[:needed]:
        code = str(row.get("priority_code") or "")
        if code == "recovery_a":
            group = "next_batch_priority_a"
            reason = "Приоритет A из предыдущего анализа; есть актуальный остаток."
        elif code == "test_b":
            group = "next_batch_priority_b"
            reason = "Приоритет B из предыдущего анализа; есть актуальный остаток."
        else:
            group = "next_batch_sales_stock"
            reason = "Добор по продажам за 30 дней и актуальному FBO-остатку."
        select([row], "B", group, reason)

    for row in eligible:
        if row["offer_id"] not in selected_offers:
            excluded.append({
                **row,
                "reason": "not_selected_after_priority_sort" if row["stock"] > 0 else "no_fbo_stock",
            })

    selected.sort(key=lambda row: (row["apply_stage"], _row_sort_key(row)))
    payload = {
        "prices": [
            {
                "offer_id": row["offer_id"],
                "price": row["target_price"],
                "old_price": row["target_old_price"],
                "currency_code": row["currency_code"],
            }
            for row in selected
        ]
    }
    forbidden_payload_keys = {
        key
        for item in payload["prices"]
        for key in item
        if key in {"min_price", "min_price_for_auto_actions_enabled", "manage_elastic_boosting_through_price"}
    }
    if forbidden_payload_keys:
        raise RuntimeError(f"Forbidden payload keys: {sorted(forbidden_payload_keys)}")
    if len(selected) < min(args.target_size, len(eligible)):
        raise RuntimeError(f"Selection is shorter than target despite eligible rows: {len(selected)}")

    checksum_payload = {
        "elastic_source_run_id": args.elastic_run_id,
        "selected_rows": [
            {key: row.get(key) for key in REPORT_FIELDS}
            for row in selected
        ],
        "write_payload": payload,
    }
    checksum = canonical_checksum(checksum_payload)
    group_distribution = Counter(str(row["selection_group"]) for row in selected)
    summary = {
        "run_id": run_id,
        "started_at": started.isoformat(timespec="seconds"),
        "mode": "dry_run",
        "overall_status": "warning" if confirmation_mismatches else "ok",
        "apply_performed": False,
        "elastic_source_run_id": args.elastic_run_id,
        "elastic_active_rows": len(active_by_offer),
        "elastic_confirmation_mismatches": len(confirmation_mismatches),
        "fresh_price_rows": len(fresh_price_rows),
        "eligible_stale_rows": len(eligible),
        "selected_rows": len(selected),
        "remaining_stale_after_plan": len(eligible) - len(selected),
        "elastic_stale_selected": sum(row["selection_group"] == "elastic_active_stale" for row in selected),
        "callsign_stale_selected": sum(row["selection_group"] == "callsign_stale" for row in selected),
        "next_batch_selected": sum(row["apply_stage"] == "B" for row in selected),
        "group_distribution": dict(group_distribution),
        "pack_distribution": dict(Counter(str(row["pack_qty"]) for row in selected)),
        "price_transition_distribution": dict(Counter(row["price_change"] for row in selected)),
        "old_price_transition_distribution": dict(Counter(row["old_price_change"] for row in selected)),
        "min_price_control_mismatches": sum(not row["min_price_matches_grid"] for row in selected),
        "payload_forbidden_keys": sorted(forbidden_payload_keys),
        "sales_period": f"{sales_from.isoformat()} - {sales_to.isoformat()}",
        "actions_checksum": checksum,
        "approval_required": True,
        "limitations": [
            "No Ozon writes were performed.",
            "min_price is control-only and absent from the write payload.",
            "Future apply must use a fresh partial drift-check and verify Elastic buyer price after stage A.",
        ],
        "artifacts": {},
    }

    plan_csv = processed_dir / "ozon_price_alignment_plan.csv"
    excluded_csv = processed_dir / "excluded_rows.csv"
    payload_path = processed_dir / "price_payload_preview.json"
    mismatch_path = processed_dir / "elastic_confirmation_mismatches.json"
    report_md = run_dir / "ozon_price_alignment_report.md"
    report_html = run_dir / "report.html"
    report_xlsx = run_dir / "ozon_price_alignment_plan.xlsx"
    _write_csv(plan_csv, selected, REPORT_FIELDS)
    _write_csv(excluded_csv, excluded, list(excluded[0]) if excluded else ["offer_id", "reason"])
    write_json(payload_path, payload)
    write_json(mismatch_path, confirmation_mismatches)
    report_md.write_text(_markdown(summary), encoding="utf-8")
    report_html.write_text(_html(summary, selected), encoding="utf-8")
    _xlsx(report_xlsx, summary, selected, excluded)

    summary["artifacts"] = {
        "run_dir": str(run_dir.relative_to(PROJECT_ROOT)),
        "report_md": str(report_md.relative_to(PROJECT_ROOT)),
        "report_html": str(report_html.relative_to(PROJECT_ROOT)),
        "report_xlsx": str(report_xlsx.relative_to(PROJECT_ROOT)),
        "plan_csv": str(plan_csv.relative_to(PROJECT_ROOT)),
        "payload_preview": str(payload_path.relative_to(PROJECT_ROOT)),
        "excluded_csv": str(excluded_csv.relative_to(PROJECT_ROOT)),
        "elastic_confirmation_mismatches": str(mismatch_path.relative_to(PROJECT_ROOT)),
        "fresh_prices": str((raw_dir / "fresh_prices.json").relative_to(PROJECT_ROOT)),
        "sales_30d": str((raw_dir / "sales_30d.json").relative_to(PROJECT_ROOT)),
    }
    manifest = write_summary_run_manifest(
        data_dir=DATA_DIR,
        run_dir=run_dir,
        summary=summary,
        task="ozon-price-alignment-plan",
        mode="dry_run",
        risk="high",
        marketplaces=["ozon"],
        inputs={
            "elastic_source_run_id": args.elastic_run_id,
            "target_size": args.target_size,
            "sales_period": summary["sales_period"],
        },
        lifecycle_status="pending_review",
        closed=False,
    )
    summary["artifacts"].update(manifest)
    write_json(run_dir / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
