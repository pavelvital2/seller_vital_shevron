#!/usr/bin/env python3
from __future__ import annotations

import argparse
from copy import copy
import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from seller_agent.config import load_credentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.ozon.adapter import OzonSellerAdapter
from seller_agent.reports.writer import ensure_dir, write_json


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data"
MOSCOW = ZoneInfo("Europe/Moscow")
# Owner-confirmed mapping that is not yet represented in the current unified catalog.
KNOWN_PACK_QTY_BY_SKU = {"2402436200": 1}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _quantity(value: Any, *, sheet: str, row: int, column: str) -> int:
    if isinstance(value, bool):
        raise RuntimeError(f"invalid quantity in {sheet}!{column}{row}: {value}")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"invalid quantity in {sheet}!{column}{row}: {value}") from exc
    if number < 0 or number != value:
        raise RuntimeError(f"invalid quantity in {sheet}!{column}{row}: {value}")
    return number


def load_pack_qty_by_sku(catalog_path: Path) -> dict[str, int]:
    result = dict(KNOWN_PACK_QTY_BY_SKU)
    with catalog_path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            sku = _text(row.get("ozon_sku"))
            pack_qty = _text(row.get("pack_qty"))
            if not sku or not pack_qty:
                continue
            value = int(pack_qty)
            previous = result.get(sku)
            if previous is not None and previous != value:
                raise RuntimeError(f"conflicting pack_qty for Ozon SKU {sku}: {previous} != {value}")
            result[sku] = value
    return result


def build_live_offer_by_sku(products: list[dict[str, Any]]) -> dict[str, str]:
    grouped: dict[str, set[str]] = {}
    for product in products:
        sku = _text(product.get("sku"))
        offer_id = _text(product.get("offer_id"))
        if sku and offer_id:
            grouped.setdefault(sku, set()).add(offer_id)
    duplicates = {sku: sorted(values) for sku, values in grouped.items() if len(values) != 1}
    if duplicates:
        raise RuntimeError(f"ambiguous live Ozon seller articles by SKU: {duplicates}")
    return {sku: next(iter(values)) for sku, values in grouped.items()}


def collect_source_rows(workbook_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    rows: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        if sheet.max_column < 6:
            raise RuntimeError(f"sheet {sheet.title} has fewer than six columns")
        for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if not any(value is not None for value in values[:6]):
                continue
            rows.append(
                {
                    "sheet": sheet.title,
                    "source_row": row_number,
                    "cluster": _text(values[0]),
                    "source_offer_id": _text(values[1]),
                    "ozon_sku": _text(values[2]),
                    "title": _text(values[3]),
                    "planned_goods_qty": _quantity(
                        values[4], sheet=sheet.title, row=row_number, column="E"
                    ),
                    "shipment_goods_qty": _quantity(
                        values[5], sheet=sheet.title, row=row_number, column="F"
                    ),
                }
            )
    return rows


def enrich_rows(
    source_rows: list[dict[str, Any]],
    *,
    live_offer_by_sku: dict[str, str],
    pack_qty_by_sku: dict[str, int],
) -> list[dict[str, Any]]:
    enriched: list[dict[str, Any]] = []
    for row in source_rows:
        sku = str(row["ozon_sku"])
        current_offer = live_offer_by_sku.get(sku)
        if not current_offer:
            raise RuntimeError(
                f"Ozon SKU is missing from fresh seller catalog: {sku} "
                f"({row['sheet']} row {row['source_row']})"
            )
        pack_qty = pack_qty_by_sku.get(sku)
        if pack_qty is None:
            raise RuntimeError(
                f"pack_qty is not confirmed for Ozon SKU {sku} "
                f"({row['sheet']} row {row['source_row']})"
            )
        enriched.append(
            {
                **row,
                "current_offer_id": current_offer,
                "seller_article_status": (
                    "updated" if current_offer != row["source_offer_id"] else "unchanged"
                ),
                "pack_qty": pack_qty,
                "planned_physical_qty": row["planned_goods_qty"] * pack_qty,
                "shipment_physical_qty": row["shipment_goods_qty"] * pack_qty,
            }
        )
    return enriched


def write_workbook(source_path: Path, output_path: Path, rows: list[dict[str, Any]]) -> None:
    workbook = load_workbook(source_path)
    rows_by_sheet: dict[str, dict[int, dict[str, Any]]] = {}
    for row in rows:
        rows_by_sheet.setdefault(str(row["sheet"]), {})[int(row["source_row"])] = row

    changed_fill = PatternFill("solid", fgColor="FFF2CC")

    for sheet in workbook.worksheets:
        original_widths = {
            letter: sheet.column_dimensions[letter].width
            for letter in ("C", "D", "E", "F")
        }
        sheet.insert_cols(3)
        sheet.column_dimensions["C"].width = 34
        sheet.column_dimensions["D"].width = original_widths["C"]
        sheet.column_dimensions["E"].width = original_widths["D"]
        sheet.column_dimensions["F"].width = original_widths["E"]
        sheet.column_dimensions["G"].width = original_widths["F"]
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_view.showGridLines = False

        for source_row, row in rows_by_sheet.get(sheet.title, {}).items():
            current_offer = str(row["current_offer_id"])
            cell = sheet.cell(source_row, 3, current_offer)
            template = sheet.cell(source_row, 2)
            cell.font = copy(template.font)
            cell.fill = copy(template.fill)
            cell.border = copy(template.border)
            cell.alignment = copy(template.alignment)
            cell.number_format = template.number_format
            cell.protection = copy(template.protection)
            if row["seller_article_status"] == "updated":
                cell.fill = changed_fill

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def build_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    sheets: dict[str, dict[str, int]] = {}
    for row in rows:
        sheet = sheets.setdefault(
            str(row["sheet"]),
            {
                "rows": 0,
                "planned_goods_qty": 0,
                "shipment_goods_qty": 0,
                "planned_physical_qty": 0,
                "shipment_physical_qty": 0,
                "updated_article_rows": 0,
                "unchanged_article_rows": 0,
            },
        )
        sheet["rows"] += 1
        for field in (
            "planned_goods_qty",
            "shipment_goods_qty",
            "planned_physical_qty",
            "shipment_physical_qty",
        ):
            sheet[field] += int(row[field])
        sheet[f"{row['seller_article_status']}_article_rows"] += 1

    unique_products: dict[str, dict[str, Any]] = {}
    for row in rows:
        unique_products.setdefault(str(row["ozon_sku"]), row)
    updated_article_mappings = sorted(
        [
            {
                "ozon_sku": sku,
                "source_offer_id": str(row["source_offer_id"]),
                "current_offer_id": str(row["current_offer_id"]),
                "title": str(row["title"]),
            }
            for sku, row in unique_products.items()
            if row["seller_article_status"] == "updated"
        ],
        key=lambda row: (row["source_offer_id"], row["ozon_sku"]),
    )
    pack_distribution: Counter[str] = Counter()
    for row in rows:
        pack_distribution[f"pack_{row['pack_qty']}"] += int(row["shipment_goods_qty"])
    return {
        "rows": len(rows),
        "unique_products": len(unique_products),
        "planned_goods_qty": sum(int(row["planned_goods_qty"]) for row in rows),
        "shipment_goods_qty": sum(int(row["shipment_goods_qty"]) for row in rows),
        "planned_physical_qty": sum(int(row["planned_physical_qty"]) for row in rows),
        "shipment_physical_qty": sum(int(row["shipment_physical_qty"]) for row in rows),
        "updated_article_rows": sum(row["seller_article_status"] == "updated" for row in rows),
        "updated_unique_products": len(
            {str(row["ozon_sku"]) for row in rows if row["seller_article_status"] == "updated"}
        ),
        "updated_article_mappings": updated_article_mappings,
        "unchanged_article_rows": sum(row["seller_article_status"] == "unchanged" for row in rows),
        "pack_qty_distribution_goods": dict(pack_distribution),
        "sheets": sheets,
    }


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _write_report(path: Path, summary: dict[str, Any], output_path: Path) -> None:
    lines = [
        "# Проверка Ozon-артикулов в отгрузке",
        "",
        f"Файл: `{output_path.name}`",
        "",
        f"- строк: `{summary['rows']}`;",
        f"- уникальных товаров: `{summary['unique_products']}`;",
        f"- товарных единиц к отгрузке: `{summary['shipment_goods_qty']}`;",
        f"- физических изделий к отгрузке: `{summary['shipment_physical_qty']}`;",
        f"- строк с обновленным артикулом: `{summary['updated_article_rows']}`;",
        f"- уникальных товаров с обновленным артикулом: `{summary['updated_unique_products']}`;",
        f"- несопоставленных строк: `0`.",
        "",
        "## По вкладкам",
        "",
        "| Вкладка | Товарных единиц | Физических изделий | Обновленных артикулов, строк |",
        "| --- | ---: | ---: | ---: |",
    ]
    for sheet, values in summary["sheets"].items():
        lines.append(
            f"| {sheet} | {values['shipment_goods_qty']} | "
            f"{values['shipment_physical_qty']} | {values['updated_article_rows']} |"
        )
    lines.extend(
        [
            "",
            "## Обновленные артикулы",
            "",
            "| SKU Ozon | Старый артикул | Актуальный артикул |",
            "| ---: | --- | --- |",
        ]
    )
    for row in summary["updated_article_mappings"]:
        lines.append(
            f"| {row['ozon_sku']} | {row['source_offer_id']} | {row['current_offer_id']} |"
        )
    lines.extend(
        [
            "",
            "Количество к отгрузке взято из последней количественной колонки исходной книги.",
            "Физические изделия рассчитаны как количество товарных единиц x pack_qty; две",
            "неразрезанные петлицы считаются одним изделием.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--run-id")
    parser.add_argument(
        "--catalog",
        type=Path,
        default=DATA_DIR / "catalog/unified/products.csv",
    )
    args = parser.parse_args()

    started = datetime.now(MOSCOW)
    run_id = args.run_id or f"ozon_shipment_sku_enrichment_{started:%Y%m%dT%H%M%S}"
    run_dir = ensure_dir(DATA_DIR / "runs" / started.date().isoformat() / run_id)
    processed_dir = ensure_dir(run_dir / "processed")
    raw_dir = ensure_dir(run_dir / "raw")
    output_path = args.output or run_dir / "Отгрузки_актуальные_артикулы_Ozon.xlsx"

    credentials = load_credentials()
    if not credentials.ozon_seller:
        raise RuntimeError("Ozon Seller API credentials are unavailable")
    live_products = OzonSellerAdapter(credentials.ozon_seller).fetch_product_list()
    write_json(raw_dir / "ozon_product_list_safe.json", live_products)

    source_rows = collect_source_rows(args.input)
    rows = enrich_rows(
        source_rows,
        live_offer_by_sku=build_live_offer_by_sku(live_products),
        pack_qty_by_sku=load_pack_qty_by_sku(args.catalog),
    )
    write_workbook(args.input, output_path, rows)
    _write_csv(processed_dir / "shipment_rows.csv", rows)
    summary = build_summary(rows)
    if summary["updated_article_mappings"]:
        _write_csv(
            processed_dir / "updated_article_mappings.csv",
            summary["updated_article_mappings"],
        )
    summary.update(
        {
            "run_id": run_id,
            "started_at": started.isoformat(timespec="seconds"),
            "overall_status": "ok",
            "mode": "read_only_local_file_enrichment",
            "source": "fresh Ozon Seller API /v3/product/list + unified pack_qty",
            "input_file": str(args.input),
            "output_file": str(output_path),
            "report": str(run_dir / "report.md"),
            "artifacts": {
                "workbook": str(output_path),
                "report": str(run_dir / "report.md"),
                "summary": str(run_dir / "summary.json"),
                "shipment_rows": str(processed_dir / "shipment_rows.csv"),
                "updated_article_mappings": str(
                    processed_dir / "updated_article_mappings.csv"
                ),
                "fresh_ozon_catalog": str(raw_dir / "ozon_product_list_safe.json"),
            },
        }
    )
    _write_report(run_dir / "report.md", summary, output_path)
    write_json(run_dir / "summary.json", summary)
    manifest = write_summary_run_manifest(
        data_dir=DATA_DIR,
        run_dir=run_dir,
        summary=summary,
        task="ozon-shipment-sku-enrichment",
        mode="read_only",
        risk="none",
        marketplaces=["ozon"],
        inputs={
            "input_file": str(args.input),
            "catalog": str(args.catalog),
            "quantity_source": "column_f_shipment_goods_qty",
        },
        lifecycle_status="closed",
        closed=True,
    )
    summary["artifacts"].update(manifest)
    write_json(run_dir / "summary.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
