#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_HALF_UP
from html import escape
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PARTICIPATING = {"да", "yes", "true", "1"}
SEGMENT_LABELS = {
    "scale_strong": "Масштабировать сильных",
    "scale_proven": "Масштабировать подтверждённых",
    "winner_top30_hold": "Сохранить ставку: уже top-30",
    "visibility_test": "Тест видимости",
    "launch_priority": "Запустить рекламу: есть спрос",
    "launch_discovery": "Запустить discovery-тест",
    "conversion_fix": "Исправить конверсию",
    "review_no_orders": "Разобрать без заказов",
    "repair_bid": "Исправить настройку ставки",
    "replenish_then_scale": "Пополнить, затем масштабировать",
    "blocked_low_stock": "Блок: низкий остаток",
    "blocked_no_stock": "Блок: нет остатка",
}


def read_csv(path: Path, *, delimiter: str = ",") -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def number(value: Any) -> float:
    try:
        return float(str(value or "0").replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def participating(row: dict[str, str]) -> bool:
    values = {
        item.strip().lower()
        for item in re.split(r"[,;]", row.get("Статусы в файлах акций", ""))
        if item.strip()
    }
    return bool(values & PARTICIPATING)


def rounded_bid(value: float) -> str:
    return str(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def target_bid(current: float, increase: float) -> str:
    if current <= 0:
        return ""
    return rounded_bid(min(max(current * (1 + increase), 1.0), 2.5))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields: list[str] = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def aggregate_campaigns(rows: list[dict[str, str]]) -> dict[str, dict[str, Any]]:
    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        grouped[row["nm_id"]].append(row)
    result: dict[str, dict[str, Any]] = {}
    for nm_id, items in grouped.items():
        totals = {
            field: sum(number(row.get(field)) for row in items)
            for field in ("views", "clicks", "atbs", "orders", "spend", "revenue")
        }
        totals.update(
            {
                "rows": items,
                "campaign_ids": sorted({row.get("advert_id", "") for row in items if row.get("advert_id")}),
                "campaign_names": sorted({row.get("campaign_name", "") for row in items if row.get("campaign_name")}),
                "has_bid": any(number(row.get("current_bid")) > 0 for row in items),
                "drr": totals["spend"] / totals["revenue"] * 100 if totals["revenue"] else None,
                "ctr": totals["clicks"] / totals["views"] * 100 if totals["views"] else None,
                "cr": totals["orders"] / totals["clicks"] * 100 if totals["clicks"] else None,
            }
        )
        result[nm_id] = totals
    return result


def merge_campaign_memberships(result: dict[str, dict[str, Any]], payload: Any) -> None:
    campaigns = payload.get("adverts", []) if isinstance(payload, dict) else payload
    for campaign in campaigns or []:
        advert_id = str(campaign.get("id") or campaign.get("advertId") or "")
        campaign_name = str((campaign.get("settings") or {}).get("name") or campaign.get("name") or "")
        for item in campaign.get("nm_settings") or campaign.get("nmSettings") or []:
            nm_id = str(item.get("nm_id") or item.get("nmId") or "")
            if not nm_id:
                continue
            bids = item.get("bids_kopecks") or item.get("bids") or {}
            current_bid = number(bids.get("search")) / 100
            if nm_id not in result:
                result[nm_id] = {
                    "views": 0,
                    "clicks": 0,
                    "atbs": 0,
                    "orders": 0,
                    "spend": 0,
                    "revenue": 0,
                    "rows": [],
                    "campaign_ids": [],
                    "campaign_names": [],
                    "has_bid": False,
                    "drr": None,
                    "ctr": None,
                    "cr": None,
                }
            target = result[nm_id]
            if advert_id and advert_id not in target["campaign_ids"]:
                target["campaign_ids"].append(advert_id)
                target["campaign_ids"].sort(key=int)
            if campaign_name and campaign_name not in target["campaign_names"]:
                target["campaign_names"].append(campaign_name)
                target["campaign_names"].sort()
            if current_bid > 0:
                target["has_bid"] = True
            if not any(str(row.get("advert_id") or "") == advert_id for row in target["rows"]):
                target["rows"].append(
                    {
                        "advert_id": advert_id,
                        "campaign_name": campaign_name,
                        "current_bid": rounded_bid(current_bid) if current_bid else "",
                    }
                )


def validate_complete_campaign_payload(campaigns_payload: Any, count_payload: Any) -> None:
    expected = {
        str(item.get("advertId") or "")
        for group in (count_payload.get("adverts", []) if isinstance(count_payload, dict) else [])
        if int(group.get("status") or 0) in {4, 9}
        for item in group.get("advert_list", [])
    }
    campaigns = campaigns_payload.get("adverts", []) if isinstance(campaigns_payload, dict) else campaigns_payload
    actual = {str(row.get("id") or row.get("advertId") or "") for row in campaigns or []}
    missing = sorted(expected - actual, key=int)
    if missing:
        raise RuntimeError(f"campaign membership snapshot is partial; missing active campaign ids: {missing}")


def classify(*, stock: float, campaign: dict[str, Any] | None, organic_sales: float, parser_position: float) -> str:
    position = parser_position or 9999
    if stock <= 0:
        return "blocked_no_stock"
    if stock < 4:
        return "blocked_low_stock"
    if not campaign:
        return "launch_priority" if organic_sales >= 1 or position <= 200 else "launch_discovery"
    if not campaign["has_bid"]:
        return "repair_bid"
    demand_30d = max(organic_sales, campaign["orders"])
    coverage_days = stock / (demand_30d / 30) if demand_30d > 0 else 9999
    if campaign["orders"] >= 1 and coverage_days < 14:
        return "replenish_then_scale"
    if campaign["orders"] >= 3 and campaign["drr"] is not None and campaign["drr"] <= 5:
        return "winner_top30_hold" if position <= 30 else "scale_strong"
    if campaign["orders"] >= 1 and campaign["drr"] is not None and campaign["drr"] <= 5:
        return "winner_top30_hold" if position <= 30 else "scale_proven"
    if campaign["orders"] == 0:
        if campaign["clicks"] >= 5 or campaign["atbs"] >= 3:
            return "conversion_fix"
        if campaign["views"] < 100 or campaign["clicks"] < 3:
            return "visibility_test"
        return "review_no_orders"
    return "review_no_orders"


def suggested_action(segment: str, *, selected: bool) -> str:
    if segment == "scale_strong":
        return "Повысить ставки на 20%; контроль 3 дня"
    if segment == "scale_proven":
        return "Повысить ставки на 10%; контроль 3 дня"
    if segment == "visibility_test" and selected:
        return "Повысить ставки на 15%; остановить после 100 показов без кликов"
    if segment == "launch_priority":
        return "Добавить в тематическую CPC-кампанию; старт 1,10 руб."
    if segment == "launch_discovery" and selected:
        return "Добавить в discovery CPC-кампанию; старт 1,00 руб."
    if segment == "conversion_fix":
        return "Ставку не повышать; проверить фото, цену, SEO и карточку"
    if segment == "winner_top30_hold":
        return "Ставку сохранить; контролировать органику и каннибализацию"
    if segment == "repair_bid":
        return "Восстановить текущую ставку/настройку кампании"
    if segment == "replenish_then_scale":
        return "Сначала пополнить запас минимум до 14 дней, затем повысить ставку"
    if segment.startswith("blocked_"):
        return "Не продвигать до пополнения остатка"
    return "Не менять в первой волне"


def bid_plan(campaign: dict[str, Any] | None, segment: str, *, selected: bool) -> str:
    if segment == "launch_priority":
        return "новая:1.10"
    if segment == "launch_discovery" and selected:
        return "новая:1.00"
    increase = {"scale_strong": 0.20, "scale_proven": 0.10}.get(segment)
    if segment == "visibility_test" and selected:
        increase = 0.15
    if increase is None or not campaign:
        return ""
    values = []
    for item in campaign["rows"]:
        current = number(item.get("current_bid"))
        if current <= 0:
            continue
        values.append(f"{item.get('advert_id')}:{rounded_bid(current)}->{target_bid(current, increase)}")
    return "; ".join(values)


def control_rule(segment: str, *, selected: bool) -> str:
    if not selected:
        return ""
    if segment in {"scale_strong", "scale_proven"}:
        return "3 дня; продолжать при ДРР <=5% и CPA <=15 руб.; снизить на 15% при >=2 заказах и ДРР >7% или CPA >25 руб."
    return "До 100 показов; 0 кликов -> карточка, не ставка. 10 кликов без заказа -> остановить рост ставки и исправить конверсию."


def write_xlsx(path: Path, sheets: list[tuple[str, list[dict[str, Any]]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title[:31])
        fields = list(rows[0]) if rows else ["Нет данных"]
        sheet.append(fields)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="365F91")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in rows:
            sheet.append([row.get(field, "") for field in fields])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, field in enumerate(fields, 1):
            max_len = max([len(str(field))] + [len(str(row.get(field, ""))) for row in rows[:500]])
            sheet.column_dimensions[get_column_letter(index)].width = min(max(max_len + 2, 10), 45)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def html_table(rows: list[dict[str, Any]], columns: list[tuple[str, str]], limit: int | None = None) -> str:
    items = rows[:limit] if limit else rows
    if not items:
        return '<p class="muted">Нет строк.</p>'
    head = "".join(f"<th>{escape(label)}</th>" for _, label in columns)
    body = []
    for row in items:
        search = " ".join(str(value) for value in row.values()).lower()
        cells = "".join(f"<td>{escape(str(row.get(key, '') or '-'))}</td>" for key, _ in columns)
        body.append(f'<tr data-search="{escape(search)}">{cells}</tr>')
    return f'<div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--before-actions", type=Path, required=True)
    parser.add_argument("--after-actions", type=Path, required=True)
    parser.add_argument("--promotion-plan", type=Path, required=True)
    parser.add_argument("--campaigns-json", type=Path, required=True)
    parser.add_argument("--campaign-count-json", type=Path, required=True)
    parser.add_argument("--parser-signals", type=Path, required=True)
    parser.add_argument("--stock-signals", type=Path, required=True)
    parser.add_argument("--sales-signals", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    before = {row["Артикул WB"]: row for row in read_csv(args.before_actions, delimiter=";")}
    after = {row["Артикул WB"]: row for row in read_csv(args.after_actions, delimiter=";")}
    exited = {nm_id for nm_id in before if participating(before[nm_id]) and not participating(after.get(nm_id, {}))}
    campaigns_payload = json.loads(args.campaigns_json.read_text(encoding="utf-8"))
    campaign_count_payload = json.loads(args.campaign_count_json.read_text(encoding="utf-8"))
    validate_complete_campaign_payload(campaigns_payload, campaign_count_payload)
    campaigns = aggregate_campaigns(read_csv(args.promotion_plan))
    merge_campaign_memberships(campaigns, campaigns_payload)
    stocks = {row["wb_nm_id"]: number(row.get("wb_stock_total")) for row in read_csv(args.stock_signals)}
    sales = {row["wb_nm_id"]: number(row.get("sales_units_30d")) for row in read_csv(args.sales_signals)}
    parser_rows = {row["wb_nm_id"]: row for row in read_csv(args.parser_signals)}

    base_rows: list[dict[str, Any]] = []
    for nm_id, action_row in after.items():
        campaign = campaigns.get(nm_id)
        stock = stocks.get(nm_id, 0)
        organic_sales = sales.get(nm_id, 0)
        parser_position = number(parser_rows.get(nm_id, {}).get("parser_best_position"))
        segment = classify(
            stock=stock,
            campaign=campaign,
            organic_sales=organic_sales,
            parser_position=parser_position,
        )
        base_rows.append(
            {
                "nmID": nm_id,
                "Артикул продавца": action_row.get("Артикул поставщика", ""),
                "Товар": action_row.get("Наименование", ""),
                "Сегмент": segment,
                "Название сегмента": SEGMENT_LABELS[segment],
                "Вышел из акции": "да" if nm_id in exited else "нет",
                "Участвует в акции сейчас": "да" if participating(action_row) else "нет",
                "Текущая скидка, %": action_row.get("Текущая скидка", ""),
                "Текущая цена, руб.": action_row.get("Текущая цена со скидкой", ""),
                "Остаток": stock,
                "Продажи API 30д": organic_sales,
                "Кампании": "; ".join(campaign["campaign_ids"]) if campaign else "",
                "Показы рекламы 30д": campaign["views"] if campaign else 0,
                "Клики рекламы 30д": campaign["clicks"] if campaign else 0,
                "Корзины рекламы 30д": campaign["atbs"] if campaign else 0,
                "Заказы рекламы 30д": campaign["orders"] if campaign else 0,
                "Расход рекламы 30д, руб.": round(campaign["spend"], 2) if campaign else 0,
                "Выручка рекламы 30д, руб.": round(campaign["revenue"], 2) if campaign else 0,
                "ДРР рекламы, %": round(campaign["drr"], 2) if campaign and campaign["drr"] is not None else "",
                "CTR, %": round(campaign["ctr"], 2) if campaign and campaign["ctr"] is not None else "",
                "CR клик-заказ, %": round(campaign["cr"], 2) if campaign and campaign["cr"] is not None else "",
                "Лучшая позиция парсера": parser_position or "",
                "Запросов парсера": parser_rows.get(nm_id, {}).get("parser_visible_queries", ""),
            }
        )

    visibility_pool = [row for row in base_rows if row["Сегмент"] == "visibility_test"]
    visibility_pool.sort(
        key=lambda row: (
            row["Вышел из акции"] == "да",
            number(row["Продажи API 30д"]),
            number(row["Лучшая позиция парсера"]) > 100,
            number(row["Запросов парсера"]),
            number(row["Остаток"]),
        ),
        reverse=True,
    )
    selected_visibility = {row["nmID"] for row in visibility_pool[:50]}
    discovery_pool = [row for row in base_rows if row["Сегмент"] == "launch_discovery"]
    discovery_pool.sort(
        key=lambda row: (row["Вышел из акции"] == "да", number(row["Остаток"])),
        reverse=True,
    )
    selected_discovery = {row["nmID"] for row in discovery_pool[:20]}

    batch_segments = {"scale_strong", "scale_proven", "launch_priority"}
    for row in base_rows:
        nm_id = row["nmID"]
        segment = row["Сегмент"]
        selected = segment in batch_segments or nm_id in selected_visibility or nm_id in selected_discovery
        row["Первая волна"] = "да" if selected else "нет"
        row["Рекомендованное действие"] = suggested_action(segment, selected=selected)
        row["План ставок"] = bid_plan(campaigns.get(nm_id), segment, selected=selected)
        row["Контроль"] = control_rule(segment, selected=selected)

    batch_rows = [row for row in base_rows if row["Первая волна"] == "да"]
    segment_counts = Counter(row["Сегмент"] for row in base_rows)
    batch_counts = Counter(row["Сегмент"] for row in batch_rows)
    summary_rows = [
        {
            "Сегмент": segment,
            "Название": SEGMENT_LABELS[segment],
            "Всего товаров": count,
            "В первой волне": batch_counts.get(segment, 0),
            "Заказы рекламы 30д": sum(number(row["Заказы рекламы 30д"]) for row in base_rows if row["Сегмент"] == segment),
            "Расход рекламы 30д, руб.": round(sum(number(row["Расход рекламы 30д, руб."]) for row in base_rows if row["Сегмент"] == segment), 2),
            "Выручка рекламы 30д, руб.": round(sum(number(row["Выручка рекламы 30д, руб."]) for row in base_rows if row["Сегмент"] == segment), 2),
        }
        for segment, count in segment_counts.most_common()
    ]

    all_csv = args.output_dir / "wb_portfolio_growth_plan_all.csv"
    batch_csv = args.output_dir / "wb_portfolio_growth_plan_wave1.csv"
    xlsx_path = args.output_dir / "wb_portfolio_growth_plan.xlsx"
    write_csv(all_csv, base_rows)
    write_csv(batch_csv, batch_rows)
    write_xlsx(
        xlsx_path,
        [
            ("Сводка", summary_rows),
            ("Первая волна", batch_rows),
            ("Все товары", base_rows),
            ("Исправить конверсию", [row for row in base_rows if row["Сегмент"] == "conversion_fix"]),
            ("Без остатка", [row for row in base_rows if row["Сегмент"].startswith("blocked_")]),
        ],
    )

    columns = [
        ("nmID", "nmID"),
        ("Товар", "Товар"),
        ("Название сегмента", "Сегмент"),
        ("Вышел из акции", "Вышел"),
        ("Остаток", "Остаток"),
        ("Заказы рекламы 30д", "Заказы"),
        ("ДРР рекламы, %", "ДРР, %"),
        ("Лучшая позиция парсера", "Позиция"),
        ("План ставок", "Ставки"),
        ("Рекомендованное действие", "Действие"),
    ]
    winners = [row for row in batch_rows if row["Сегмент"] in {"scale_strong", "scale_proven"}]
    visibility = [row for row in batch_rows if row["Сегмент"] == "visibility_test"]
    launches = [row for row in batch_rows if row["Сегмент"] in {"launch_priority", "launch_discovery"}]
    conversion = [row for row in base_rows if row["Сегмент"] == "conversion_fix"]
    blocked = [row for row in base_rows if row["Сегмент"].startswith("blocked_")]
    html = f"""<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>WB: портфельный план роста</title><style>
:root{{--bg:#f3f5f7;--panel:#fff;--text:#17212b;--muted:#65717d;--line:#d8dee4;--accent:#1769aa;--good:#14794c;--warn:#9b6500;--bad:#b3261e}}
*{{box-sizing:border-box}}html,body{{max-width:100%;overflow-x:hidden}}body{{margin:0;background:var(--bg);color:var(--text);font:14px/1.45 Arial,sans-serif;letter-spacing:0}}main{{max-width:1450px;margin:auto;padding:18px}}
h1{{font-size:25px;margin:0 0 6px}}h2{{font-size:18px;margin:24px 0 10px}}h3{{font-size:15px}}.muted{{color:var(--muted)}}.notice{{background:#fff8e6;border-left:4px solid var(--warn);padding:12px;margin:14px 0}}
.grid{{display:grid;grid-template-columns:repeat(5,minmax(130px,1fr));gap:9px;margin:14px 0}}.metric,.section{{background:var(--panel);border:1px solid var(--line);border-radius:6px}}.metric{{padding:11px}}.metric b{{display:block;font-size:22px}}.metric span{{color:var(--muted)}}
.section{{min-width:0;max-width:100%;padding:15px;margin:12px 0}}.steps{{display:grid;grid-template-columns:repeat(4,1fr);gap:9px}}.step{{border:1px solid var(--line);padding:10px;border-radius:5px}}.step b{{display:block;color:var(--accent);font-size:18px}}
input{{width:100%;max-width:460px;padding:9px;border:1px solid var(--line);border-radius:4px;margin:5px 0 12px}}.table-wrap{{width:100%;max-width:100%;overflow:auto;border:1px solid var(--line)}}table{{border-collapse:collapse;width:100%;min-width:1100px}}th,td{{padding:8px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}}th{{position:sticky;top:0;background:#edf2f5}}tr:hover{{background:#f7fafc}}
.good{{color:var(--good)}}.bad{{color:var(--bad)}}code{{background:#edf2f5;padding:1px 4px;border-radius:3px}}@media(max-width:800px){{main{{padding:11px}}.grid{{grid-template-columns:repeat(2,1fr)}}.steps{{grid-template-columns:1fr}}h1{{font-size:21px}}}}
</style></head><body><main><h1>WB: портфельный план роста продаж</h1><p class="muted">Read-only план. Ставки, кампании, цены и акции не изменялись.</p>
<div class="notice"><b>Цель:</b> компенсировать потерю акционного охвата не пятью товарами, а управляемой первой волной из <b>{len(batch_rows)}</b> товаров. Изменения разбиты по причине, чтобы не повышать ставки карточкам с плохой конверсией или без остатка.</div>
<div class="grid"><div class="metric"><b>{len(base_rows)}</b><span>товаров в магазине</span></div><div class="metric"><b>{len(exited)}</b><span>вышли из акций</span></div><div class="metric"><b class="good">{len(batch_rows)}</b><span>первая волна</span></div><div class="metric"><b>{len(winners)}</b><span>победителей масштабировать</span></div><div class="metric"><b>{len(launches)}</b><span>новых рекламных тестов</span></div></div>
<section class="section"><h2>Первая волна</h2><div class="steps"><div class="step"><b>{batch_counts['scale_strong']}</b>сильных: ставка +20%</div><div class="step"><b>{batch_counts['scale_proven']}</b>подтверждённых: ставка +10%</div><div class="step"><b>{batch_counts['visibility_test']}</b>низкий охват: тест +15%</div><div class="step"><b>{batch_counts['launch_priority'] + batch_counts['launch_discovery']}</b>новых подключений к CPC</div></div>
<p>Контрольный фактический расход: сначала до <b>100 руб./день</b> на весь портфель. Если три дня подряд ДРР не выше 5% и CPA не выше 15 руб., поднять контрольный уровень до <b>150 руб./день</b>. Для новых и visibility-тестов: 100 показов без кликов означают проблему карточки, а 10 кликов без заказа — проблему конверсии; ставку дальше не повышать. Для победителей hard-stop после минимум 2 заказов: ДРР выше 7% или CPA выше 25 руб.</p></section>
<section class="section"><h2>Что не масштабируем</h2><p><b>{len(conversion)}</b> карточек получили клики/корзины, но не дали заказов: сначала исправить карточку, цену и конверсию. <b>{len(blocked)}</b> товаров без достаточного остатка: реклама заблокирована до пополнения. <b>{segment_counts['winner_top30_hold']}</b> эффективных товаров уже находятся в top-30: ставки пока сохранить.</p></section>
<input id="filter" placeholder="Поиск по nmID, товару или сегменту" oninput="filterRows(this.value)">
<section class="section"><h2 class="good">Масштабировать победителей: {len(winners)}</h2>{html_table(winners, columns)}</section>
<section class="section"><h2>Тест недостатка показов: {len(visibility)}</h2>{html_table(visibility, columns)}</section>
<section class="section"><h2>Подключить новые товары: {len(launches)}</h2>{html_table(launches, columns)}</section>
<section class="section"><h2 class="bad">Исправить конверсию, не повышать ставку: {len(conversion)}</h2>{html_table(conversion, columns)}</section>
<section class="section"><h2>Файлы</h2><p><code>{escape(str(xlsx_path))}</code></p><p><code>{escape(str(batch_csv))}</code></p><p><code>{escape(str(all_csv))}</code></p></section>
<script>function filterRows(q){{q=q.toLowerCase().trim();document.querySelectorAll('tbody tr').forEach(r=>r.hidden=q&&!r.dataset.search.includes(q))}}</script></main></body></html>"""
    (args.output_dir / "report.html").write_text(html, encoding="utf-8")

    winner_spend = sum(number(row["Расход рекламы 30д, руб."]) for row in winners)
    winner_orders = sum(number(row["Заказы рекламы 30д"]) for row in winners)
    winner_revenue = sum(number(row["Выручка рекламы 30д, руб."]) for row in winners)
    summary = f"""# WB Portfolio Growth Plan

## Итог

- товаров в магазине: `{len(base_rows)}`;
- вышли из акций: `{len(exited)}`;
- первая волна: `{len(batch_rows)}`;
- масштабировать победителей: `{len(winners)}`;
- visibility-test: `{len(visibility)}`;
- новые рекламные тесты: `{len(launches)}`;
- исправить конверсию без повышения ставки: `{len(conversion)}`;
- блок по остатку: `{len(blocked)}`.

Победители за 30 дней: расход `{winner_spend:.2f}` руб., заказы `{winner_orders:.0f}`, выручка `{winner_revenue:.2f}` руб.

## Файлы

- `{args.output_dir / 'report.html'}`
- `{xlsx_path}`
- `{batch_csv}`
- `{all_csv}`
"""
    (args.output_dir / "summary.md").write_text(summary, encoding="utf-8")
    print(summary)
    print("segments", dict(segment_counts))
    print("wave1", dict(batch_counts))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
