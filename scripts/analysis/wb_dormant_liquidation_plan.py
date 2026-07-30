#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import date, datetime
from decimal import Decimal, ROUND_CEILING, ROUND_HALF_UP
from html import escape
import json
from pathlib import Path
import sys
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from seller_agent.config import load_credentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.wb.adapter import WbContentAdapter
from seller_agent.marketplaces.wb.analytics_adapter import WbAnalyticsAdapter
from seller_agent.marketplaces.wb.statistics_adapter import WbStatisticsAdapter
from seller_agent.reports.writer import write_json

from scripts.analysis.wb_dormant_inventory import (
    COST_PER_PHYSICAL_ITEM,
    LOGISTICS_PER_GOOD,
    RETAINED_PRICE_SHARE,
    _actual_price,
    _build_rows,
    _minimum_safe_discount,
    _promo_rows,
)


DEFAULT_MISSING_CAMPAIGN_ID = 37041670
DEFAULT_MISSING_CAMPAIGN_NAME = "Прочие форменные 01.06"
PRICE_DROP_LIMIT = Decimal("0.33")
DISCOUNT_STEP_LIMIT = 35


def _args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--action-run-dir", type=Path, required=True)
    parser.add_argument("--promotion-run-dir", type=Path, required=True)
    parser.add_argument("--parser-json", type=Path, required=True)
    parser.add_argument("--parser-meta", type=Path, required=True)
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--analysis-date", default=date.today().isoformat())
    parser.add_argument("--ad-reserve-per-good", default="20")
    parser.add_argument("--run-id", default="")
    return parser.parse_args()


def _read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _int(value: Any) -> int:
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


def _float(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _practical_floor(pack_qty: int, ad_reserve: Decimal) -> Decimal:
    exact = (
        COST_PER_PHYSICAL_ITEM * Decimal(pack_qty)
        + LOGISTICS_PER_GOOD
        + ad_reserve
    ) / RETAINED_PRICE_SHARE
    return (
        exact / Decimal("10")
    ).to_integral_value(rounding=ROUND_CEILING) * Decimal("10")


def _max_safe_upload_discount(
    *,
    base_price: Decimal,
    current_price: Decimal,
    current_discount: int,
    target_discount: int,
) -> int:
    minimum_upload_price = current_price * (Decimal("1") - PRICE_DROP_LIMIT)
    price_limited = _minimum_safe_discount(base_price, minimum_upload_price)
    return min(target_discount, current_discount + DISCOUNT_STEP_LIMIT, price_limited)


def _promotion_index(
    promotion_run_dir: Path,
) -> tuple[
    dict[int, dict[str, Any]],
    dict[tuple[int, int], dict[str, Any]],
    dict[int, float | None],
]:
    campaigns_payload = _read_json(promotion_run_dir / "raw" / "campaigns.json")
    campaigns = campaigns_payload.get("adverts") or []
    campaign_by_id: dict[int, dict[str, Any]] = {}
    membership: dict[tuple[int, int], dict[str, Any]] = {}
    for campaign in campaigns:
        if _int(campaign.get("status")) != 9:
            continue
        settings = campaign.get("settings") or {}
        if str(settings.get("payment_type") or "") != "cpc":
            continue
        advert_id = _int(campaign.get("id"))
        campaign_by_id[advert_id] = {
            "advert_id": advert_id,
            "campaign_name": str(settings.get("name") or ""),
        }
        for item in campaign.get("nm_settings") or []:
            nm_id = _int(item.get("nm_id"))
            membership[(advert_id, nm_id)] = {
                "advert_id": advert_id,
                "nm_id": nm_id,
                "campaign_name": str(settings.get("name") or ""),
                "current_bid": _float(
                    (item.get("bids_kopecks") or {}).get("search")
                )
                / 100,
            }

    stats: dict[tuple[int, int], dict[str, Any]] = defaultdict(
        lambda: {
            "views": 0,
            "clicks": 0,
            "atbs": 0,
            "orders": 0,
            "spend": 0.0,
            "revenue": 0.0,
        }
    )
    with (promotion_run_dir / "wb_promotion_products.csv").open(
        encoding="utf-8-sig", newline=""
    ) as source:
        for row in csv.DictReader(source):
            key = (_int(row.get("advert_id")), _int(row.get("nm_id")))
            target = stats[key]
            for field in ("views", "clicks", "atbs", "orders"):
                target[field] += _int(row.get(field))
            for field in ("spend", "revenue"):
                target[field] = round(target[field] + _float(row.get(field)), 2)

    winning_bids: dict[int, list[float]] = defaultdict(list)
    for key, metric in stats.items():
        member = membership.get(key)
        if member and metric["orders"] > 0 and member["current_bid"] > 0:
            winning_bids[key[0]].append(member["current_bid"])
    campaign_winner_median: dict[int, float | None] = {}
    for advert_id in campaign_by_id:
        values = sorted(Decimal(str(value)) for value in winning_bids.get(advert_id) or [])
        if not values:
            campaign_winner_median[advert_id] = None
            continue
        midpoint = len(values) // 2
        median = (
            values[midpoint]
            if len(values) % 2
            else (values[midpoint - 1] + values[midpoint]) / Decimal("2")
        )
        campaign_winner_median[advert_id] = float(
            median.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        )
    return membership, stats, campaign_winner_median


def _fresh_sources(run_dir: Path) -> None:
    raw = run_dir / "raw"
    expected = (
        raw / "sales_history.json",
        raw / "orders_30d.json",
        raw / "warehouse_stocks.json",
        raw / "cards.json",
    )
    if all(path.exists() and path.stat().st_size > 0 for path in expected):
        return
    credentials = load_credentials()
    if credentials.wb is None:
        raise RuntimeError("WB credentials are missing")
    raw.mkdir(parents=True, exist_ok=True)
    statistics_api = WbStatisticsAdapter(credentials.wb)
    write_json(
        raw / "sales_history.json",
        statistics_api.fetch_sales(date_from="2025-07-30T00:00:00", flag=0),
    )
    write_json(
        raw / "orders_30d.json",
        statistics_api.fetch_orders(date_from="2026-06-30T00:00:00", flag=0),
    )
    write_json(
        raw / "warehouse_stocks.json",
        WbAnalyticsAdapter(credentials.wb).fetch_wb_warehouse_stocks(),
    )
    write_json(
        raw / "cards.json",
        WbContentAdapter(credentials.wb).fetch_cards(limit=100),
    )


def _campaign_for_row(
    row: dict[str, Any],
    memberships: dict[tuple[int, int], dict[str, Any]],
) -> dict[str, Any]:
    matches = [
        item for (advert_id, nm_id), item in memberships.items() if nm_id == row["nm_id"]
    ]
    if matches:
        return sorted(matches, key=lambda item: item["advert_id"])[0]
    return {
        "advert_id": DEFAULT_MISSING_CAMPAIGN_ID,
        "nm_id": row["nm_id"],
        "campaign_name": DEFAULT_MISSING_CAMPAIGN_NAME,
        "current_bid": 0.0,
    }


def _content_priority(row: dict[str, Any]) -> tuple[str, str]:
    if (
        row["completed_sales_history"] == 0
        and not row["parser_visible_latest"]
        and row["ad_views_30d"] < 20
    ):
        return (
            "high",
            "никогда не продавался, нет parser-видимости и почти нет рекламных показов",
        )
    if not row["parser_visible_latest"] and row["ad_views_30d"] < 100:
        return ("medium", "нет свежей parser-видимости и меньше 100 рекламных показов")
    if row["ad_clicks_30d"] >= 10 and row["ad_orders_30d"] == 0:
        return ("medium", "есть 10+ кликов без рекламного заказа: проверить конверсию карточки")
    return ("normal", "сначала проверить результат цены, акции и CPC")


def _build_plan(
    *,
    dormant_rows: list[dict[str, Any]],
    action_run_dir: Path,
    promotion_run_dir: Path,
    ad_reserve: Decimal,
) -> list[dict[str, Any]]:
    offers_by_nm = _promo_rows(action_run_dir)
    memberships, stats, winner_medians = _promotion_index(promotion_run_dir)
    output: list[dict[str, Any]] = []

    for source in dormant_rows:
        nm_id = source["nm_id"]
        pack_qty = source["pack_qty"]
        base_price = Decimal(str(source["base_price"]))
        current_price = Decimal(str(source["current_price"]))
        current_discount = source["current_discount"]
        floor = _practical_floor(pack_qty, ad_reserve)
        offers = offers_by_nm.get(nm_id) or []
        eligible_offers = [
            offer
            for offer in offers
            if Decimal(str(offer["actual_price"])) >= floor
        ]
        chosen_offer = max(
            eligible_offers,
            key=lambda row: (
                Decimal(str(row["actual_price"])),
                -_int(row["required_discount"]),
            ),
            default=None,
        )
        if chosen_offer:
            price_mode = "action"
            target_discount = _int(chosen_offer["required_discount"])
            target_price = Decimal(str(chosen_offer["actual_price"]))
            target_action_id = _int(chosen_offer["action_id"])
            target_action_name = str(chosen_offer["action_name"])
        else:
            price_mode = "manual_clearance"
            target_discount = _minimum_safe_discount(base_price, floor)
            target_price = _actual_price(base_price, target_discount)
            target_action_id = 0
            target_action_name = ""

        upload_discount = _max_safe_upload_discount(
            base_price=base_price,
            current_price=current_price,
            current_discount=current_discount,
            target_discount=target_discount,
        )
        upload_price = _actual_price(base_price, upload_discount)
        stages_remaining = int(upload_discount != target_discount)

        campaign = _campaign_for_row(source, memberships)
        key = (campaign["advert_id"], nm_id)
        metric = stats.get(key) or {
            "views": source["ad_views_30d"],
            "clicks": source["ad_clicks_30d"],
            "atbs": source["ad_atbs_30d"],
            "orders": source["ad_orders_30d"],
            "spend": source["ad_spend_30d"],
            "revenue": source["ad_revenue_30d"],
        }
        current_bid = _float(campaign.get("current_bid"))
        empirical_bid = winner_medians.get(campaign["advert_id"])
        if current_bid <= 0:
            target_bid = empirical_bid or Decimal("2.10")
            bid_action = "add"
            bid_reason = "добавить отсутствующую карточку по медиане заказов кампании"
        elif metric["orders"] > 0 or source["open_orders_30d"] > 0:
            target_bid = current_bid
            bid_action = "keep"
            bid_reason = "ставка уже дала рекламный или открытый заказ"
        elif metric["clicks"] >= 10:
            target_bid = current_bid
            bid_action = "keep"
            bid_reason = "трафик уже есть; сначала проверить конверсию после снижения цены"
        elif empirical_bid is not None:
            target_bid = empirical_bid
            bid_action = (
                "increase"
                if target_bid > current_bid
                else "decrease"
                if target_bid < current_bid
                else "keep"
            )
            bid_reason = "медианная текущая ставка товаров с заказами в этой кампании"
        else:
            target_bid = current_bid
            bid_action = "keep"
            bid_reason = "в кампании нет доказанной winning bid; цену меняем, ставку не разгоняем"

        priority, priority_reason = _content_priority(source)
        output.append(
            {
                **source,
                "clearance_ad_reserve": float(ad_reserve),
                "clearance_floor": float(floor),
                "target_minimum": float(floor),
                "price_mode": price_mode,
                "target_action_id": target_action_id,
                "target_action_name": target_action_name,
                "target_discount": target_discount,
                "target_price": float(target_price),
                "upload_discount_stage1": upload_discount,
                "upload_price_stage1": float(upload_price),
                "requires_second_price_stage": bool(stages_remaining),
                "price_reason": (
                    "лучшая активная акция с ценой не ниже zero-margin floor"
                    if chosen_offer
                    else "все action-price ниже floor или акция не предложена; ручная распродажа"
                ),
                "advert_id": campaign["advert_id"],
                "campaign_name": campaign["campaign_name"],
                "current_bid_recalc": round(current_bid, 2),
                "campaign_winner_median_bid": empirical_bid or "",
                "target_bid_recalc": round(float(target_bid), 2),
                "bid_action": bid_action,
                "bid_reason": bid_reason,
                "post_apply_click_stop": 10,
                "post_apply_spend_stop": float(ad_reserve),
                "content_priority": priority,
                "content_priority_reason": priority_reason,
            }
        )
    return sorted(
        output,
        key=lambda row: (
            0 if row["price_mode"] == "action" else 1,
            -row["stock_physical_items"],
            row["nm_id"],
        ),
    )


def _summary(rows: list[dict[str, Any]], ad_reserve: Decimal) -> dict[str, Any]:
    stock_goods = sum(row["stock_goods"] for row in rows)
    stock_physical = sum(row["stock_physical_items"] for row in rows)
    target_seller_value = sum(
        Decimal(str(row["target_price"])) * Decimal(row["stock_goods"]) for row in rows
    )
    retained_value = target_seller_value * RETAINED_PRICE_SHARE
    logistics = LOGISTICS_PER_GOOD * Decimal(stock_goods)
    production_cost = COST_PER_PHYSICAL_ITEM * Decimal(stock_physical)
    cpc_reserve_total = ad_reserve * Decimal(stock_goods)
    forecast_buffer = retained_value - logistics - production_cost - cpc_reserve_total
    return {
        "cards": len(rows),
        "stock_goods": stock_goods,
        "stock_physical_items": stock_physical,
        "zero_margin_floors": {
            str(pack_qty): float(_practical_floor(pack_qty, ad_reserve))
            for pack_qty in sorted({row["pack_qty"] for row in rows})
        },
        "minimum_price_changes": sum(
            row["minimum_price"] != row["target_minimum"] for row in rows
        ),
        "action_cards": sum(row["price_mode"] == "action" for row in rows),
        "manual_clearance_cards": sum(
            row["price_mode"] == "manual_clearance" for row in rows
        ),
        "action_keep": sum(
            row["price_mode"] == "action" and row["currently_in_action"] for row in rows
        ),
        "action_join": sum(
            row["price_mode"] == "action" and not row["currently_in_action"] for row in rows
        ),
        "action_below_floor_or_missing": sum(
            row["price_mode"] == "manual_clearance" for row in rows
        ),
        "action_below_floor": sum(
            row["best_action_price"] > 0
            and row["best_action_price"] < row["clearance_floor"]
            for row in rows
        ),
        "action_not_offered": sum(row["best_action_price"] <= 0 for row in rows),
        "single_stage_price_rows": sum(
            not row["requires_second_price_stage"] for row in rows
        ),
        "two_stage_price_rows": sum(
            row["requires_second_price_stage"] for row in rows
        ),
        "price_target_distribution": dict(
            Counter(str(row["target_price"]) for row in rows)
        ),
        "cpc_add": sum(row["bid_action"] == "add" for row in rows),
        "cpc_increase": sum(row["bid_action"] == "increase" for row in rows),
        "cpc_decrease": sum(row["bid_action"] == "decrease" for row in rows),
        "cpc_keep": sum(row["bid_action"] == "keep" for row in rows),
        "cpc_target_distribution": dict(
            Counter(f"{row['target_bid_recalc']:.2f}" for row in rows)
        ),
        "maximum_test_budget": float(ad_reserve * Decimal(len(rows))),
        "forecast_target_seller_value": float(target_seller_value),
        "forecast_retained_after_variable_expenses": float(
            retained_value.quantize(Decimal("0.01"))
        ),
        "forecast_logistics": float(logistics.quantize(Decimal("0.01"))),
        "forecast_production_cost": float(production_cost),
        "forecast_cpc_reserve_all_sales": float(cpc_reserve_total),
        "forecast_buffer_after_all_modelled_costs": float(
            forecast_buffer.quantize(Decimal("0.01"))
        ),
        "content_high": sum(row["content_priority"] == "high" for row in rows),
        "content_medium": sum(row["content_priority"] == "medium" for row in rows),
        "open_order_cards": sum(row["open_orders_30d"] > 0 for row in rows),
    }


def _write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str] | None = None) -> None:
    headers = fields or (list(rows[0]) if rows else [])
    with path.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _write_xlsx(path: Path, rows: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "План распродажи"
    headers = list(rows[0]) if rows else []
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(field, "") for field in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="275D50")
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
            max(width + 2, 11), 55
        )
    summary_sheet = workbook.create_sheet("Сводка")
    summary_sheet.append(["Показатель", "Значение"])
    for key, value in summary.items():
        summary_sheet.append(
            [key, json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value]
        )
    summary_sheet.column_dimensions["A"].width = 40
    summary_sheet.column_dimensions["B"].width = 90
    workbook.save(path)


def _write_markdown(
    path: Path,
    *,
    run_id: str,
    summary: dict[str, Any],
    action_run_dir: Path,
    promotion_run_dir: Path,
    parser_meta: dict[str, Any],
) -> None:
    floors = ", ".join(
        f"{pack} шт. = {price:.0f} ₽"
        for pack, price in summary["zero_margin_floors"].items()
    )
    lines = [
        "# WB: dry-run распродажи залежалого остатка",
        "",
        f"Run ID: `{run_id}`. Изменений в WB нет.",
        "",
        "## Цель",
        "",
        "Продать неходовой остаток с нулевой маржой, покрывая себестоимость, "
        "средние расходы/логистику WB и 20 руб. CPC-резерва на проданный товар.",
        "",
        "## Сводка",
        "",
        f"- карточек: `{summary['cards']}`;",
        f"- остаток: `{summary['stock_goods']}` товаров / "
        f"`{summary['stock_physical_items']}` изделий;",
        f"- zero-margin floor: {floors};",
        f"- изменить минимальную цену: `{summary['minimum_price_changes']}`;",
        f"- включить/сохранить в лучшей допустимой акции: "
        f"`{summary['action_cards']}` (`{summary['action_keep']}` keep + "
        f"`{summary['action_join']}` join);",
        f"- ручная распродажная цена вне акции: "
        f"`{summary['manual_clearance_cards']}`;",
        f"- из них action-price ниже floor: `{summary['action_below_floor']}`, "
        f"акция не предложена: `{summary['action_not_offered']}`;",
        f"- один price upload: `{summary['single_stage_price_rows']}`; "
        f"два согласуемых upload: `{summary['two_stage_price_rows']}`;",
        f"- CPC: add `{summary['cpc_add']}`, increase `{summary['cpc_increase']}`, "
        f"decrease `{summary['cpc_decrease']}`, keep `{summary['cpc_keep']}`;",
        f"- максимальный тестовый CPC-бюджет: `{summary['maximum_test_budget']:.0f} ₽`;",
        f"- расчётная стоимость всего остатка по target price: "
        f"`{summary['forecast_target_seller_value']:.0f} ₽`; модельный остаток после "
        f"переменных расходов, логистики, себестоимости и CPC-резерва: "
        f"`{summary['forecast_buffer_after_all_modelled_costs']:.0f} ₽`;",
        f"- приоритетный content/SEO review: `{summary['content_high']}` high + "
        f"`{summary['content_medium']}` medium.",
        "",
        "## Логика",
        "",
        "1. Для всех карточек временно снизить minimum до zero-margin floor.",
        "2. Если активная акция даёт цену не ниже floor, выбрать акцию с самой "
        "высокой фактической ценой.",
        "3. Если все action-price ниже floor, установить ручную скидку с ценой "
        "не ниже floor.",
        "4. CPC-ставку брать не произвольно, а по медианной текущей ставке товаров "
        "с заказами в той же кампании. Если ставка уже дала заказ или 10+ кликов, "
        "сначала оставить её и проверить влияние новой цены.",
        "5. После применения для каждой карточки новый hard stop: 10 кликов или "
        "20 руб. расхода без заказа. При достижении стопа карточку убрать из CPC.",
        "6. Снижение цены, превышающее безопасный порог WB, выполнять двумя "
        "отдельно согласуемыми upload.",
        "",
        "## Источники",
        "",
        f"- actions: `{action_run_dir}`;",
        f"- promotion: `{promotion_run_dir}`;",
        f"- parser: `{parser_meta.get('run_date_max')}`, "
        f"`{', '.join(parser_meta.get('regions') or [])}`, "
        f"`{', '.join(parser_meta.get('query_packs') or [])}`;",
        "- WB Statistics API, Analytics API и Content API собраны непосредственно "
        "при построении плана.",
        "",
        "## Ограничения",
        "",
        "- Налог не учитывается.",
        "- 20 руб. — ограниченный CPC-резерв на проданный товар; неуспешные тесты "
        "могут дать общий расход до указанного максимального бюджета.",
        "- Parser покрывает только сохранённый набор запросов и не является полной "
        "выдачей WB.",
        "- План read-only. Для apply нужны отдельные checksummed price/minimum и CPC "
        "dry-run, owner approval, drift-check и verify.",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_html(
    path: Path,
    *,
    run_id: str,
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    parser_meta: dict[str, Any],
) -> None:
    table_rows = []
    for row in rows:
        search = escape(
            " ".join(str(value) for value in row.values()).lower(), quote=True
        )
        table_rows.append(
            f"""<tr data-mode="{escape(row['price_mode'])}" data-content="{escape(row['content_priority'])}" data-search="{search}">
<td><a href="https://www.wildberries.ru/catalog/{row['nm_id']}/detail.aspx">{row['nm_id']}</a><br><small>{escape(row['vendor_code'])}</small></td>
<td><b>{escape(row['title'])}</b><br><small>{escape(row['warehouses'])}</small></td>
<td class="num">{row['stock_goods']}<br><small>{row['stock_physical_items']} изделий</small></td>
<td class="num">{row['completed_sales_history']}<br><small>{escape(row['last_completed_sale'] or 'не было')}</small></td>
<td><b>{'Акция' if row['price_mode']=='action' else 'Ручная цена'}</b><br><small>{escape(row['target_action_name'] or row['price_reason'])}</small></td>
<td class="num">{row['current_price']:.0f} → <b>{row['target_price']:.0f} ₽</b><br><small>minimum {row['minimum_price']:.0f} → {row['target_minimum']:.0f} ₽</small></td>
<td class="num">{row['current_discount']}% → <b>{row['target_discount']}%</b><br><small>upload 1: {row['upload_discount_stage1']}% {'· нужен второй' if row['requires_second_price_stage'] else ''}</small></td>
<td>{escape(row['campaign_name'])}<br><small>{row['advert_id']}</small></td>
<td class="num">{row['current_bid_recalc']:.2f} → <b>{row['target_bid_recalc']:.2f} ₽</b><br><small>{escape(row['bid_action'])}: {escape(row['bid_reason'])}</small></td>
<td class="num">{row['ad_spend_30d']:.2f} ₽<br><small>{row['ad_clicks_30d']} кликов · {row['ad_orders_30d']} заказов</small></td>
<td><b>{escape(row['content_priority'])}</b><br><small>{escape(row['content_priority_reason'])}</small></td>
</tr>"""
        )
    floors = " · ".join(
        f"{pack} шт.: {price:.0f} ₽"
        for pack, price in summary["zero_margin_floors"].items()
    )
    html = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>WB: план распродажи залежалого остатка</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;background:#f3f5f4;color:#17201c;font:14px/1.45 Arial,sans-serif;letter-spacing:0}}main{{max-width:1700px;margin:auto;padding:16px}}
h1{{font-size:25px;margin:0}}h2{{font-size:18px;margin:0 0 10px}}small,.meta{{color:#66716c}}.band{{background:#fff;border:1px solid #d6ddda;border-radius:6px;padding:14px;margin:12px 0}}
.metrics{{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));gap:8px}}.metric{{background:#fff;border:1px solid #d6ddda;border-left:4px solid #28604f;padding:10px;min-width:0}}.metric b{{display:block;font-size:22px}}.warn{{border-left-color:#a36615}}
.steps{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:10px}}.step{{background:#f6f9f7;border-left:4px solid #28604f;padding:10px}}
.controls{{display:grid;grid-template-columns:2fr 1fr 1fr;gap:8px}}input,select{{width:100%;padding:9px;border:1px solid #aab5b0;border-radius:4px;background:#fff}}
.table{{overflow:auto;max-height:72vh;border:1px solid #d6ddda;background:#fff}}table{{border-collapse:collapse;width:100%;min-width:1900px}}th,td{{padding:8px;border-bottom:1px solid #e1e6e4;text-align:left;vertical-align:top}}th{{position:sticky;top:0;background:#e7eeeb;z-index:1;font-size:12px}}.num{{text-align:right;white-space:nowrap}}a{{color:#075e48}}
@media(max-width:900px){{main{{padding:9px}}h1{{font-size:21px}}.metrics{{grid-template-columns:repeat(2,minmax(0,1fr))}}.steps,.controls{{grid-template-columns:1fr}}}}
</style></head><body><main>
<h1>WB: план распродажи залежалого остатка</h1><div class="meta">{escape(run_id)} · dry-run · изменений в WB нет</div>
<section class="band"><b>Цель: продать неходовой остаток без прибыли, но не ниже полной себестоимости.</b> В floor включены изделие, средние расходы/логистика WB и 20 ₽ CPC-резерва на проданный товар.</section>
<section class="metrics">
<div class="metric"><span>Карточки</span><b>{summary['cards']}</b></div>
<div class="metric"><span>Остаток</span><b>{summary['stock_goods']}</b><span>{summary['stock_physical_items']} изделий</span></div>
<div class="metric"><span>В акции</span><b>{summary['action_cards']}</b><span>{summary['action_join']} добавить</span></div>
<div class="metric"><span>Ручная цена</span><b>{summary['manual_clearance_cards']}</b></div>
<div class="metric warn"><span>Два price upload</span><b>{summary['two_stage_price_rows']}</b></div>
<div class="metric"><span>Макс. CPC-тест</span><b>{summary['maximum_test_budget']:.0f} ₽</b></div>
</section>
<section class="band"><h2>Zero-margin floor</h2><p><b>{floors}</b></p></section>
<section class="band"><p>Расчётная стоимость всего остатка по целевым ценам: <b>{summary['forecast_target_seller_value']:.0f} ₽</b>. После переменных расходов, логистики, себестоимости и CPC-резерва модель оставляет около <b>{summary['forecast_buffer_after_all_modelled_costs']:.0f} ₽</b> на весь хвост. Это небольшой защитный запас, а не целевая прибыль.</p></section>
<section class="band"><div class="steps">
<div class="step"><b>Цена и minimum</b><p>Minimum опускается до floor. Если акция остаётся выше floor — выбирается акция с максимальной ценой. Иначе устанавливается ручная распродажная цена.</p></div>
<div class="step"><b>CPC</b><p>Ориентир — медианная ставка товаров с заказами в той же кампании. Стоп по карточке: 10 кликов или 20 ₽ после запуска без заказа.</p></div>
<div class="step"><b>Контроль</b><p>Ежедневно: заказы, выкупы, расход, клики и остаток. При hard stop товар исключается из CPC; цена и акция продолжают распродажу.</p></div>
</div></section>
<section class="band"><div class="controls"><input id="search" type="search" placeholder="Поиск по nmID, артикулу, названию"><select id="mode"><option value="">Все цены</option><option value="action">Акция</option><option value="manual_clearance">Ручная цена</option></select><select id="content"><option value="">Все content-приоритеты</option><option value="high">high</option><option value="medium">medium</option><option value="normal">normal</option></select></div></section>
<div class="table"><table><thead><tr><th>WB</th><th>Товар / склады</th><th>Остаток</th><th>Выкупы</th><th>Режим цены</th><th>Цена</th><th>Скидка</th><th>Кампания</th><th>CPC</th><th>Реклама 30д</th><th>Контент</th></tr></thead><tbody>{''.join(table_rows)}</tbody></table></div>
<section class="band"><p>Parser: {escape(str(parser_meta.get('run_date_max')))}, регион {escape(', '.join(parser_meta.get('regions') or []))}, пакет {escape(', '.join(parser_meta.get('query_packs') or []))}. Покрытие ограничено сохранённым набором запросов. Налог не учтён. Apply не выполнялся.</p></section>
</main><script>
const q=document.getElementById('search'),m=document.getElementById('mode'),c=document.getElementById('content');
function filter(){{const t=q.value.toLowerCase(),mv=m.value,cv=c.value;document.querySelectorAll('tbody tr').forEach(r=>r.hidden=!(r.dataset.search.includes(t)&&(!mv||r.dataset.mode===mv)&&(!cv||r.dataset.content===cv)));}}
[q,m,c].forEach(x=>x.addEventListener('input',filter));
</script></body></html>"""
    path.write_text(html, encoding="utf-8")


def main() -> int:
    args = _args()
    started = datetime.now().astimezone()
    run_id = args.run_id or f"wb_dormant_liquidation_plan_{started:%Y%m%dT%H%M%S}"
    run_dir = args.data_dir / "runs" / started.strftime("%Y-%m-%d") / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    processed = run_dir / "processed"
    processed.mkdir(exist_ok=True)
    _fresh_sources(run_dir)
    parser_rows = _read_json(args.parser_json)
    parser_meta = _read_json(args.parser_meta)
    write_json(processed / "parser_latest_visibility_all_store.json", parser_rows)
    write_json(processed / "parser_meta.json", parser_meta)

    analysis_date = date.fromisoformat(args.analysis_date)
    dormant_rows = _build_rows(
        run_dir=run_dir,
        action_run_dir=args.action_run_dir,
        promotion_run_dir=args.promotion_run_dir,
        analysis_date=analysis_date,
    )
    ad_reserve = Decimal(str(args.ad_reserve_per_good))
    rows = _build_plan(
        dormant_rows=dormant_rows,
        action_run_dir=args.action_run_dir,
        promotion_run_dir=args.promotion_run_dir,
        ad_reserve=ad_reserve,
    )
    summary = _summary(rows, ad_reserve)

    all_csv = run_dir / "liquidation_plan.csv"
    price_csv = run_dir / "price_minimum_plan.csv"
    cpc_csv = run_dir / "cpc_plan.csv"
    xlsx = run_dir / "liquidation_plan.xlsx"
    html = run_dir / "report.html"
    markdown = run_dir / "report.md"
    summary_path = run_dir / "summary.json"
    _write_csv(all_csv, rows)
    _write_csv(
        price_csv,
        rows,
        [
            "nm_id",
            "vendor_code",
            "internal_sku",
            "title",
            "pack_qty",
            "stock_goods",
            "minimum_price",
            "target_minimum",
            "current_price",
            "price_mode",
            "target_action_id",
            "target_action_name",
            "current_discount",
            "target_discount",
            "target_price",
            "upload_discount_stage1",
            "upload_price_stage1",
            "requires_second_price_stage",
            "price_reason",
        ],
    )
    _write_csv(
        cpc_csv,
        rows,
        [
            "advert_id",
            "campaign_name",
            "nm_id",
            "vendor_code",
            "title",
            "stock_goods",
            "current_bid_recalc",
            "campaign_winner_median_bid",
            "target_bid_recalc",
            "bid_action",
            "bid_reason",
            "ad_views_30d",
            "ad_clicks_30d",
            "ad_atbs_30d",
            "ad_orders_30d",
            "ad_spend_30d",
            "post_apply_click_stop",
            "post_apply_spend_stop",
        ],
    )
    _write_xlsx(xlsx, rows, summary)
    _write_markdown(
        markdown,
        run_id=run_id,
        summary=summary,
        action_run_dir=args.action_run_dir,
        promotion_run_dir=args.promotion_run_dir,
        parser_meta=parser_meta,
    )
    _write_html(
        html,
        run_id=run_id,
        rows=rows,
        summary=summary,
        parser_meta=parser_meta,
    )
    result = {
        "run_id": run_id,
        "mode": "dry_run",
        "overall_status": "pending_review",
        "goal": "liquidate dormant WB stock with zero margin and controlled CPC",
        "analysis_date": analysis_date.isoformat(),
        "collected_at": started.isoformat(timespec="seconds"),
        "summary": summary,
        "sources": {
            "action_run_dir": str(args.action_run_dir),
            "promotion_run_dir": str(args.promotion_run_dir),
            "parser_json": str(args.parser_json),
            "wb_api": [
                "Statistics API sales/orders",
                "Analytics API warehouse stocks",
                "Content API cards",
            ],
        },
        "artifacts": {
            "report_html": str(html),
            "report_md": str(markdown),
            "xlsx": str(xlsx),
            "all_csv": str(all_csv),
            "price_csv": str(price_csv),
            "cpc_csv": str(cpc_csv),
            "summary": str(summary_path),
        },
        "marketplace_write_performed": False,
    }
    write_json(summary_path, result)
    write_summary_run_manifest(
        data_dir=args.data_dir,
        run_dir=run_dir,
        summary=result,
        task="wb-dormant-liquidation-plan",
        mode="dry_run",
        risk="normal",
        marketplaces=["wb"],
        inputs={
            "action_run_dir": str(args.action_run_dir),
            "promotion_run_dir": str(args.promotion_run_dir),
            "ad_reserve_per_good": str(ad_reserve),
        },
        source_run_ids=[args.action_run_dir.name, args.promotion_run_dir.name],
        lifecycle_status="pending_review",
        closed=False,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
