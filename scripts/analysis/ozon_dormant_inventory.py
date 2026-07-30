#!/usr/bin/env python3
from __future__ import annotations

import argparse
import calendar
from collections import Counter, defaultdict
import csv
from datetime import date, datetime, timedelta
from html import escape
import json
import os
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from seller_agent.config import load_credentials
from seller_agent.core.run_manifest import write_summary_run_manifest
from seller_agent.marketplaces.ozon.adapter import OzonSellerAdapter


MOSCOW = ZoneInfo("Europe/Moscow")
DELIVERED = "OperationAgentDeliveredToCustomer"


def _args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--candidates",
        default=(
            "data/runs/2026-07-29/ozon_expense_reduction_20260729T1730/"
            "zero_buyouts_stock.csv"
        ),
    )
    parser.add_argument("--history-from", default="2023-01-01")
    parser.add_argument("--date-to", default=(date.today() - timedelta(days=1)).isoformat())
    parser.add_argument("--run-id", default="")
    parser.add_argument("--data-dir", default="data")
    parser.add_argument(
        "--cpc-csv",
        default=(
            "data/runs/2026-07-29/ozon_cpc_efficiency_20260729T171524/"
            "processed/by_sku_30d.csv"
        ),
    )
    parser.add_argument("--parser-env", default="/home/pavel/.parser-data-api.env")
    return parser.parse_args()


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as source:
        return list(csv.DictReader(source))


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    fields = list(rows[0]) if rows else []
    with path.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _int(value: Any, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _float(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _money(value: Any) -> float:
    return round(_float(value) + 1e-9, 2)


def _date(value: Any) -> date | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _month_chunks(start: date, end: date) -> list[tuple[date, date]]:
    chunks: list[tuple[date, date]] = []
    cursor = start
    while cursor <= end:
        month_end = date(
            cursor.year,
            cursor.month,
            calendar.monthrange(cursor.year, cursor.month)[1],
        )
        chunk_end = min(month_end, end)
        chunks.append((cursor, chunk_end))
        cursor = chunk_end + timedelta(days=1)
    return chunks


def _load_env(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip("\"").strip("'")
    return values


def _product_index(
    adapter: OzonSellerAdapter,
    candidate_skus: set[str],
) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    listed = adapter.fetch_product_list(visibility="ALL")
    by_sku = {
        str(row.get("sku") or ""): row
        for row in listed
        if str(row.get("sku") or "") in candidate_skus
    }
    product_ids = [
        str(row.get("product_id") or "")
        for row in by_sku.values()
        if str(row.get("product_id") or "")
    ]
    info = adapter.fetch_product_info(product_ids)
    info_by_sku = {
        str(row.get("sku") or ""): row
        for row in info
        if str(row.get("sku") or "") in candidate_skus
    }
    return info_by_sku, listed


def _finance_history(
    adapter: OzonSellerAdapter,
    *,
    start: date,
    end: date,
    candidate_skus: set[str],
) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "buyouts_total": 0,
            "buyouts_30d": 0,
            "buyouts_60d": 0,
            "buyouts_90d": 0,
            "buyouts_180d": 0,
            "first_buyout_date": "",
            "last_buyout_date": "",
        }
    )
    boundaries = {
        "buyouts_30d": end - timedelta(days=29),
        "buyouts_60d": end - timedelta(days=59),
        "buyouts_90d": end - timedelta(days=89),
        "buyouts_180d": end - timedelta(days=179),
    }
    for chunk_start, chunk_end in _month_chunks(start, end):
        operations = adapter.fetch_finance_transactions(
            date_from=f"{chunk_start.isoformat()}T00:00:00.000Z",
            date_to=f"{chunk_end.isoformat()}T23:59:59.999Z",
            operation_type=[DELIVERED],
        )
        for operation in operations:
            operation_day = _date(operation.get("operation_date"))
            if operation_day is None:
                continue
            items = operation.get("items") if isinstance(operation.get("items"), list) else []
            for item in items:
                if not isinstance(item, dict):
                    continue
                sku = str(item.get("sku") or "")
                if sku not in candidate_skus:
                    continue
                quantity = max(1, _int(item.get("quantity"), 1))
                row = result[sku]
                row["buyouts_total"] += quantity
                for field, boundary in boundaries.items():
                    if operation_day >= boundary:
                        row[field] += quantity
                current_first = _date(row["first_buyout_date"])
                current_last = _date(row["last_buyout_date"])
                if current_first is None or operation_day < current_first:
                    row["first_buyout_date"] = operation_day.isoformat()
                if current_last is None or operation_day > current_last:
                    row["last_buyout_date"] = operation_day.isoformat()
    return result


def _analytics_30d(
    adapter: OzonSellerAdapter,
    *,
    start: date,
    end: date,
) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for offset in range(0, 5000, 1000):
        payload = adapter.fetch_analytics_data(
            date_from=start.isoformat(),
            date_to=end.isoformat(),
            metrics=["ordered_units", "revenue", "cancellations", "returns"],
            dimensions=["sku"],
            limit=1000,
            offset=offset,
        )
        rows = ((payload.get("result") or {}).get("data") or []) if isinstance(payload, dict) else []
        for row in rows:
            dimensions = row.get("dimensions") or []
            metrics = row.get("metrics") or []
            if not dimensions:
                continue
            sku = str(dimensions[0].get("id") or "")
            result[sku] = {
                "orders_30d": _int(metrics[0] if metrics else 0),
                "orders_revenue_30d": _money(metrics[1] if len(metrics) > 1 else 0),
                "cancellations_30d": _int(metrics[2] if len(metrics) > 2 else 0),
                "returns_30d": _int(metrics[3] if len(metrics) > 3 else 0),
            }
        if len(rows) < 1000:
            break
    return result


def _fbo_statuses_30d(
    adapter: OzonSellerAdapter,
    *,
    start: date,
    end: date,
    candidate_skus: set[str],
) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = defaultdict(lambda: defaultdict(int))
    postings = adapter.fetch_fbo_postings(
        since=f"{start.isoformat()}T00:00:00.000Z",
        to=f"{end.isoformat()}T23:59:59.999Z",
    )
    for posting in postings:
        status = str(posting.get("status") or "unknown")
        products = posting.get("products") if isinstance(posting.get("products"), list) else []
        for item in products:
            if not isinstance(item, dict):
                continue
            sku = str(item.get("sku") or "")
            if sku not in candidate_skus:
                continue
            result[sku][status] += max(1, _int(item.get("quantity"), 1))
    return {sku: dict(statuses) for sku, statuses in result.items()}


def _cpc_index(path: Path) -> dict[str, dict[str, str]]:
    if not path.exists():
        return {}
    return {str(row.get("sku") or ""): row for row in _read_csv(path)}


def _parser_positions(
    *,
    skus: list[str],
    env_path: Path,
    run_date: date,
) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    env = {**_load_env(env_path), **os.environ}
    token = env.get("PARSER_DATA_API_TOKEN", "")
    base_url = env.get("PARSER_DATA_API_URL") or env.get(
        "PARSER_DATA_API_BASE_URL", "http://127.0.0.1:8787"
    )
    if not token:
        return {}, {"status": "not_confirmed", "reason": "Parser Data API token is missing"}

    result: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"parser_query_count": 0, "parser_best_position": 0, "parser_queries": []}
    )
    for start in range(0, len(skus), 30):
        params: list[tuple[str, str]] = [
            ("date", run_date.isoformat()),
            ("limit", "500"),
        ]
        params.extend(("sku", sku) for sku in skus[start : start + 30])
        url = f"{base_url.rstrip('/')}/warehouse/ozon/query-positions?{urlencode(params)}"
        request = Request(url, headers={"Authorization": f"Bearer {token}"})
        with urlopen(request, timeout=60) as response:
            payload = json.load(response)
        for row in payload.get("rows") or []:
            normalized = str(row.get("normalized_sku") or row.get("sku") or "")
            sku = normalized.removeprefix("OZN")
            if sku not in skus:
                continue
            target = result[sku]
            position = _int(row.get("absolute_position"))
            target["parser_query_count"] += 1
            if position and (
                not target["parser_best_position"]
                or position < target["parser_best_position"]
            ):
                target["parser_best_position"] = position
            target["parser_queries"].append(
                f"{row.get('query')}: {position or 'н/д'}"
            )
    return dict(result), {
        "status": "ok",
        "endpoint": "/warehouse/ozon/query-positions",
        "date": run_date.isoformat(),
        "filters": "exact sku list",
        "query_pack_scope": "30 collected queries",
    }


def _stock(info: dict[str, Any]) -> tuple[int, int]:
    stocks = info.get("stocks") if isinstance(info.get("stocks"), dict) else {}
    rows = stocks.get("stocks") if isinstance(stocks.get("stocks"), list) else []
    present = sum(
        _int(row.get("present"))
        for row in rows
        if isinstance(row, dict) and str(row.get("source") or "").lower() == "fbo"
    )
    reserved = sum(
        _int(row.get("reserved"))
        for row in rows
        if isinstance(row, dict) and str(row.get("source") or "").lower() == "fbo"
    )
    return present, reserved


def _price_index(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {str(row.get("offer_id") or ""): row for row in rows}


def _price_values(row: dict[str, Any]) -> dict[str, Any]:
    price = row.get("price") if isinstance(row.get("price"), dict) else {}
    marketing = (
        row.get("marketing_actions")
        if isinstance(row.get("marketing_actions"), dict)
        else {}
    )
    actions = marketing.get("actions") if isinstance(marketing.get("actions"), list) else []
    action_titles = [
        str(action.get("title") or "")
        for action in actions
        if isinstance(action, dict) and str(action.get("title") or "")
    ]
    return {
        "current_price": _money(price.get("price")),
        "current_buyer_price": _money(price.get("marketing_seller_price")),
        "current_old_price": _money(price.get("old_price")),
        "current_min_price": _money(price.get("min_price")),
        "active_actions": "; ".join(action_titles[:8]),
        "ozon_actions_exist": bool(marketing.get("ozon_actions_exist")),
    }


def _classify(
    *,
    end: date,
    created: date | None,
    history_start: date,
    history: dict[str, Any],
) -> tuple[str, str, str]:
    age = (end - created).days if created else None
    last_buyout = _date(history.get("last_buyout_date"))
    total = _int(history.get("buyouts_total"))
    history_complete = bool(created and created >= history_start)

    if _int(history.get("buyouts_30d")) > 0:
        return (
            "recovered_after_cutoff",
            "Оставить и наблюдать",
            "После исходного среза появился завершённый выкуп.",
        )
    if total == 0:
        qualifier = "за всё время существования карточки" if history_complete else "в доступной истории"
        if age is not None and age < 30:
            return (
                "new_under_30d_no_buyouts",
                "Наблюдать до 30 дней",
                f"Новая карточка, выкупов {qualifier} пока нет.",
            )
        if age is not None and age < 90:
            return (
                "never_sold_30_89d",
                "Доработать карточку и дать ограниченный тест",
                f"Выкупов {qualifier} нет, карточке {age} дней.",
            )
        if age is not None and age < 180:
            return (
                "never_sold_90_179d",
                "Не пополнять; SEO/контент-тест или распродажа",
                f"Выкупов {qualifier} нет, карточке {age} дней.",
            )
        return (
            "never_sold_180d_plus",
            "Не пополнять; распродать или вывести",
            f"Выкупов {qualifier} нет, карточка существует не менее 180 дней.",
        )

    days_since = (end - last_buyout).days if last_buyout else 9999
    if days_since >= 180:
        return (
            "dormant_180d_plus",
            "Не пополнять; распродать или вывести",
            f"Последний выкуп был {days_since} дней назад.",
        )
    if days_since >= 90:
        return (
            "dormant_90_179d",
            "Не пополнять; перезапуск карточки или распродажа",
            f"Последний выкуп был {days_since} дней назад.",
        )
    if days_since >= 60:
        return (
            "dormant_60_89d",
            "Доработать карточку/SEO и дать тест",
            f"Последний выкуп был {days_since} дней назад.",
        )
    return (
        "dormant_30_59d",
        "Оставить под наблюдением",
        f"Последний выкуп был {days_since} дней назад.",
    )


def _decision_adjustment(
    *,
    base_group: str,
    base_decision: str,
    base_reason: str,
    history: dict[str, Any],
    analytics: dict[str, Any],
    postings: dict[str, Any],
    cpc: dict[str, str],
    parser: dict[str, Any],
) -> tuple[str, str]:
    decision = base_decision
    reasons = [base_reason]
    spend = _float(cpc.get("spend"))
    ad_orders = _int(cpc.get("orders"))
    orders = _int(analytics.get("orders_30d"))
    total = _int(history.get("buyouts_total"))
    delivering = _int(postings.get("delivering"))
    delivered = _int(postings.get("delivered"))
    cancelled = _int(postings.get("cancelled"))

    if delivering or delivered:
        decision = "Дождаться завершения текущих заказов; затем оценить"
        reasons.append(
            f"Сейчас в доставке {delivering}, доставлено без завершённой финансовой операции {delivered}."
        )
    elif str(cpc.get("signal") or "") == "расход без заказов":
        decision = "Сначала убрать из CPC; затем исправить карточку"
        reasons.append(f"CPC израсходовал {spend:.2f} ₽ без рекламных заказов.")
    elif orders > 0 and _int(history.get("buyouts_30d")) == 0:
        if base_group.startswith("never_sold"):
            decision = "Сначала проверить статусы заказов; затем принимать решение"
        reasons.append(
            f"За 30 дней было {orders} заказов, но завершённых выкупов нет: проверить невыкупы."
        )
    if cancelled:
        reasons.append(f"За 30 дней отменено отправлений: {cancelled}.")
    if not _int(parser.get("parser_query_count")):
        reasons.append("Не найден в свежем срезе по 30 запросам; это не равно полной невидимости Ozon.")
    if total >= 20 and "распродать или вывести" in decision.lower():
        decision = "Не пополнять; попробовать точечный перезапуск"
        reasons.append(f"Исторически товар выкупали {total} раз, спрос ранее был подтверждён.")
    return decision, " ".join(reasons)


def _build_rows(
    *,
    candidates: list[dict[str, str]],
    info_by_sku: dict[str, dict[str, Any]],
    price_by_offer: dict[str, dict[str, Any]],
    history_by_sku: dict[str, dict[str, Any]],
    analytics_by_sku: dict[str, dict[str, Any]],
    postings_by_sku: dict[str, dict[str, Any]],
    cpc_by_sku: dict[str, dict[str, str]],
    parser_by_sku: dict[str, dict[str, Any]],
    history_start: date,
    end: date,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for candidate in candidates:
        sku = str(candidate.get("sku") or "")
        offer_id = str(candidate.get("offer_id") or "")
        info = info_by_sku.get(sku) or {}
        history = history_by_sku.get(sku) or {}
        analytics = analytics_by_sku.get(sku) or {}
        postings = postings_by_sku.get(sku) or {}
        cpc = cpc_by_sku.get(sku) or {}
        parser = parser_by_sku.get(sku) or {}
        created = _date(info.get("created_at"))
        last_buyout = _date(history.get("last_buyout_date"))
        stock_products, reserved = _stock(info)
        pack_qty = max(1, _int(candidate.get("pack_qty"), 1))
        base_group, base_decision, base_reason = _classify(
            end=end,
            created=created,
            history_start=history_start,
            history=history,
        )
        decision, reason = _decision_adjustment(
            base_group=base_group,
            base_decision=base_decision,
            base_reason=base_reason,
            history=history,
            analytics=analytics,
            postings=postings,
            cpc=cpc,
            parser=parser,
        )
        prices = _price_values(price_by_offer.get(offer_id) or {})
        statuses = info.get("statuses") if isinstance(info.get("statuses"), dict) else {}
        rows.append(
            {
                "group": base_group,
                "recommendation": decision,
                "reason": reason,
                "sku": sku,
                "offer_id": offer_id,
                "title": str(info.get("name") or candidate.get("title") or ""),
                "pack_qty": pack_qty,
                "created_date": created.isoformat() if created else "",
                "card_age_days": (end - created).days if created else "",
                "history_complete_from_card_creation": bool(
                    created and created >= history_start
                ),
                "status": str(statuses.get("status_name") or ""),
                "stock_products": stock_products,
                "reserved_products": reserved,
                "stock_physical_pieces": stock_products * pack_qty,
                "buyouts_total": _int(history.get("buyouts_total")),
                "first_buyout_date": str(history.get("first_buyout_date") or ""),
                "last_buyout_date": str(history.get("last_buyout_date") or ""),
                "days_since_last_buyout": (
                    (end - last_buyout).days if last_buyout else ""
                ),
                "buyouts_30d": _int(history.get("buyouts_30d")),
                "buyouts_60d": _int(history.get("buyouts_60d")),
                "buyouts_90d": _int(history.get("buyouts_90d")),
                "buyouts_180d": _int(history.get("buyouts_180d")),
                "orders_30d": _int(analytics.get("orders_30d")),
                "cancellations_30d": _int(analytics.get("cancellations_30d")),
                "returns_30d": _int(analytics.get("returns_30d")),
                "fbo_cancelled_30d": _int(postings.get("cancelled")),
                "fbo_delivering_30d": _int(postings.get("delivering")),
                "fbo_delivered_pending_finance_30d": _int(postings.get("delivered")),
                "cpc_active": str(cpc.get("in_current_campaign") or ""),
                "cpc_views_30d": _int(cpc.get("views")),
                "cpc_clicks_30d": _int(cpc.get("clicks")),
                "cpc_orders_30d": _int(cpc.get("orders")),
                "cpc_spend_30d": _money(cpc.get("spend")),
                "cpc_drr_pct": _money(cpc.get("drr_percent")),
                "parser_query_count": _int(parser.get("parser_query_count")),
                "parser_best_position": _int(parser.get("parser_best_position")),
                "parser_queries": "; ".join(parser.get("parser_queries") or []),
                **prices,
            }
        )
    group_order = {
        "never_sold_180d_plus": 0,
        "dormant_180d_plus": 1,
        "never_sold_90_179d": 2,
        "dormant_90_179d": 3,
        "dormant_60_89d": 4,
        "never_sold_30_89d": 5,
        "dormant_30_59d": 6,
        "new_under_30d_no_buyouts": 7,
        "recovered_after_cutoff": 8,
    }
    return sorted(
        rows,
        key=lambda row: (
            group_order.get(str(row["group"]), 99),
            -_int(row["stock_physical_pieces"]),
        ),
    )


def _summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    groups = Counter(str(row["group"]) for row in rows)
    never_sold = [row for row in rows if str(row["group"]).startswith("never_sold")]
    dormant_90 = [
        row
        for row in rows
        if row["group"] in {
            "never_sold_180d_plus",
            "never_sold_90_179d",
            "dormant_180d_plus",
            "dormant_90_179d",
        }
    ]
    recovered = [row for row in rows if row["group"] == "recovered_after_cutoff"]
    orders_without_buyout = [
        row
        for row in rows
        if _int(row["orders_30d"]) > 0 and _int(row["buyouts_30d"]) == 0
    ]
    cpc_without_orders = [
        row
        for row in rows
        if str(row["recommendation"]).startswith("Сначала убрать из CPC")
    ]
    parser_visible = [row for row in rows if _int(row["parser_query_count"]) > 0]
    current_delivery = [
        row
        for row in rows
        if _int(row["fbo_delivering_30d"])
        or _int(row["fbo_delivered_pending_finance_30d"])
    ]
    return {
        "candidate_skus": len(rows),
        "current_stock_products": sum(_int(row["stock_products"]) for row in rows),
        "current_stock_physical_pieces": sum(
            _int(row["stock_physical_pieces"]) for row in rows
        ),
        "never_sold_skus": len(never_sold),
        "never_sold_stock_products": sum(_int(row["stock_products"]) for row in never_sold),
        "decision_priority_skus_90d_plus": len(dormant_90),
        "decision_priority_stock_products": sum(
            _int(row["stock_products"]) for row in dormant_90
        ),
        "recovered_after_cutoff": len(recovered),
        "orders_without_completed_buyout_skus": len(orders_without_buyout),
        "orders_without_completed_buyout_units": sum(
            _int(row["orders_30d"]) for row in orders_without_buyout
        ),
        "cpc_spend_without_orders_skus": len(cpc_without_orders),
        "cpc_spend_without_orders": _money(
            sum(_float(row["cpc_spend_30d"]) for row in cpc_without_orders)
        ),
        "parser_visible_skus": len(parser_visible),
        "parser_not_found_skus": len(rows) - len(parser_visible),
        "current_delivery_skus": len(current_delivery),
        "current_delivery_units": sum(
            _int(row["fbo_delivering_30d"])
            + _int(row["fbo_delivered_pending_finance_30d"])
            for row in current_delivery
        ),
        "cancelled_postings_30d": sum(_int(row["fbo_cancelled_30d"]) for row in rows),
        "groups": dict(groups),
    }


def _write_xlsx(path: Path, rows: list[dict[str, Any]], metrics: dict[str, Any]) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Сводка"
    summary_sheet.append(["Показатель", "Значение"])
    for key, value in metrics.items():
        summary_sheet.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value])

    sheet = workbook.create_sheet("Товары")
    fields = list(rows[0]) if rows else []
    sheet.append(fields)
    for row in rows:
        sheet.append([row.get(field, "") for field in fields])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="275D4A")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        header = str(column_cells[0].value or "")
        if header in {"title", "reason", "active_actions", "parser_queries"}:
            width = 48
        elif header in {"offer_id", "recommendation"}:
            width = 30
        else:
            width = min(22, max(11, len(header) + 2))
        sheet.column_dimensions[letter].width = width
    summary_sheet.column_dimensions["A"].width = 38
    summary_sheet.column_dimensions["B"].width = 32
    workbook.save(path)


def _write_markdown(
    path: Path,
    *,
    run_id: str,
    collected_at: str,
    history_start: date,
    end: date,
    metrics: dict[str, Any],
    parser_meta: dict[str, Any],
) -> None:
    groups = metrics["groups"]
    lines = [
        "# Ozon: товары без выкупов за исходные 30 дней",
        "",
        "## Итог",
        "",
        (
            f"- Исходный список: **{metrics['candidate_skus']} SKU**. "
            f"Он означает отсутствие выкупов только в срезе 29.06-28.07, а не за всё время."
        ),
        f"- Никогда не продавались в подтверждённой истории: **{metrics['never_sold_skus']} SKU**.",
        (
            f"- Требуют первоочередного решения из-за 90+ дней без выкупа: "
            f"**{metrics['decision_priority_skus_90d_plus']} SKU**."
        ),
        f"- После исходного среза снова начали продаваться: **{metrics['recovered_after_cutoff']} SKU**.",
        (
            f"- Есть заказы без завершённого выкупа: "
            f"**{metrics['orders_without_completed_buyout_skus']} SKU / "
            f"{metrics['orders_without_completed_buyout_units']} заказов**."
        ),
        (
            f"- Из них ещё доставляются или ожидают финансового завершения: "
            f"**{metrics['current_delivery_skus']} SKU / "
            f"{metrics['current_delivery_units']} товаров**; отменено "
            f"**{metrics['cancelled_postings_30d']}**."
        ),
        (
            f"- Подтверждённый CPC без рекламных заказов: "
            f"**{metrics['cpc_spend_without_orders_skus']} SKU / "
            f"{metrics['cpc_spend_without_orders']:.2f} ₽**."
        ),
        (
            f"- Текущий остаток целевой группы: **{metrics['current_stock_products']} товаров / "
            f"{metrics['current_stock_physical_pieces']} физических изделий**."
        ),
        "",
        "## Группы",
        "",
    ]
    for group, count in sorted(groups.items()):
        lines.append(f"- `{group}`: `{count}` SKU")
    lines.extend(
        [
            "",
            "## Источники",
            "",
            f"- Ozon Seller API, завершённые выкупы: `{history_start}` - `{end}`.",
            "- Ozon Seller API: карточки, дата создания, цены и текущий FBO-остаток.",
            "- Ozon Seller API `/v1/analytics/data`: заказы/отмены/возвраты за последние 30 дней.",
            "- Ozon Performance API: последний готовый CPC-срез за 30 дней.",
            (
                f"- Parser Data API: `{parser_meta.get('endpoint', 'не подтверждено')}`, "
                f"дата `{parser_meta.get('date', 'не подтверждено')}`."
            ),
            "",
            "## Ограничения",
            "",
            "- Дата создания карточки не равна дате первой поставки на склад.",
            "- Отсутствие в парсере означает только отсутствие в собранных 30 запросах и глубине 500.",
            "- Рекомендации read-only. Вывоз, цены, акции, реклама и карточки не изменялись.",
            "",
            f"Run ID: `{run_id}`. Собрано: `{collected_at}`.",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_html(
    path: Path,
    *,
    run_id: str,
    collected_at: str,
    history_start: date,
    end: date,
    rows: list[dict[str, Any]],
    metrics: dict[str, Any],
    parser_meta: dict[str, Any],
) -> None:
    group_labels = {
        "never_sold_180d_plus": "Никогда не продавался, 180+ дней",
        "dormant_180d_plus": "Раньше продавался, пауза 180+ дней",
        "never_sold_90_179d": "Никогда не продавался, 90-179 дней",
        "dormant_90_179d": "Пауза 90-179 дней",
        "dormant_60_89d": "Пауза 60-89 дней",
        "never_sold_30_89d": "Никогда не продавался, 30-89 дней",
        "dormant_30_59d": "Пауза 30-59 дней",
        "new_under_30d_no_buyouts": "Новый, до 30 дней",
        "recovered_after_cutoff": "Продажи восстановились",
    }
    table_rows = []
    for row in rows:
        search = escape(" ".join(str(value) for value in row.values()).lower())
        table_rows.append(
            f"""<tr data-group="{escape(str(row['group']))}" data-search="{search}">
<td><span class="status">{escape(group_labels.get(str(row['group']), str(row['group'])))}</span></td>
<td><b>{escape(str(row['recommendation']))}</b><br><small>{escape(str(row['reason']))}</small></td>
<td><a href="https://www.ozon.ru/product/{escape(str(row['sku']))}/">{escape(str(row['offer_id']))}</a><br><small>SKU {escape(str(row['sku']))}</small></td>
<td>{escape(str(row['title']))}</td>
<td>{row['created_date']}<br><small>{row['card_age_days']} дней</small></td>
<td class="num">{row['stock_products']}<br><small>{row['stock_physical_pieces']} изделий</small></td>
<td class="num">{row['buyouts_total']}<br><small>последний: {escape(str(row['last_buyout_date'] or 'не было'))}</small></td>
<td class="num">{row['orders_30d']}<br><small>в доставке {row['fbo_delivering_30d']} / отменено {row['fbo_cancelled_30d']}</small></td>
<td class="num">{row['cpc_spend_30d']:.2f} ₽<br><small>{row['cpc_clicks_30d']} кликов / {row['cpc_orders_30d']} заказов</small></td>
<td class="num">{row['parser_best_position'] or '—'}<br><small>{row['parser_query_count']} запросов</small></td>
<td class="num">{row['current_buyer_price']:.0f} ₽<br><small>min {row['current_min_price']:.0f} ₽</small></td>
</tr>"""
        )
    options = "\n".join(
        f'<option value="{escape(group)}">{escape(group_labels.get(group, group))} ({count})</option>'
        for group, count in sorted(metrics["groups"].items())
    )
    parser_endpoint = escape(str(parser_meta.get("endpoint", "не подтверждено")))
    parser_date = escape(str(parser_meta.get("date", "не подтверждено")))
    html = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Ozon: товары без выкупов</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;background:#f4f6f7;color:#18201d;font:14px/1.45 Arial,sans-serif;letter-spacing:0}}
main{{max-width:1600px;margin:auto;padding:18px}}h1{{font-size:25px;margin:0 0 4px}}h2{{font-size:18px;margin:0 0 10px}}
.meta,small{{color:#63706a}}.band{{background:#fff;border:1px solid #d8dfdc;border-radius:6px;padding:14px;margin:12px 0}}
.metrics{{display:grid;grid-template-columns:repeat(5,minmax(0,1fr));gap:9px}}.metric{{border-left:4px solid #28624e;background:#f7faf8;padding:10px;min-width:0}}
.metric b{{display:block;font-size:22px}}.warn{{border-left-color:#a25d12}}.danger{{border-left-color:#a53b35}}
.controls{{display:grid;grid-template-columns:2fr 1fr;gap:8px}}input,select{{width:100%;padding:9px;border:1px solid #aab6b0;border-radius:4px;background:#fff}}
.table{{overflow:auto;max-height:72vh;border:1px solid #d8dfdc}}table{{border-collapse:collapse;width:100%;min-width:1500px;background:#fff}}
th,td{{padding:8px;border-bottom:1px solid #e2e7e5;text-align:left;vertical-align:top}}th{{position:sticky;top:0;background:#e9efec;z-index:1;font-size:12px}}
.num{{text-align:right;white-space:nowrap}}.status{{font-size:12px;font-weight:bold}}a{{color:#0b5d49}}ul{{margin:8px 0;padding-left:20px}}
@media(max-width:800px){{main{{padding:10px}}h1{{font-size:21px}}.metrics{{grid-template-columns:repeat(2,minmax(0,1fr))}}.controls{{grid-template-columns:1fr}}}}
</style></head><body><main>
<h1>Ozon: товары без выкупов</h1>
<div class="meta">{escape(run_id)} · read-only · история {history_start}–{end} · собрано {escape(collected_at)}</div>
<section class="band"><p><b>Исходные 138 SKU означали отсутствие выкупов только за 30 дней.</b> Этот отчёт проверяет всю доступную историю с даты создания карточки и отделяет никогда не продававшиеся товары от временно остановившихся.</p></section>
<section class="metrics">
<div class="metric"><span>Исходный список</span><b>{metrics['candidate_skus']}</b></div>
<div class="metric danger"><span>Никогда не продавались</span><b>{metrics['never_sold_skus']}</b></div>
<div class="metric warn"><span>Без выкупа 90+ дней</span><b>{metrics['decision_priority_skus_90d_plus']}</b></div>
<div class="metric warn"><span>Ещё доставляются</span><b>{metrics['current_delivery_units']}</b><span>{metrics['current_delivery_skus']} SKU</span></div>
<div class="metric"><span>Остаток</span><b>{metrics['current_stock_products']}</b><span>{metrics['current_stock_physical_pieces']} изделий</span></div>
</section>
<section class="band"><h2>Как принимать решения</h2><ul>
<li>Никогда не продавались 180+ дней: не пополнять, выбирать распродажу или вывод.</li>
<li>Раньше продавались: учитывать исторический спрос; сильные товары сначала перезапустить.</li>
<li>CPC без заказов: убрать платный трафик до исправления карточки.</li>
<li>Отсутствие в парсере относится только к 30 запросам, регион Ярославль, глубина 500.</li>
</ul></section>
<section class="band"><div class="controls"><input id="search" type="search" placeholder="Поиск по артикулу, названию или рекомендации">
<select id="group"><option value="">Все группы</option>{options}</select></div></section>
<div class="table"><table><thead><tr><th>Группа</th><th>Рекомендация</th><th>Артикул</th><th>Товар</th><th>Карточка создана</th><th>Остаток</th><th>Выкупы за историю</th><th>Заказы 30д</th><th>CPC 30д</th><th>Парсер</th><th>Цена покупателя</th></tr></thead>
<tbody>{''.join(table_rows)}</tbody></table></div>
<section class="band"><h2>Источники и ограничения</h2><p>Ozon Seller API: карточки, цены, FBO-остатки, аналитика и завершённые финансовые операции. Ozon Performance API: CPC 30 дней. Parser Data API: {parser_endpoint}, дата {parser_date}. Дата создания карточки не равна дате первой поставки на склад. Изменений в Ozon не выполнялось.</p></section>
</main><script>
const q=document.getElementById('search'),g=document.getElementById('group');
function filter(){{const text=q.value.toLowerCase(),group=g.value;document.querySelectorAll('tbody tr').forEach(r=>r.hidden=!(r.dataset.search.includes(text)&&(!group||r.dataset.group===group)));}}
q.addEventListener('input',filter);g.addEventListener('change',filter);
</script></body></html>"""
    path.write_text(html, encoding="utf-8")


def main() -> int:
    args = _args()
    started = datetime.now(MOSCOW)
    data_dir = Path(args.data_dir)
    run_id = args.run_id or f"ozon_dormant_inventory_{started:%Y%m%dT%H%M%S}"
    run_dir = data_dir / "runs" / started.date().isoformat() / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    candidates_path = Path(args.candidates)
    candidates = _read_csv(candidates_path)
    candidate_skus = {str(row.get("sku") or "") for row in candidates}
    history_start = date.fromisoformat(args.history_from)
    end = date.fromisoformat(args.date_to)
    analytics_start = end - timedelta(days=29)

    credentials = load_credentials()
    if credentials.ozon_seller is None:
        raise RuntimeError("Ozon Seller API credentials are missing")
    adapter = OzonSellerAdapter(credentials.ozon_seller)

    info_by_sku, listed = _product_index(adapter, candidate_skus)
    history_by_sku = _finance_history(
        adapter,
        start=history_start,
        end=end,
        candidate_skus=candidate_skus,
    )
    analytics_by_sku = _analytics_30d(adapter, start=analytics_start, end=end)
    postings_by_sku = _fbo_statuses_30d(
        adapter,
        start=analytics_start,
        end=end,
        candidate_skus=candidate_skus,
    )
    cpc_by_sku = _cpc_index(Path(args.cpc_csv))
    parser_by_sku, parser_meta = _parser_positions(
        skus=sorted(candidate_skus),
        env_path=Path(args.parser_env),
        run_date=end,
    )
    offers = [str(row.get("offer_id") or "") for row in candidates]
    price_rows = adapter.fetch_product_info_prices_by_offer_ids(offers)
    price_by_offer = _price_index(price_rows)

    rows = _build_rows(
        candidates=candidates,
        info_by_sku=info_by_sku,
        price_by_offer=price_by_offer,
        history_by_sku=history_by_sku,
        analytics_by_sku=analytics_by_sku,
        postings_by_sku=postings_by_sku,
        cpc_by_sku=cpc_by_sku,
        parser_by_sku=parser_by_sku,
        history_start=history_start,
        end=end,
    )
    metrics = _summary(rows)
    collected_at = datetime.now(MOSCOW).isoformat(timespec="seconds")

    csv_path = run_dir / "dormant_inventory.csv"
    xlsx_path = run_dir / "dormant_inventory.xlsx"
    html_path = run_dir / "report.html"
    md_path = run_dir / "report.md"
    summary_path = run_dir / "summary.json"
    _write_csv(csv_path, rows)
    _write_xlsx(xlsx_path, rows, metrics)
    _write_html(
        html_path,
        run_id=run_id,
        collected_at=collected_at,
        history_start=history_start,
        end=end,
        rows=rows,
        metrics=metrics,
        parser_meta=parser_meta,
    )
    _write_markdown(
        md_path,
        run_id=run_id,
        collected_at=collected_at,
        history_start=history_start,
        end=end,
        metrics=metrics,
        parser_meta=parser_meta,
    )
    summary = {
        "run_id": run_id,
        "started_at": started.isoformat(timespec="seconds"),
        "finished_at": collected_at,
        "overall_status": "ok",
        "mode": "read_only",
        "marketplace": "ozon",
        "history_period": {
            "date_from": history_start.isoformat(),
            "date_to": end.isoformat(),
        },
        "source_candidate_period": "2026-06-29..2026-07-28",
        "metrics": metrics,
        "source_coverage": {
            "candidate_info_found": len(info_by_sku),
            "candidate_info_missing": len(candidate_skus - set(info_by_sku)),
            "listed_products_scanned": len(listed),
            "cpc_rows_joined": sum(sku in cpc_by_sku for sku in candidate_skus),
            "fbo_posting_rows_joined": len(postings_by_sku),
            "parser_rows_joined": sum(sku in parser_by_sku for sku in candidate_skus),
            "parser": parser_meta,
        },
        "limitations": [
            "Product created_at is not the first warehouse receipt date.",
            "Parser visibility is limited to 30 queries, Yaroslavl and depth 500.",
            "Recommendations are read-only and require owner review before marketplace changes.",
        ],
        "writes_performed": 0,
        "artifacts": {
            "report": str(md_path),
            "html": str(html_path),
            "csv": str(csv_path),
            "xlsx": str(xlsx_path),
            "summary": str(summary_path),
        },
    }
    _write_json(summary_path, summary)
    manifest_artifacts = write_summary_run_manifest(
        data_dir=data_dir,
        run_dir=run_dir,
        summary=summary,
        task="ozon-dormant-inventory",
        mode="read_only",
        risk="low",
        marketplaces=["ozon"],
        inputs={
            "candidate_source": str(candidates_path),
            "history_from": history_start.isoformat(),
            "date_to": end.isoformat(),
            "candidate_skus": len(candidate_skus),
        },
        lifecycle_status="closed",
        closed=True,
    )
    summary["artifacts"].update(manifest_artifacts)
    _write_json(summary_path, summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
