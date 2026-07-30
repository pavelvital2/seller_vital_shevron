#!/usr/bin/env python3
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_FLOOR
from html import escape
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from seller_agent.core.run_manifest import write_summary_run_manifest


RETAINED_PRICE_SHARE = Decimal("0.4317489500")
LOGISTICS_PER_GOOD = Decimal("78.3190598291")
COST_PER_PHYSICAL_ITEM = Decimal("85")
TARGET_MARGIN_PER_PHYSICAL_ITEM = Decimal("50")


def _args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--action-run-dir", type=Path, required=True)
    parser.add_argument("--promotion-run-dir", type=Path, required=True)
    parser.add_argument("--analysis-date", default=date.today().isoformat())
    return parser.parse_args()


def _read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def _int(value: Any) -> int:
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


def _float(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _day(value: Any) -> date | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _actual_price(base_price: Decimal, discount: int) -> Decimal:
    return (base_price * (Decimal("100") - Decimal(discount)) / Decimal("100")).quantize(
        Decimal("0.01")
    )


def _minimum_safe_discount(base_price: Decimal, minimum: Decimal) -> int:
    raw = (Decimal("100") * (Decimal("1") - minimum / base_price)).to_integral_value(
        rounding=ROUND_FLOOR
    )
    discount = max(0, min(99, int(raw)))
    while discount > 0 and _actual_price(base_price, discount) < minimum:
        discount -= 1
    return discount


def _estimated_margin(price: Decimal, pack_qty: int) -> tuple[Decimal, Decimal]:
    total = (
        price * RETAINED_PRICE_SHARE
        - LOGISTICS_PER_GOOD
        - COST_PER_PHYSICAL_ITEM * Decimal(pack_qty)
    )
    return total.quantize(Decimal("0.01")), (
        total / Decimal(pack_qty)
    ).quantize(Decimal("0.01"))


def _promo_rows(action_run_dir: Path) -> dict[int, list[dict[str, Any]]]:
    snapshot = _read_json(action_run_dir / "raw" / "cabinet-actions-snapshot.json")
    promos = {int(row["actionID"]): row for row in snapshot.get("promos") or []}
    result: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for path in sorted((action_run_dir / "raw" / "excel").glob("*.xlsx")):
        match = re.search(r"promo-(\d+)-", path.name)
        if not match:
            continue
        action_id = int(match.group(1))
        # WB exports can contain a stale worksheet dimension that hides data in
        # openpyxl read-only mode. Normal mode reads the actual populated rows.
        workbook = load_workbook(path, read_only=False, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        index = {
            re.sub(r"\s+", " ", str(name or "").replace("\u00a0", " ")).strip(): pos
            for pos, name in enumerate(header)
        }
        for source in rows:
            if source[index["Артикул WB"]] in (None, ""):
                continue
            nm_id = int(source[index["Артикул WB"]])
            base = Decimal(str(source[index["Текущая розничная цена"]]))
            plan = Decimal(str(source[index["Плановая цена для акции"]]))
            discount = int(source[index["Загружаемая скидка для участия в акции"]])
            result[nm_id].append(
                {
                    "action_id": action_id,
                    "action_name": str((promos.get(action_id) or {}).get("name") or ""),
                    "plan_price": float(plan),
                    "required_discount": discount,
                    "actual_price": float(_actual_price(base, discount)),
                    "currently_participates": str(
                        source[index["Товар уже участвует в акции"]] or ""
                    ).strip().lower()
                    == "да",
                }
            )
        workbook.close()
    return result


def _promotion_sources(
    promotion_run_dir: Path,
) -> tuple[dict[int, list[dict[str, Any]]], dict[int, dict[str, Any]]]:
    active: dict[int, list[dict[str, Any]]] = defaultdict(list)
    campaigns = _read_json(promotion_run_dir / "raw" / "campaigns.json").get("adverts") or []
    for campaign in campaigns:
        if _int(campaign.get("status")) != 9:
            continue
        settings = campaign.get("settings") or {}
        for item in campaign.get("nm_settings") or []:
            active[_int(item.get("nm_id"))].append(
                {
                    "advert_id": _int(campaign.get("id")),
                    "campaign_name": str(settings.get("name") or ""),
                    "payment_type": str(settings.get("payment_type") or ""),
                    "search_bid": _float((item.get("bids_kopecks") or {}).get("search")) / 100,
                }
            )

    metrics: dict[int, dict[str, Any]] = defaultdict(
        lambda: {
            "views": 0,
            "clicks": 0,
            "atbs": 0,
            "orders": 0,
            "spend": 0.0,
            "revenue": 0.0,
        }
    )
    with (promotion_run_dir / "wb_promotion_products.csv").open(
        encoding="utf-8-sig", newline=""
    ) as source:
        for row in csv.DictReader(source):
            target = metrics[_int(row.get("nm_id"))]
            for field in ("views", "clicks", "atbs", "orders"):
                target[field] += _int(row.get(field))
            for field in ("spend", "revenue"):
                target[field] = round(target[field] + _float(row.get(field)), 2)
    return active, metrics


def _build_rows(
    *,
    run_dir: Path,
    action_run_dir: Path,
    promotion_run_dir: Path,
    analysis_date: date,
) -> list[dict[str, Any]]:
    sales = _read_json(run_dir / "raw" / "sales_history.json")
    orders = _read_json(run_dir / "raw" / "orders_30d.json")
    stocks = _read_json(run_dir / "raw" / "warehouse_stocks.json")
    cards = _read_json(run_dir / "raw" / "cards.json")
    action_report = _read_json(action_run_dir / "report.json")
    parser_latest = {
        _int(row.get("nm_id")): row
        for row in _read_json(run_dir / "processed" / "parser_latest_visibility_all_store.json")
    }
    offers_by_nm = _promo_rows(action_run_dir)
    active_ads, ad_metrics = _promotion_sources(promotion_run_dir)

    returned_srids = {
        str(row.get("srid") or "")
        for row in sales
        if str(row.get("saleID") or "").startswith("R")
    }
    completed_rows = [
        row
        for row in sales
        if str(row.get("saleID") or "").startswith("S")
        and str(row.get("srid") or "") not in returned_srids
    ]
    completed_srids = {str(row.get("srid") or "") for row in completed_rows}
    completed_by_nm: dict[int, list[dict[str, Any]]] = defaultdict(list)
    returns_by_nm: Counter[int] = Counter()
    for row in completed_rows:
        completed_by_nm[_int(row.get("nmId"))].append(row)
    for row in sales:
        if str(row.get("saleID") or "").startswith("R"):
            returns_by_nm[_int(row.get("nmId"))] += 1

    order_signals: dict[int, Counter[str]] = defaultdict(Counter)
    for row in orders:
        nm_id = _int(row.get("nmId"))
        order_signals[nm_id]["orders"] += 1
        if bool(row.get("isCancel")):
            order_signals[nm_id]["cancelled"] += 1
        elif str(row.get("srid") or "") in completed_srids:
            order_signals[nm_id]["completed"] += 1
        else:
            order_signals[nm_id]["open"] += 1

    stock_by_nm: Counter[int] = Counter()
    warehouses_by_nm: dict[int, Counter[str]] = defaultdict(Counter)
    for row in stocks:
        nm_id = _int(row.get("nmId"))
        quantity = _int(row.get("quantity"))
        stock_by_nm[nm_id] += quantity
        warehouses_by_nm[nm_id][str(row.get("warehouseName") or "")] += quantity

    card_by_nm = {_int(row.get("nmID")): row for row in cards}
    action_products = {
        _int(row.get("nm_id")): row for row in action_report.get("products") or []
    }
    cutoff = analysis_date - timedelta(days=30)
    rows: list[dict[str, Any]] = []

    for nm_id, product in action_products.items():
        stock_goods = stock_by_nm[nm_id]
        if stock_goods <= 0:
            continue
        completed = completed_by_nm[nm_id]
        completed_days = [_day(row.get("date")) for row in completed]
        completed_days = [value for value in completed_days if value is not None]
        if any(value >= cutoff for value in completed_days):
            continue

        last_sale = max(completed_days) if completed_days else None
        days_without_sale = (analysis_date - last_sale).days if last_sale else 999
        card = card_by_nm.get(nm_id) or {}
        created = _day(card.get("createdAt"))
        card_age = (analysis_date - created).days if created else 0
        pack_qty = max(1, _int(product.get("pack_qty")))
        title = str(product.get("title") or card.get("title") or "")
        internal_sku = str(product.get("internal_sku") or "")
        callsign = bool(re.search(r"позывн", title, re.IGNORECASE) or "_pz_" in internal_sku)
        order_signal = order_signals[nm_id]
        ads = active_ads.get(nm_id) or []
        metric = ad_metrics[nm_id]
        parser = parser_latest.get(nm_id) or {}
        offers = offers_by_nm.get(nm_id) or []
        best_offer = max(offers, key=lambda row: row["actual_price"], default=None)

        minimum = Decimal(str(product.get("minimum") or 0))
        base_price = Decimal(str(product.get("base_price") or 0))
        safe_discount = _minimum_safe_discount(base_price, minimum)
        safe_price = _actual_price(base_price, safe_discount)
        safe_margin_total, safe_margin_piece = _estimated_margin(safe_price, pack_qty)
        best_action_price = Decimal(str((best_offer or {}).get("actual_price") or 0))
        action_margin_total, action_margin_piece = (
            _estimated_margin(best_action_price, pack_qty)
            if best_action_price > 0
            else (Decimal("0"), Decimal("0"))
        )

        if order_signal["open"] > 0:
            group = "pending_orders"
            recommendation = "Не выводить: дождаться текущих заказов, цену снизить до safe-level"
            deadline = "3 дня на завершение заказов, затем 14 дней контроля"
        elif callsign:
            group = "callsign_clearance"
            recommendation = "Убрать CPC, снизить цену до safe-level, производство остановить"
            deadline = "14 дней; без выкупа перейти к нулевой марже или выводу"
        elif len(completed) >= 3 and days_without_sale < 90:
            group = "proven_relaunch"
            recommendation = "Перезапуск: safe-level цены и заметное усиление CPC"
            deadline = "14 дней; стоп рекламы после 10 кликов или 20 ₽ без заказа"
        elif len(completed) >= 3:
            group = "historical_clearance"
            recommendation = "Распродажа с контролем: safe-level цены и короткий CPC-тест"
            deadline = "14 дней; без заказа перейти к нулевой марже"
        else:
            group = "weak_clearance"
            recommendation = "Распродажа: safe-level цены; CPC не повышать без видимости/корзин"
            deadline = "14 дней; без заказа нулевая маржа, затем вывод"

        if group == "proven_relaunch":
            cpc_target = 3.0 if len(completed) >= 6 else 2.5
            cpc_rule = f"{cpc_target:.2f} ₽; лимит 20 ₽ или 10 кликов без заказа"
        elif group == "historical_clearance":
            cpc_target = 2.5
            cpc_rule = "2.50 ₽; лимит 15 ₽ без заказа"
        elif group == "weak_clearance" and (
            _int(parser.get("query_count")) > 0 or metric["atbs"] > 0
        ):
            cpc_target = 2.0
            cpc_rule = "2.00 ₽; лимит 15 ₽ без заказа"
        elif group == "pending_orders" and not callsign:
            cpc_target = max([_float(row.get("search_bid")) for row in ads] or [0])
            cpc_rule = "оставить текущую ставку до завершения открытых заказов"
        else:
            cpc_target = 0.0
            cpc_rule = "не повышать" if not callsign else "убрать из CPC"

        current_bid = max([_float(row.get("search_bid")) for row in ads] or [0])
        warehouse_text = ", ".join(
            f"{name}: {quantity}"
            for name, quantity in warehouses_by_nm[nm_id].most_common()
            if quantity > 0
        )
        rows.append(
            {
                "group": group,
                "recommendation": recommendation,
                "deadline": deadline,
                "nm_id": nm_id,
                "vendor_code": str(product.get("vendor_code") or ""),
                "internal_sku": internal_sku,
                "title": title,
                "pack_qty": pack_qty,
                "stock_goods": stock_goods,
                "stock_physical_items": stock_goods * pack_qty,
                "warehouses": warehouse_text,
                "card_created": created.isoformat() if created else "",
                "card_age_days": card_age,
                "completed_sales_history": len(completed),
                "completed_sales_60d": sum(
                    value >= analysis_date - timedelta(days=60) for value in completed_days
                ),
                "completed_sales_90d": sum(
                    value >= analysis_date - timedelta(days=90) for value in completed_days
                ),
                "completed_sales_180d": len(completed),
                "returns_history": returns_by_nm[nm_id],
                "last_completed_sale": last_sale.isoformat() if last_sale else "",
                "days_without_completed_sale": days_without_sale,
                "orders_30d": order_signal["orders"],
                "open_orders_30d": order_signal["open"],
                "cancelled_orders_30d": order_signal["cancelled"],
                "base_price": float(base_price),
                "current_discount": _int(product.get("current_discount")),
                "current_price": _float(product.get("current_price")),
                "minimum_price": float(minimum),
                "safe_discount": safe_discount,
                "safe_price": float(safe_price),
                "safe_margin_total_estimate": float(safe_margin_total),
                "safe_margin_per_item_estimate": float(safe_margin_piece),
                "currently_in_action": bool(product.get("currently_participates")),
                "offered_actions": len(offers),
                "best_action_name": str((best_offer or {}).get("action_name") or ""),
                "best_action_discount": _int((best_offer or {}).get("required_discount")),
                "best_action_price": float(best_action_price),
                "best_action_margin_total_estimate": float(action_margin_total),
                "best_action_margin_per_item_estimate": float(action_margin_piece),
                "best_action_above_current_minimum": bool(
                    best_action_price and best_action_price >= minimum
                ),
                "best_action_above_zero_margin": bool(
                    best_action_price and action_margin_total >= 0
                ),
                "active_cpc": any(row.get("payment_type") == "cpc" for row in ads),
                "active_campaigns": "; ".join(
                    f"{row['advert_id']} {row['campaign_name']}" for row in ads
                ),
                "current_bid": round(current_bid, 2),
                "recommended_bid": round(cpc_target, 2),
                "cpc_rule": cpc_rule,
                "ad_views_30d": metric["views"],
                "ad_clicks_30d": metric["clicks"],
                "ad_atbs_30d": metric["atbs"],
                "ad_orders_30d": metric["orders"],
                "ad_spend_30d": metric["spend"],
                "ad_revenue_30d": metric["revenue"],
                "parser_visible_latest": bool(parser),
                "parser_query_count": _int(parser.get("query_count")),
                "parser_best_position": _int(parser.get("best_position")),
                "parser_rating": _float(parser.get("rating")),
                "parser_feedbacks": _int(parser.get("feedbacks")),
                "parser_queries": ", ".join(parser.get("sample_queries") or []),
                "is_callsign": callsign,
            }
        )

    priority = {
        "pending_orders": 0,
        "proven_relaunch": 1,
        "historical_clearance": 2,
        "weak_clearance": 3,
        "callsign_clearance": 4,
    }
    return sorted(
        rows,
        key=lambda row: (
            priority[row["group"]],
            -row["stock_physical_items"],
            row["nm_id"],
        ),
    )


def _summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    groups = Counter(row["group"] for row in rows)
    group_goods = Counter()
    group_physical = Counter()
    for row in rows:
        group_goods[row["group"]] += row["stock_goods"]
        group_physical[row["group"]] += row["stock_physical_items"]
    offered = [row for row in rows if row["best_action_price"] > 0]
    return {
        "dormant_cards": len(rows),
        "stock_goods": sum(row["stock_goods"] for row in rows),
        "stock_physical_items": sum(row["stock_physical_items"] for row in rows),
        "never_sold_cards": sum(row["completed_sales_history"] == 0 for row in rows),
        "no_sale_90d_plus_cards": sum(
            row["days_without_completed_sale"] >= 90 for row in rows
        ),
        "open_order_cards": sum(row["open_orders_30d"] > 0 for row in rows),
        "open_orders": sum(row["open_orders_30d"] for row in rows),
        "active_cpc_cards": sum(row["active_cpc"] for row in rows),
        "ad_spend_30d": round(sum(row["ad_spend_30d"] for row in rows), 2),
        "ad_clicks_30d": sum(row["ad_clicks_30d"] for row in rows),
        "ad_orders_30d": sum(row["ad_orders_30d"] for row in rows),
        "parser_visible_latest": sum(row["parser_visible_latest"] for row in rows),
        "phase1_price_changes": sum(
            not row["currently_in_action"] and row["current_price"] != row["safe_price"]
            for row in rows
        ),
        "keep_current_action": sum(
            row["currently_in_action"] and row["best_action_above_current_minimum"]
            for row in rows
        ),
        "stage2_action_positive_margin": sum(
            row["best_action_price"] > 0
            and not row["best_action_above_current_minimum"]
            and row["best_action_above_zero_margin"]
            for row in rows
        ),
        "action_below_cost": sum(
            row["best_action_price"] > 0 and not row["best_action_above_zero_margin"]
            for row in rows
        ),
        "not_offered_action": len(rows) - len(offered),
        "remove_callsigns_from_cpc": sum(
            row["is_callsign"] and row["active_cpc"] for row in rows
        ),
        "groups": dict(groups),
        "group_goods": dict(group_goods),
        "group_physical": dict(group_physical),
    }


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=list(rows[0]) if rows else [])
        writer.writeheader()
        writer.writerows(rows)


def _write_xlsx(path: Path, rows: list[dict[str, Any]], summary: dict[str, Any]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Решения"
    headers = list(rows[0]) if rows else []
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(field, "") for field in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5D50")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
            max(width + 2, 11), 55
        )
    summary_sheet = workbook.create_sheet("Сводка")
    summary_sheet.append(["Показатель", "Значение"])
    for key, value in summary.items():
        summary_sheet.append(
            [key, json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value]
        )
    summary_sheet.column_dimensions["A"].width = 42
    summary_sheet.column_dimensions["B"].width = 80
    workbook.save(path)


def _write_markdown(
    path: Path,
    *,
    run_id: str,
    analysis_date: date,
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    parser_meta: dict[str, Any],
) -> None:
    lines = [
        "# WB: залежалые товары",
        "",
        f"Run ID: `{run_id}`. Режим: `read-only`. Дата анализа: `{analysis_date}`.",
        "",
        "## Итог",
        "",
        f"- без завершённого выкупа 30 дней: `{summary['dormant_cards']}` карточки;",
        f"- остаток: `{summary['stock_goods']}` товаров / `{summary['stock_physical_items']}` изделий;",
        f"- никогда не продавались в доступной истории: `{summary['never_sold_cards']}`;",
        f"- без выкупа 90+ дней: `{summary['no_sale_90d_plus_cards']}`;",
        f"- открытые заказы: `{summary['open_orders']}` по `{summary['open_order_cards']}` карточкам;",
        f"- в CPC: `{summary['active_cpc_cards']}`, расход `{summary['ad_spend_30d']}` руб., "
        f"клики `{summary['ad_clicks_30d']}`, рекламные заказы `{summary['ad_orders_30d']}`;",
        f"- видны в свежем parser-срезе: `{summary['parser_visible_latest']}`.",
        "",
        "## Рекомендация",
        "",
        "1. Остановить производство и пополнение всех карточек из списка до выхода из хвоста.",
        f"2. Для `{summary['phase1_price_changes']}` карточек вне акции снизить обычную цену "
        "до ближайшей цены не ниже действующего минимума: 507 руб. для 1 изделия, "
        "819 руб. для 2 изделий, 1 120 руб. для 3 изделий.",
        f"3. Сохранить текущую допустимую акцию для `{summary['keep_current_action']}` карточек.",
        f"4. Убрать из CPC `{summary['remove_callsigns_from_cpc']}` залежалых позывных.",
        "5. Для доказанно продававшихся товаров провести 14-дневный CPC-тест со ставкой "
        "2,50-3,00 руб. и стопом 10 кликов/20 руб. без заказа.",
        f"6. Если за 14 дней нет заказа: `{summary['stage2_action_positive_margin']}` карточек "
        "можно перевести в доступную акцию ниже текущего минимума, но ещё без убытка.",
        f"7. Не включать автоматически `{summary['action_below_cost']}` карточек: предложенная "
        "WB цена акции уже ниже себестоимости с логистикой. Для них второй шаг - ручная "
        "цена нулевой маржи; затем вывод остатка, если продажи не появились.",
        "",
        "## Группы",
        "",
    ]
    labels = {
        "pending_orders": "Есть открытые заказы",
        "proven_relaunch": "Доказанный спрос, перезапуск",
        "historical_clearance": "Исторический спрос, распродажа",
        "weak_clearance": "Слабый/неподтверждённый спрос",
        "callsign_clearance": "Позывные без массового CPC",
    }
    for group, count in summary["groups"].items():
        lines.append(
            f"- {labels.get(group, group)}: `{count}` карточек, "
            f"`{summary['group_goods'].get(group, 0)}` товаров, "
            f"`{summary['group_physical'].get(group, 0)}` изделий."
        )
    lines.extend(
        [
            "",
            "## Источники и ограничения",
            "",
            "- WB Statistics API: доступная история завершённых продаж фактически начинается "
            "`2026-02-01`; возвраты исключены по `srid`.",
            "- WB Statistics API orders: `2026-06-30` - дата сбора; незавершённые заказы "
            "показаны отдельно и не названы выкупами.",
            "- WB Analytics API: свежий продаваемый остаток по складам.",
            "- WB Promotion API: `2026-06-30` - `2026-07-29`.",
            "- WB LK actions snapshot: две активные автоматические акции; сравнение с "
            "действующими минимальными ценами.",
            f"- Parser Data API: `{parser_meta.get('endpoint')}`, дата "
            f"`{parser_meta.get('run_date_max')}`, регион `{', '.join(parser_meta.get('regions') or [])}`, "
            f"пакет `{', '.join(parser_meta.get('query_packs') or [])}`. Отсутствие означает "
            "невидимость только по отслеживаемым 31 запросам в глубине сбора.",
            "- Маржа акции рассчитана по согласованной 30-дневной модели: retained share "
            "`43.174895%`, логистика `78.32 руб.` на товар, себестоимость `85 руб.` на изделие.",
            "- Цены, скидки, акции, ставки, бюджеты и карточки не изменялись.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_html(
    path: Path,
    *,
    run_id: str,
    analysis_date: date,
    rows: list[dict[str, Any]],
    summary: dict[str, Any],
    parser_meta: dict[str, Any],
) -> None:
    labels = {
        "pending_orders": "Есть открытые заказы",
        "proven_relaunch": "Доказанный спрос",
        "historical_clearance": "Исторический спрос",
        "weak_clearance": "Слабый спрос",
        "callsign_clearance": "Позывные",
    }
    group_options = "".join(
        f'<option value="{escape(group)}">{escape(labels.get(group, group))} ({count})</option>'
        for group, count in summary["groups"].items()
    )
    table_rows = []
    for row in rows:
        searchable = escape(
            " ".join(str(value) for value in row.values()).lower(), quote=True
        )
        action_text = (
            f"{row['best_action_price']:.0f} ₽ / {row['best_action_discount']}%"
            if row["best_action_price"]
            else "не предложена"
        )
        table_rows.append(
            f"""<tr data-group="{escape(row['group'])}" data-search="{searchable}">
<td><b>{escape(labels.get(row['group'], row['group']))}</b><br><small>{escape(row['deadline'])}</small></td>
<td><a href="https://www.wildberries.ru/catalog/{row['nm_id']}/detail.aspx">{row['nm_id']}</a><br><small>{escape(row['vendor_code'])}</small></td>
<td><b>{escape(row['title'])}</b><br><small>{escape(row['warehouses'])}</small></td>
<td class="num">{row['stock_goods']}<br><small>{row['stock_physical_items']} изделий</small></td>
<td class="num">{row['completed_sales_history']}<br><small>{escape(row['last_completed_sale'] or 'не было')} · {row['days_without_completed_sale'] if row['days_without_completed_sale'] < 999 else 'вся история'} дн.</small></td>
<td class="num">{row['open_orders_30d']}<br><small>отменено {row['cancelled_orders_30d']}</small></td>
<td class="num">{row['current_price']:.0f} → <b>{row['safe_price']:.0f} ₽</b><br><small>min {row['minimum_price']:.0f} ₽ · скидка {row['safe_discount']}%</small></td>
<td class="num">{action_text}<br><small>маржа/изд. {row['best_action_margin_per_item_estimate']:.0f} ₽</small></td>
<td class="num">{row['current_bid']:.2f} → <b>{row['recommended_bid']:.2f} ₽</b><br><small>{escape(row['cpc_rule'])}</small></td>
<td class="num">{row['ad_spend_30d']:.2f} ₽<br><small>{row['ad_clicks_30d']} кликов · {row['ad_orders_30d']} заказов</small></td>
<td class="num">{row['parser_best_position'] or '—'}<br><small>{row['parser_query_count']} запросов</small></td>
<td>{escape(row['recommendation'])}</td>
</tr>"""
        )
    html = f"""<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>WB: залежалые товары</title>
<style>
*{{box-sizing:border-box}}body{{margin:0;background:#f3f5f4;color:#17201c;font:14px/1.45 Arial,sans-serif;letter-spacing:0}}
main{{max-width:1660px;margin:auto;padding:16px}}h1{{font-size:25px;margin:0}}h2{{font-size:18px;margin:0 0 10px}}small,.meta{{color:#66716c}}
.band{{background:#fff;border:1px solid #d6ddda;border-radius:6px;padding:14px;margin:12px 0}}
.metrics{{display:grid;grid-template-columns:repeat(6,minmax(0,1fr));gap:8px}}.metric{{background:#fff;border:1px solid #d6ddda;border-left:4px solid #2f6654;padding:10px;min-width:0}}
.metric b{{display:block;font-size:22px}}.danger{{border-left-color:#9b3832}}.warn{{border-left-color:#a36a18}}
.steps{{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:10px}}.step{{border-left:4px solid #2f6654;padding:10px;background:#f7faf8}}
.controls{{display:grid;grid-template-columns:2fr 1fr;gap:8px}}input,select{{width:100%;padding:9px;border:1px solid #aab5b0;border-radius:4px;background:#fff}}
.table{{overflow:auto;max-height:74vh;border:1px solid #d6ddda;background:#fff}}table{{border-collapse:collapse;width:100%;min-width:1900px}}
th,td{{padding:8px;border-bottom:1px solid #e1e6e4;text-align:left;vertical-align:top}}th{{position:sticky;top:0;background:#e7eeeb;z-index:1;font-size:12px}}
.num{{text-align:right;white-space:nowrap}}a{{color:#075e48}}ul{{margin:6px 0;padding-left:20px}}
@media(max-width:900px){{main{{padding:9px}}h1{{font-size:21px}}.metrics{{grid-template-columns:repeat(2,minmax(0,1fr))}}.steps{{grid-template-columns:1fr}}.controls{{grid-template-columns:1fr}}}}
</style></head><body><main>
<h1>WB: залежалые товары</h1><div class="meta">{escape(run_id)} · read-only · дата анализа {analysis_date}</div>
<section class="band"><b>Это не список «удалить всё».</b> Завершённые выкупы сверены с возвратами, открытые заказы вынесены отдельно. Решение строится по остатку, истории спроса, цене, акции, CPC и свежей поисковой видимости.</section>
<section class="metrics">
<div class="metric danger"><span>Без выкупа 30 дней</span><b>{summary['dormant_cards']}</b></div>
<div class="metric danger"><span>Остаток</span><b>{summary['stock_goods']}</b><span>{summary['stock_physical_items']} изделий</span></div>
<div class="metric warn"><span>Никогда не продавались</span><b>{summary['never_sold_cards']}</b></div>
<div class="metric"><span>Есть открытые заказы</span><b>{summary['open_order_cards']}</b><span>{summary['open_orders']} заказов</span></div>
<div class="metric warn"><span>В свежем поиске</span><b>{summary['parser_visible_latest']}</b><span>из {summary['dormant_cards']}</span></div>
<div class="metric"><span>CPC 30 дней</span><b>{summary['ad_spend_30d']:.0f} ₽</b><span>{summary['ad_clicks_30d']} кликов</span></div>
</section>
<section class="band"><h2>Решение в три шага</h2><div class="steps">
<div class="step"><b>Сейчас</b><p>Остановить производство хвоста. Для {summary['phase1_price_changes']} карточек снизить обычную цену до ближайшей цены не ниже минимума. Сохранить {summary['keep_current_action']} уже допустимые акции.</p></div>
<div class="step"><b>14 дней</b><p>Доказанный спрос: CPC 2,50–3,00 ₽, стоп 10 кликов/20 ₽ без заказа. Позывные: убрать CPC. Слабый спрос без видимости: сначала цена и карточка, не разгонять ставки вслепую.</p></div>
<div class="step"><b>После 14 дней</b><p>{summary['stage2_action_positive_margin']} карточек можно опустить до цены акции без убытка. Для {summary['action_below_cost']} карточек акция уже убыточна: цена нулевой маржи, затем вывод остатка.</p></div>
</div></section>
<section class="band"><p><b>Почему не включать все товары в акции сразу:</b> у {summary['action_below_cost']} карточек лучшая доступная цена акции уже ниже себестоимости с логистикой. Это не ускорение продаж, а автоматическая фиксация убытка до рекламных расходов.</p></section>
<section class="band"><div class="controls"><input id="search" type="search" placeholder="Поиск по nmID, артикулу, названию, складу"><select id="group"><option value="">Все группы</option>{group_options}</select></div></section>
<div class="table"><table><thead><tr><th>Группа</th><th>WB</th><th>Товар / склады</th><th>Остаток</th><th>Выкупы</th><th>Открытые заказы</th><th>Цена сейчас → safe</th><th>Лучшая акция</th><th>CPC сейчас → рекомендовано</th><th>Реклама 30д</th><th>Парсер</th><th>Решение</th></tr></thead><tbody>{''.join(table_rows)}</tbody></table></div>
<section class="band"><h2>Источники и ограничения</h2><ul>
<li>История WB Statistics API фактически доступна с 01.02.2026; возвраты исключены по srid.</li>
<li>Promotion API: 30.06–29.07.2026. Остатки и карточки собраны 30.07.2026.</li>
<li>Parser Data API: {escape(str(parser_meta.get('run_date_max')))}, регион {escape(', '.join(parser_meta.get('regions') or []))}, пакет {escape(', '.join(parser_meta.get('query_packs') or []))}. Это 31 отслеживаемый запрос, а не весь поиск WB.</li>
<li>Экономика: 43,174895% цены остаётся после переменных расходов, логистика 78,32 ₽ на товар, себестоимость 85 ₽ на изделие. Налог не включён.</li>
<li>Никаких изменений на WB не выполнялось.</li>
</ul></section>
</main><script>
const q=document.getElementById('search'),g=document.getElementById('group');
function filter(){{const text=q.value.toLowerCase(),group=g.value;document.querySelectorAll('tbody tr').forEach(r=>r.hidden=!(r.dataset.search.includes(text)&&(!group||r.dataset.group===group)));}}
q.addEventListener('input',filter);g.addEventListener('change',filter);
</script></body></html>"""
    path.write_text(html, encoding="utf-8")


def main() -> int:
    args = _args()
    run_dir = args.run_dir
    analysis_date = date.fromisoformat(args.analysis_date)
    run_id = run_dir.name
    rows = _build_rows(
        run_dir=run_dir,
        action_run_dir=args.action_run_dir,
        promotion_run_dir=args.promotion_run_dir,
        analysis_date=analysis_date,
    )
    summary = _summary(rows)
    parser_meta = _read_json(run_dir / "processed" / "parser_meta.json")

    csv_path = run_dir / "dormant_inventory.csv"
    xlsx_path = run_dir / "dormant_inventory.xlsx"
    html_path = run_dir / "report.html"
    md_path = run_dir / "report.md"
    summary_path = run_dir / "summary.json"
    _write_csv(csv_path, rows)
    _write_xlsx(xlsx_path, rows, summary)
    _write_markdown(
        md_path,
        run_id=run_id,
        analysis_date=analysis_date,
        rows=rows,
        summary=summary,
        parser_meta=parser_meta,
    )
    _write_html(
        html_path,
        run_id=run_id,
        analysis_date=analysis_date,
        rows=rows,
        summary=summary,
        parser_meta=parser_meta,
    )
    result = {
        "run_id": run_id,
        "mode": "read_only",
        "overall_status": "warning" if rows else "ok",
        "analysis_date": analysis_date.isoformat(),
        "summary": summary,
        "sources": {
            "sales": str(run_dir / "raw" / "sales_history.json"),
            "orders": str(run_dir / "raw" / "orders_30d.json"),
            "stocks": str(run_dir / "raw" / "warehouse_stocks.json"),
            "cards": str(run_dir / "raw" / "cards.json"),
            "actions": str(args.action_run_dir),
            "promotion": str(args.promotion_run_dir),
            "parser": parser_meta,
        },
        "artifacts": {
            "report_html": str(html_path),
            "report_md": str(md_path),
            "csv": str(csv_path),
            "xlsx": str(xlsx_path),
            "summary": str(summary_path),
        },
        "marketplace_write_performed": False,
    }
    _write_json(summary_path, result)
    write_summary_run_manifest(
        data_dir=run_dir.parents[2],
        run_dir=run_dir,
        summary=result,
        task="wb-dormant-inventory",
        mode="read_only",
        risk="low",
        marketplaces=["wb"],
        inputs={
            "analysis_date": analysis_date.isoformat(),
            "action_run_dir": str(args.action_run_dir),
            "promotion_run_dir": str(args.promotion_run_dir),
        },
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
