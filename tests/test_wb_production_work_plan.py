from __future__ import annotations

import csv
from datetime import date
import json
from pathlib import Path

from openpyxl import load_workbook

from seller_agent.tasks.wb_production_work_plan import (
    _classify_product,
    _load_catalog,
    record_wb_work_plan_decision,
    run_wb_production_work_plan,
)


class FakeAnalyticsAdapter:
    def __init__(self, *, anomalous: bool = False) -> None:
        self.anomalous = anomalous

    def fetch_wb_warehouse_stocks(self):  # type: ignore[no-untyped-def]
        rows = [
            {"nmId": 100, "warehouseName": "Волгоград", "regionName": "Южный и Северо-Кавказский", "quantity": 1},
            {"nmId": 200, "warehouseName": "Коледино", "regionName": "Центральный", "quantity": 0},
            {"nmId": 300, "warehouseName": "Коледино", "regionName": "Центральный", "quantity": 0},
            {"nmId": 400, "warehouseName": "Коледино", "regionName": "Центральный", "quantity": 0},
        ]
        if self.anomalous:
            rows.append(
                {
                    "nmId": 200,
                    "warehouseName": "Электросталь",
                    "regionName": "Центральный",
                    "quantity": 1,
                    "inWayFromClient": 200,
                }
            )
        return rows


class FakeStatisticsAdapter:
    def fetch_sales(self, *, date_from: str, flag: int):  # type: ignore[no-untyped-def]
        rows = []
        products = [
            ("old_nr", 100, "Южный федеральный округ", "Владимир"),
            ("ng", 200, "Центральный федеральный округ", "Владимир"),
            ("kit", 300, "Центральный федеральный округ", "Владимир"),
            ("unknown", 400, "Центральный федеральный округ", "Владимир"),
        ]
        for vendor, nm_id, buyer_region, ship_from in products:
            for number in range(12):
                rows.append(
                    {
                        "date": f"2026-07-{(number % 12) + 1:02d}T10:00:00",
                        "supplierArticle": vendor,
                        "nmId": nm_id,
                        "oblastOkrugName": buyer_region,
                        "warehouseName": ship_from,
                        "saleID": f"S-{vendor}-{number}",
                        "isRealization": True,
                    }
                )
        return rows


class FakeContentAdapter:
    def fetch_cards(self, *, limit: int):  # type: ignore[no-untyped-def]
        return [
            {"vendorCode": "nr", "nmID": 100, "title": "Шеврон нарукавный", "sizes": [{"skus": ["b100"]}]},
            {"vendorCode": "ng", "nmID": 200, "title": "Шеврон нагрудный", "sizes": [{"skus": ["b200"]}]},
            {
                "vendorCode": "kit",
                "nmID": 300,
                "title": "Комплект шевронов Позывной 2 шт.",
                "sizes": [{"skus": ["b300"]}],
            },
            {"vendorCode": "unknown", "nmID": 400, "title": "Нашивка Прикол", "sizes": [{"skus": ["b400"]}]},
        ]


class FakeSuppliesAdapter:
    def fetch_supplies(self):  # type: ignore[no-untyped-def]
        return [{"supplyID": 10, "statusID": 3}]

    def fetch_supply(self, supply_id):  # type: ignore[no-untyped-def]
        return {"statusID": 3, "warehouseName": "Волгоград", "quantity": 2}

    def fetch_supply_goods(self, supply_id):  # type: ignore[no-untyped-def]
        return [{"vendorCode": "nr", "nmID": 100, "quantity": 2, "readyForSaleQuantity": 0}]


def _write_catalog(data_dir: Path) -> None:
    path = data_dir / "catalog" / "unified" / "products.csv"
    path.parent.mkdir(parents=True)
    path.write_text(
        "internal_sku,product_name,product_group,pack_qty,wb_vendor_code,wb_nm_id,wb_barcode\n"
        "chev_nr_test,Шеврон нарукавный,chev,1,nr,100,b100\n"
        "chev_ng_test,Шеврон нагрудный,chev,1,ng,200,b200\n"
        "chev_kit2_pz_test,Комплект Позывной,chev,2,kit,300,b300\n"
        "chev_prikol_test,Нашивка Прикол,chev,1,unknown,400,b400\n",
        encoding="utf-8",
    )


def _run(
    tmp_path: Path,
    *,
    mode: str,
    value: int,
    anomalous: bool = False,
    cluster_count: int = 6,
):  # type: ignore[no-untyped-def]
    data_dir = tmp_path / "data"
    _write_catalog(data_dir)
    return run_wb_production_work_plan(
        data_dir=data_dir,
        run_id=f"wb_production_work_plan_{mode}_{value}_{cluster_count}",
        mode=mode,
        value=value,
        cluster_count=cluster_count,
        today=date(2026, 7, 18),
        analytics_adapter=FakeAnalyticsAdapter(anomalous=anomalous),  # type: ignore[arg-type]
        statistics_adapter=FakeStatisticsAdapter(),  # type: ignore[arg-type]
        content_adapter=FakeContentAdapter(),  # type: ignore[arg-type]
        supplies_adapter=FakeSuppliesAdapter(),  # type: ignore[arg-type]
    )


def test_product_type_multiples_and_unknown_default() -> None:
    assert _classify_product(title="Шеврон нарукавный", internal_sku="chev_nr_x", vendor_code="x") == (
        "Нарукавный",
        8,
        True,
    )


def test_confirmed_mapping_is_catalog_fallback_for_stable_nm_id(tmp_path: Path) -> None:
    mapping = tmp_path / "confirmed.csv"
    mapping.write_text(
        "internal_sku,wb_vendor_code,wb_nm_id,wb_barcode,wb_title\n"
        "chev_kit2_pz_test,old_vendor,123,b123,Комплект Позывной 2 шт.\n",
        encoding="utf-8",
    )
    catalog = _load_catalog(tmp_path / "missing.csv", mapping)

    assert catalog["nm:123"]["internal_sku"] == "chev_kit2_pz_test"
    assert catalog["nm:123"]["pack_qty"] == "2"
    assert _classify_product(title="Шеврон нагрудный", internal_sku="chev_ng_x", vendor_code="x")[1] == 12
    assert _classify_product(title="Шеврон на спину", internal_sku="chev_back_x", vendor_code="x")[1] == 2
    assert _classify_product(title="Шеврон на кепку", internal_sku="chev_kp_x", vendor_code="x")[1] == 9
    assert _classify_product(title="Комплект Позывной", internal_sku="chev_kit2_pz_x", vendor_code="x")[1] == 6
    assert _classify_product(title="Нашивка Прикол", internal_sku="chev_x", vendor_code="x") == (
        "Тип не определен",
        8,
        False,
    )


def test_coverage_plan_uses_buyer_region_multiples_inbound_and_three_sheets(tmp_path: Path) -> None:
    result = _run(tmp_path, mode="coverage_days", value=30)

    assert result["overall_status"] == "warning"
    workbook_path = Path(result["artifacts"]["report"])
    assert workbook_path.exists()
    assert load_workbook(workbook_path, read_only=True).sheetnames == ["Артикулы", "ВБ регионы", "Контроль"]
    with Path(result["artifacts"]["rows_csv"]).open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    by_vendor = {row["vendor_code"]: row for row in rows}
    assert int(by_vendor["nr"]["quantity"]) % 8 == 0
    assert int(by_vendor["ng"]["quantity"]) % 12 == 0
    assert int(by_vendor["kit"]["quantity"]) % 6 == 0
    assert int(by_vendor["unknown"]["quantity"]) % 8 == 0
    assert by_vendor["nr"]["region"] == "Южный и Северо-Кавказский"
    assert int(by_vendor["nr"]["confirmed_inbound"]) == 2
    controls = Path(result["artifacts"]["controls_csv"]).read_text(encoding="utf-8")
    assert "unknown_product_type_default_8" in controls


def test_capacity_plan_never_exceeds_physical_capacity(tmp_path: Path) -> None:
    result = _run(tmp_path, mode="capacity", value=25)

    metrics = result["metrics"]
    assert 0 < metrics["physical_pieces"] <= 25
    assert metrics["physical_pieces"] + metrics["unused_capacity_physical"] == 25
    with Path(result["artifacts"]["rows_csv"]).open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    for row in rows:
        assert int(row["quantity"]) % int(row["batch_multiple"]) == 0


def test_cluster_count_selects_regions_by_cluster_local_net_need(tmp_path: Path) -> None:
    result = _run(tmp_path, mode="coverage_days", value=30, cluster_count=1)

    assert result["metrics"]["selected_clusters"] == 1
    assert result["cluster_priorities"][0]["region"] == "Центральный"
    assert {row["region"] for row in result["region_totals"]} == {"Центральный"}
    with Path(result["artifacts"]["rows_csv"]).open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert {row["region"] for row in rows} == {"Центральный"}
    assert all(int(row["current_stock"]) == 0 for row in rows)


def test_stock_state_anomaly_excludes_affected_product_region(tmp_path: Path) -> None:
    result = _run(tmp_path, mode="coverage_days", value=30, anomalous=True)

    with Path(result["artifacts"]["rows_csv"]).open(encoding="utf-8", newline="") as handle:
        vendors = {row["vendor_code"] for row in csv.DictReader(handle)}
    assert "ng" not in vendors
    assert "stock_state_anomaly" in Path(result["artifacts"]["controls_csv"]).read_text(encoding="utf-8")


def test_owner_decision_is_local_and_checksum_bound(tmp_path: Path) -> None:
    result = _run(tmp_path, mode="capacity", value=25)
    decision = record_wb_work_plan_decision(
        data_dir=tmp_path / "data",
        plan_run_id=result["run_id"],
        approved=True,
    )

    assert decision["status"] == "owner_approved"
    assert decision["marketplace_write_performed"] is False
    summary = json.loads(Path(result["artifacts"]["summary"]).read_text(encoding="utf-8"))
    assert summary["approval_status"] == "owner_approved"
    assert Path(summary["artifacts"]["owner_decision"]).exists()
