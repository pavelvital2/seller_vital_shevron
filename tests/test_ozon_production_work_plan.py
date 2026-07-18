from __future__ import annotations

import csv
from datetime import date
import json
from pathlib import Path

from openpyxl import load_workbook

from seller_agent.tasks.ozon_production_work_plan import (
    record_ozon_work_plan_decision,
    run_ozon_production_work_plan,
)


class FakeOzonAdapter:
    def fetch_product_list(self, *, visibility: str):  # type: ignore[no-untyped-def]
        return [{"offer_id": "nr", "product_id": 100, "sku": 1000}]

    def fetch_product_info(self, product_ids):  # type: ignore[no-untyped-def]
        return [
            {
                "offer_id": "nr",
                "product_id": 100,
                "sku": 1000,
                "name": "Шеврон нарукавный",
            }
        ]

    def fetch_fbo_clusters(self):  # type: ignore[no-untyped-def]
        return [
            {
                "macrolocal_cluster_id": 1,
                "data": {
                    "macrolocal_cluster": {"name": "Кластер А"},
                    "fulfillments": [{"name": "СКЛАД-А", "warehouse_id": 11}],
                },
            },
            {
                "macrolocal_cluster_id": 2,
                "data": {
                    "macrolocal_cluster": {"name": "Кластер Б"},
                    "fulfillments": [{"name": "СКЛАД-Б", "warehouse_id": 22}],
                },
            },
        ]

    def fetch_stock_on_warehouses(self):  # type: ignore[no-untyped-def]
        return [
            {
                "item_code": "nr",
                "sku": 1000,
                "warehouse_name": "СКЛАД-А",
                "free_to_sell_amount": 100,
                "reserved_amount": 0,
                "promised_amount": 0,
            },
            {
                "item_code": "nr",
                "sku": 1000,
                "warehouse_name": "СКЛАД-Б",
                "free_to_sell_amount": 0,
                "reserved_amount": 2,
                "promised_amount": 3,
            },
        ]

    def fetch_fbo_postings(self, *, since: str, to: str, status: str, limit: int):  # type: ignore[no-untyped-def]
        rows = []
        for cluster, quantity in (("Кластер А", 60), ("Кластер Б", 30)):
            for number in range(quantity):
                rows.append(
                    {
                        "posting_number": f"{cluster}-{number}",
                        "status": "delivered",
                        "created_at": "2026-07-01T10:00:00Z",
                        "financial_data": {"cluster_to": cluster},
                        "products": [{"offer_id": "nr", "sku": 1000, "quantity": 1}],
                    }
                )
        return rows

    def fetch_supply_order_ids(self, *, states):  # type: ignore[no-untyped-def]
        return [10]

    def fetch_supply_orders(self, order_ids):  # type: ignore[no-untyped-def]
        return [
            {
                "order_id": 10,
                "order_tags": {"is_virtual": False},
                "state": "IN_TRANSIT",
                "supplies": [
                    {
                        "state": "IN_TRANSIT",
                        "macrolocal_cluster_id": 2,
                        "bundle_id": "bundle-b",
                    }
                ],
            }
        ]

    def fetch_supply_order_bundle(self, bundle_id):  # type: ignore[no-untyped-def]
        return [{"offer_id": "nr", "sku": 1000, "quantity": 6}]


def _write_catalog(data_dir: Path) -> None:
    path = data_dir / "catalog" / "unified" / "products.csv"
    path.parent.mkdir(parents=True)
    path.write_text(
        "internal_sku,product_name,product_group,pack_qty,ozon_offer_id,ozon_product_id,ozon_sku\n"
        "chev_nr_test,Шеврон нарукавный,chev,1,nr,100,1000\n",
        encoding="utf-8",
    )


def _run(tmp_path: Path, *, mode: str, value: int, cluster_count: int):  # type: ignore[no-untyped-def]
    data_dir = tmp_path / "data"
    _write_catalog(data_dir)
    return run_ozon_production_work_plan(
        data_dir=data_dir,
        run_id=f"ozon_production_work_plan_{mode}_{value}_{cluster_count}",
        mode=mode,
        value=value,
        cluster_count=cluster_count,
        today=date(2026, 7, 18),
        adapter=FakeOzonAdapter(),  # type: ignore[arg-type]
    )


def test_plan_uses_stock_and_inbound_only_inside_same_cluster(tmp_path: Path) -> None:
    result = _run(tmp_path, mode="coverage_days", value=30, cluster_count=1)

    assert result["metrics"]["selected_clusters"] == 1
    assert result["cluster_priorities"][0]["cluster"] == "Кластер Б"
    assert result["cluster_totals"][0]["cluster"] == "Кластер Б"
    workbook_path = Path(result["artifacts"]["report"])
    assert load_workbook(workbook_path, read_only=True).sheetnames == [
        "Артикулы",
        "Ozon кластеры",
        "Контроль",
    ]
    with Path(result["artifacts"]["rows_csv"]).open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 1
    assert rows[0]["cluster"] == "Кластер Б"
    assert int(rows[0]["free_stock_in_cluster"]) == 0
    assert int(rows[0]["confirmed_inbound_to_cluster"]) == 6
    assert int(rows[0]["promised_in_cluster_reference_only"]) == 3
    assert int(rows[0]["quantity"]) == 16


def test_capacity_plan_respects_physical_limit_and_requested_cluster_count(tmp_path: Path) -> None:
    result = _run(tmp_path, mode="capacity", value=10, cluster_count=2)

    assert result["metrics"]["physical_pieces"] == 8
    assert result["metrics"]["unused_capacity_physical"] == 2
    assert result["metrics"]["selected_clusters"] == 1
    assert any("Запрошено кластеров: 2" in warning for warning in result["warnings"])


def test_owner_decision_is_local_and_checksum_bound(tmp_path: Path) -> None:
    result = _run(tmp_path, mode="capacity", value=10, cluster_count=1)
    decision = record_ozon_work_plan_decision(
        data_dir=tmp_path / "data",
        plan_run_id=result["run_id"],
        approved=True,
    )

    assert decision["status"] == "owner_approved"
    assert decision["marketplace_write_performed"] is False
    summary = json.loads(Path(result["artifacts"]["summary"]).read_text(encoding="utf-8"))
    assert summary["approval_status"] == "owner_approved"
    assert Path(summary["artifacts"]["owner_decision"]).exists()
